"""Local deterministic export check; uses no AI, account or payment API."""
import io
import json
from pathlib import Path

import fitz
from pptx import Presentation

from app.engine.qc.visual import pptx_to_pdf
from app.engine.render.renderer import DeckRenderer
from app.engine.style.pptx_analyzer import analyze_pptx
from app.engine.template.builder import build_native
from app.generation.specs import Bullet, SlideSpec


output = Path('/tmp/clastio-footer-smoke')
output.mkdir(exist_ok=True)
sample = Path('/samples/science_ms_sara.pptx')
base, spec = build_native(sample, analyze_pptx(sample))
slides = [SlideSpec(number=n, layout='concept', title=f'Energy transfer {n}', purpose='Explain',
                    bullets=[Bullet(text='Heating transfers energy into a cooler object.')]) for n in (1, 2)]
evidence = []
for enabled in (False, True):
    name = 'enabled' if enabled else 'disabled'
    props = {'author': 'Ms Audit Teacher', 'teacher_footer_text': 'Ms Audit Teacher' if enabled else '',
             'page_numbers': 'true' if enabled else 'false'}
    pptx = DeckRenderer(base, spec).render(slides, core_props=props)
    restored = Presentation(io.BytesIO(pptx))
    for index, slide in enumerate(restored.slides, start=1):
        footers = [shape for shape in slide.shapes if shape.name == 'Teacher name footer']
        numbers = [shape for shape in slide.shapes if shape.name == 'Teacher page number']
        assert len(footers) == len(numbers) == int(enabled)
        if enabled:
            assert footers[0].text == 'Ms Audit Teacher'
            assert numbers[0].text == f'{index} / 2'
            assert footers[0].left > restored.slide_width / 2
            assert footers[0].top > restored.slide_height * .9
    pdf = pptx_to_pdf(pptx)
    document = fitz.open(stream=pdf, filetype='pdf')
    legacy_numbers = []
    for index, page in enumerate(document, start=1):
        text = page.get_text()
        assert ('Ms Audit Teacher' in text) == enabled
        assert (f'{index} / 2' in text) == enabled
        footer_spans = [span['text'].strip() for block in page.get_text('dict')['blocks']
                        for line in block.get('lines', []) for span in line['spans']
                        if span['bbox'][1] >= page.rect.height * .9]
        legacy_numbers.extend({'page': index, 'text': text} for text in footer_spans if text.isdigit())
    (output / f'footer-{name}.pptx').write_bytes(pptx)
    (output / f'footer-{name}.pdf').write_bytes(pdf)
    document[0].get_pixmap(dpi=120).save(str(output / f'footer-{name}.png'))
    evidence.append({'mode': name, 'slides': len(restored.slides), 'pdf_pages': len(document),
                     'teacher_name_and_numbers': 'present' if enabled else 'absent',
                     'legacy_numeric_footer_text': legacy_numbers})
(output / 'evidence.json').write_text(json.dumps(evidence, indent=2))
print(json.dumps(evidence))
assert not any(item['legacy_numeric_footer_text'] for item in evidence), 'A legacy template footer number remains visible'
