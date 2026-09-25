"""Generate realistic sample teacher decks used by tests and demos.

Run from apps/api:  .venv/bin/python ../../samples/make_samples.py
Creates:
  samples/science_ms_sara.pptx      text-box based design, header band, logo, cover background
  samples/maths_mr_raj.pptx         placeholder-based design, custom master background, serif headings
  samples/primary_colourful.pptx    background images, rounded cards, playful fonts
  samples/science_ms_sara.pdf       PDF export of the science deck (for the PDF analyzer)
"""

from __future__ import annotations

import io
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "apps" / "api"))
from app.engine.pptx_xml import write_theme  # noqa: E402


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def png(img: Image.Image) -> io.BytesIO:
    b = io.BytesIO()
    img.save(b, "PNG")
    b.seek(0)
    return b


def logo_image(color=(15, 118, 110)) -> io.BytesIO:
    img = Image.new("RGBA", (256, 256), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((8, 8, 248, 248), fill=color + (255,))
    d.polygon([(128, 50), (190, 150), (128, 210), (66, 150)], fill=(255, 255, 255, 255))
    d.line([(128, 70), (128, 205)], fill=color + (255,), width=8)
    return png(img)


def photo_like(seed: int, w=960, h=720) -> io.BytesIO:
    import random

    rnd = random.Random(seed)
    img = Image.new("RGB", (w, h), (rnd.randint(120, 200), rnd.randint(150, 220), rnd.randint(150, 230)))
    d = ImageDraw.Draw(img)
    for _ in range(18):
        x, y = rnd.randint(0, w), rnd.randint(0, h)
        r = rnd.randint(40, 200)
        col = (rnd.randint(40, 255), rnd.randint(80, 255), rnd.randint(40, 200))
        d.ellipse((x - r, y - r, x + r, y + r), fill=col)
    return png(img.filter(ImageFilter.GaussianBlur(6)))


def gradient_bg(c1, c2, w=1600, h=900) -> io.BytesIO:
    img = Image.new("RGB", (w, h))
    d = ImageDraw.Draw(img)
    for y in range(h):
        t = y / h
        d.line([(0, y), (w, y)], fill=tuple(int(c1[i] * (1 - t) + c2[i] * t) for i in range(3)))
    for i in range(6):
        d.ellipse((1300 - i * 40, 650 - i * 30, 1700, 1000), outline=(255, 255, 255), width=2)
    return png(img)


def textbox(slide, x, y, w, h, paragraphs, size=20, color="#1F2937", bold=False, font=None, align=None,
            bullets=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " + text) if bullets else text
        if align:
            p.alignment = align
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = rgb(color)
            if font:
                r.font.name = font
        p.space_after = Pt(8)
    return tb


# --------------------------------------------------------------------------- science deck


def science_deck(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    write_theme(prs, colors={"dk1": "#1F2937", "lt1": "#FFFFFF", "dk2": "#0F766E", "lt2": "#ECFDF5",
                             "accent1": "#0F766E", "accent2": "#F59E0B", "accent3": "#10B981",
                             "accent4": "#0EA5E9", "accent5": "#6366F1", "accent6": "#EF4444"},
                fonts={"major": "Montserrat", "minor": "Open Sans"})
    blank = prs.slide_layouts[6]
    logo = logo_image()

    def chrome(slide, n):
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2))
        band.fill.solid()
        band.fill.fore_color.rgb = rgb("#0F766E")
        band.line.fill.background()
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.2), prs.slide_width, Inches(0.08))
        accent.fill.solid()
        accent.fill.fore_color.rgb = rgb("#F59E0B")
        accent.line.fill.background()
        logo.seek(0)
        slide.shapes.add_picture(logo, Inches(12.35), Inches(6.55), Inches(0.7), Inches(0.7))
        textbox(slide, 0.5, 6.85, 5, 0.4, ["Grade 8 Science  |  Ms Sara"], size=11, color="#6B7280")
        textbox(slide, 11.3, 6.85, 0.9, 0.4, [str(n)], size=11, color="#6B7280", align=PP_ALIGN.RIGHT)

    def title(slide, text):
        textbox(slide, 0.5, 0.22, 11.5, 0.8, [text], size=32, color="#FFFFFF", bold=True, font="Montserrat")

    # cover
    s = prs.slides.add_slide(blank)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb("#0F766E")
    deco = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1.5), Inches(6), Inches(6))
    deco.fill.solid()
    deco.fill.fore_color.rgb = rgb("#115E59")
    deco.line.fill.background()
    textbox(s, 0.9, 2.4, 9, 1.4, ["States of Matter"], size=48, color="#FFFFFF", bold=True, font="Montserrat")
    textbox(s, 0.9, 3.8, 9, 0.8, ["Grade 8 Science  •  Lesson 1"], size=22, color="#F59E0B")
    logo.seek(0)
    s.shapes.add_picture(logo, Inches(0.9), Inches(5.6), Inches(1.0), Inches(1.0))

    content = [
        ("Learning objectives", ["Describe the three states of matter", "Explain particle arrangement in each state",
                                 "Use the particle model to explain changes of state"]),
        ("What is matter?", ["Matter is anything that has mass and takes up space",
                             "Everything around us is made of tiny particles",
                             "Particles are always moving", "How they move decides the state"]),
    ]
    n = 2
    for t, bl in content:
        s = prs.slides.add_slide(blank)
        chrome(s, n)
        title(s, t)
        textbox(s, 0.7, 1.7, 11.5, 4.8, bl, size=22, bullets=True)
        n += 1

    for t, bl, seed in [("Solids", ["Particles packed closely in a fixed pattern", "Vibrate in place",
                                    "Fixed shape and volume"], 3),
                        ("Liquids", ["Particles close but can slide past each other", "Take the shape of the container",
                                     "Fixed volume"], 4),
                        ("Gases", ["Particles far apart and fast moving", "Fill any container", "Easy to compress"], 5)]:
        s = prs.slides.add_slide(blank)
        chrome(s, n)
        title(s, t)
        textbox(s, 0.7, 1.7, 6.3, 4.8, bl, size=22, bullets=True)
        s.shapes.add_picture(photo_like(seed), Inches(7.4), Inches(1.75), Inches(5.3), Inches(4.0))
        n += 1

    s = prs.slides.add_slide(blank)
    chrome(s, n)
    title(s, "Compare: solid vs liquid")
    textbox(s, 0.7, 1.7, 5.8, 0.6, ["Solid"], size=24, color="#0F766E", bold=True)
    textbox(s, 0.7, 2.3, 5.8, 4, ["Fixed shape", "Particles vibrate", "Hard to compress"], size=20, bullets=True)
    textbox(s, 6.9, 1.7, 5.8, 0.6, ["Liquid"], size=24, color="#0F766E", bold=True)
    textbox(s, 6.9, 2.3, 5.8, 4, ["Flows", "Particles slide", "Hard to compress"], size=20, bullets=True)
    n += 1

    s = prs.slides.add_slide(blank)
    chrome(s, n)
    title(s, "Quick quiz")
    textbox(s, 0.7, 1.7, 11.5, 1, ["Which state of matter has a fixed shape?"], size=24, bold=True)
    for i, opt in enumerate(["A) Solid", "B) Liquid", "C) Gas", "D) Plasma"]):
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7 + (i % 2) * 6),
                                 Inches(3 + (i // 2) * 1.5), Inches(5.6), Inches(1.2))
        box.fill.solid()
        box.fill.fore_color.rgb = rgb("#ECFDF5")
        box.line.color.rgb = rgb("#0F766E")
        box.text_frame.text = opt
        for r in box.text_frame.paragraphs[0].runs:
            r.font.size = Pt(20)
            r.font.color.rgb = rgb("#1F2937")
    n += 1

    s = prs.slides.add_slide(blank)
    chrome(s, n)
    title(s, "Summary")
    textbox(s, 0.7, 1.7, 11.5, 4.8, ["Matter is made of particles", "Solids, liquids and gases differ in particle "
                                     "arrangement and movement", "Heating and cooling change states"],
            size=22, bullets=True)
    prs.save(path)


# --------------------------------------------------------------------------- maths deck (placeholders)


def maths_deck(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    write_theme(prs, colors={"dk1": "#111827", "lt1": "#FFFFFF", "dk2": "#1E3A8A", "lt2": "#FFFBF5",
                             "accent1": "#1E3A8A", "accent2": "#F97316", "accent3": "#FBBF24",
                             "accent4": "#14B8A6", "accent5": "#8B5CF6", "accent6": "#EC4899"},
                fonts={"major": "Georgia", "minor": "Calibri"})
    master = prs.slide_master
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = rgb("#FFFBF5")
    # scale placeholders for 16:9
    own_xfrm = [ph for layout in prs.slide_layouts for ph in layout.placeholders if ph._element.spPr.xfrm is not None]
    for ph in list(master.placeholders) + own_xfrm:  # layouts without xfrm inherit the scaled master
        left, top, width, height = ph.left, ph.top, ph.width, ph.height
        if width and height:
            ph.left, ph.top = int(left * 13.333 / 10), top
            ph.width, ph.height = int(width * 13.333 / 10), height
    # coral underline under titles on the master
    tmp = prs.slides.add_slide(prs.slide_layouts[6])
    line = tmp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.55), Inches(2.0), Inches(0.07))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb("#F97316")
    line.line.fill.background()
    master.shapes._spTree.append(line._element)  # move onto the master
    from app.engine.pptx_xml import delete_slide

    delete_slide(prs, 0)

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Fractions: Adding and Subtracting"
    s.placeholders[1].text = "Grade 6 Mathematics • Mr Raj"
    for t, bl in [("Learning objectives", ["Add fractions with the same denominator",
                                            "Find a common denominator", "Subtract mixed numbers"]),
                  ("Recap: what is a fraction?", ["A fraction shows part of a whole", "Numerator: parts we have",
                                                   "Denominator: equal parts in the whole"]),
                  ("Same denominators", ["Add the numerators", "Keep the denominator", "Simplify if possible"]),
                  ("Different denominators", ["Find the lowest common multiple", "Convert both fractions",
                                              "Then add the numerators"]),
                  ("Practice", ["1/4 + 2/4 = ?", "2/3 + 1/6 = ?", "3/5 − 1/10 = ?"]),
                  ("Summary", ["Same denominator: add numerators", "Different: find a common denominator first"])]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        tf = s.placeholders[1].text_frame
        for i, b in enumerate(bl):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
    s = prs.slides.add_slide(prs.slide_layouts[3])
    s.shapes.title.text = "Compare methods"
    s.placeholders[1].text = "Drawing fraction bars"
    s.placeholders[2].text = "Using common denominators"
    s = prs.slides.add_slide(prs.slide_layouts[2])
    s.shapes.title.text = "Part 2: Mixed numbers"
    prs.save(path)


# --------------------------------------------------------------------------- primary deck


def primary_deck(path: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    write_theme(prs, colors={"dk1": "#3B0764", "lt1": "#FFFFFF", "dk2": "#7C3AED", "lt2": "#FDF4FF",
                             "accent1": "#7C3AED", "accent2": "#F472B6", "accent3": "#22C55E",
                             "accent4": "#FACC15", "accent5": "#38BDF8", "accent6": "#FB923C"},
                fonts={"major": "Comic Sans MS", "minor": "Verdana"})
    blank = prs.slide_layouts[6]
    bgimg = gradient_bg((253, 244, 255), (224, 231, 255))
    for i, (t, bl) in enumerate([("Animals and Their Homes", []),
                                 ("Where do animals live?", ["Forests", "Oceans", "Deserts", "Polar lands"]),
                                 ("Desert animals", ["Camels store fat in their humps", "Oryx can go without water",
                                                     "Desert foxes have big ears"]),
                                 ("Ocean animals", ["Dolphins breathe air", "Turtles lay eggs on beaches"]),
                                 ("Let's think!", ["Why does a camel have long eyelashes?"])]):
        s = prs.slides.add_slide(blank)
        bgimg.seek(0)
        pic = s.shapes.add_picture(bgimg, 0, 0, prs.slide_width, prs.slide_height)
        s.shapes._spTree.remove(pic._element)
        s.shapes._spTree.insert(2, pic._element)
        if i == 0:
            textbox(s, 1.5, 2.6, 10.3, 1.6, [t], size=54, color="#7C3AED", bold=True, font="Comic Sans MS",
                    align=PP_ALIGN.CENTER)
            textbox(s, 1.5, 4.2, 10.3, 0.8, ["Grade 2 Science"], size=26, color="#F472B6", align=PP_ALIGN.CENTER)
            continue
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = rgb("#7C3AED")
        card.line.fill.background()
        textbox(s, 0.9, 0.55, 11.5, 0.9, [t], size=34, color="#FFFFFF", bold=True, font="Comic Sans MS")
        body = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.9), Inches(12.1), Inches(5.0))
        body.fill.solid()
        body.fill.fore_color.rgb = rgb("#FFFFFF")
        body.line.color.rgb = rgb("#F472B6")
        textbox(s, 1.0, 2.2, 11.3, 4.5, bl, size=26, color="#3B0764", bullets=True)
    prs.save(path)


def to_pdf(pptx: Path) -> None:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                   check=True, capture_output=True, timeout=180)


if __name__ == "__main__":
    science_deck(HERE / "science_ms_sara.pptx")
    maths_deck(HERE / "maths_mr_raj.pptx")
    primary_deck(HERE / "primary_colourful.pptx")
    to_pdf(HERE / "science_ms_sara.pptx")
    print("samples written to", HERE)
