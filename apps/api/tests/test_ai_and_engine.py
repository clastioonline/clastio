from __future__ import annotations

import io
import json
from types import SimpleNamespace

import httpx
import pytest
from pptx import Presentation

from app.ai.base import AIError, AIRequest, ChatMessage, Usage
from app.ai.pricing import cost_usd
from app.ai.schema_utils import extract_json, strict_schema
from app.ai.service import AIService, Route
from app.engine.render.renderer import DeckRenderer
from app.engine.render.textfit import Para, fit
from app.engine.style.pdf_analyzer import analyze_pdf
from app.engine.style.pptx_analyzer import analyze_pptx
from app.engine.template.builder import build_builtin, build_native
from app.generation.budgets import compute_budgets
from app.generation.pipeline import enforce_budgets
from app.generation.specs import Bullet, CoursePlan, LessonDeck, SlideSpec
from tests.conftest import SAMPLES

# --------------------------------------------------------------------------- style analysis


def test_pptx_analyzer_extracts_design_system():
    a = analyze_pptx(SAMPLES / "science_ms_sara.pptx")
    assert a["colors"]["primary"] == "#0F766E"
    assert a["fonts"]["heading"]["family"] == "Montserrat"
    assert a["fonts"]["body"]["family"] == "Open Sans"
    kinds = {i["kind"] for i in a["decorations"]["content"]["items"]}
    assert {"autoshape", "picture", "text", "slide_number"} <= kinds
    assert a["zones"]["image_side"] == "right"
    labels = {x["key"] for x in a["layouts_found"]}
    assert {"cover", "image_text", "quiz", "summary"} <= labels


def test_pptx_analyzer_placeholder_deck():
    a = analyze_pptx(SAMPLES / "maths_mr_raj.pptx")
    assert a["fonts"]["heading"]["family"] == "Georgia"
    assert a["colors"]["background"] == "#FFFBF5"
    assert a["layout_map"]["section_layout_index"] is not None


def test_pdf_analyzer():
    a = analyze_pdf(SAMPLES / "science_ms_sara.pdf")
    assert a["colors"]["primary"] == "#0F766E"
    assert a["stats"]["text_layer"] is True
    assert any(i["kind"] == "slide_number" for i in a["decorations"]["content"]["items"])


# --------------------------------------------------------------------------- rendering


def all_layouts_deck() -> list[SlideSpec]:
    from app.engine.template.preview import PREVIEW_SLIDES

    extra = [
        SlideSpec(number=7, layout="summary", purpose="s", title="Summary",
                  bullets=[Bullet(text="One"), Bullet(text="Two"), Bullet(text="Three")]),
        SlideSpec(number=8, layout="section", purpose="s", title="Part 2", subtitle="Next steps"),
    ]
    return list(PREVIEW_SLIDES) + extra


@pytest.mark.parametrize("sample", ["science_ms_sara.pptx", "maths_mr_raj.pptx", "primary_colourful.pptx"])
def test_native_templates_render_every_layout(sample):
    a = analyze_pptx(SAMPLES / sample)
    base, spec = build_native(SAMPLES / sample, a)
    r = DeckRenderer(base, spec)
    data = r.render(all_layouts_deck(), core_props={"title": "t", "author": "a"})
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == len(all_layouts_deck())  # donor slides removed
    assert not [rep for rep in r.reports if rep.overflow]
    ids = [[int(el.get("id")) for el in s._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")]
           for s in prs.slides]
    assert all(len(x) == len(set(x)) for x in ids), "duplicate shape ids would trigger PowerPoint repair"


def test_arabic_rtl_rendering():
    base, spec = build_builtin("clean-blue")
    slides = [SlideSpec(number=1, layout="concept", purpose="c", title="البناء الضوئي", language="ar",
                        bullets=[Bullet(text="تصنع النباتات غذاءها باستخدام ضوء الشمس"),
                                 Bullet(text="يحدث في البلاستيدات الخضراء")])]
    data = DeckRenderer(base, spec, language="ar").render(slides)
    prs = Presentation(io.BytesIO(data))
    xml = prs.slides[0]._element.xml
    assert 'rtl="1"' in xml and "Noto Sans Arabic" in xml


def test_text_fit_shrinks_then_reports_overflow():
    paras = [Para("A short line")]
    assert fit(paras, "Open Sans", 400, 100, 28, 16).fits
    long = [Para("word " * 400)]
    r = fit(long, "Open Sans", 300, 80, 28, 16)
    assert not r.fits and r.size_pt == 16


def test_overflow_is_repaired_by_budget_enforcement():
    base, spec = build_builtin("clean-blue")
    budgets = compute_budgets(spec)
    long = SlideSpec(number=1, layout="concept", purpose="c", title="A very long title " * 3,
                     bullets=[Bullet(text="This bullet is far too long for a slide " * 6) for _ in range(9)])
    fixed = enforce_budgets(long.model_copy(deep=True), budgets, strict=True)
    r = DeckRenderer(base, spec)
    r.render([fixed])
    assert not r.reports[0].overflow
    assert len(fixed.bullets) <= budgets["bullets_max"]


# --------------------------------------------------------------------------- AI layer


def test_strict_schema_is_provider_safe():
    s = strict_schema(LessonDeck)
    assert "$defs" not in json.dumps(s)

    def walk(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node["additionalProperties"] is False
                assert set(node["required"]) == set(node.get("properties", {}))
            assert "minimum" not in node and "title" not in node
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)

    walk(s)


def test_extract_json_tolerates_fences():
    assert extract_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert extract_json('Here you go: {"a": [1,2]} thanks') == {"a": [1, 2]}


def test_cost_calculation():
    assert cost_usd("claude-opus-5", Usage(input_tokens=1_000_000, output_tokens=0)) == 5.0
    assert cost_usd("claude-haiku-4-5", Usage(input_tokens=0, output_tokens=1_000_000)) == 5.0
    assert cost_usd("offline", Usage(input_tokens=10**6)) == 0.0
    assert cost_usd("claude-sonnet-5", Usage(output_tokens=10**6), {"tokens": {"claude-sonnet-5": [1, 1, 0]}}) == 1.0


class _FakeStream:
    def __init__(self, message):
        self.message = message

    async def __aenter__(self):
        return self

    async def __aexit__(self, *a):
        return False

    async def get_final_message(self):
        return self.message


async def test_anthropic_adapter_structured_output(monkeypatch):
    from app.ai.anthropic_provider import AnthropicProvider

    plan = {"title": "T", "overview": "o", "big_idea": "b", "learning_outcomes": [{"code": None, "text": "x"}],
            "prerequisites": [], "diagnostic_questions": [], "lectures": [], "revision_strategy": "r",
            "assessment_strategy": "a", "cross_curricular_links": [], "local_context_links": []}
    msg = SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=json.dumps(plan))],
                          usage=SimpleNamespace(input_tokens=1200, output_tokens=300, cache_read_input_tokens=800,
                                                cache_creation_input_tokens=0))
    captured = {}

    class FakeMessages:
        def stream(self, **params):
            captured.update(params)
            return _FakeStream(msg)

    provider = AnthropicProvider.__new__(AnthropicProvider)
    provider._client = SimpleNamespace(messages=FakeMessages())
    req = AIRequest(task="course_plan", system="sys", messages=[ChatMessage("user", "plan it")], effort="medium")
    res = await provider.generate_structured("claude-opus-5", req, CoursePlan)
    assert res.data.title == "T"
    assert res.usage.cached_tokens == 800
    assert captured["output_config"]["format"]["type"] == "json_schema"
    assert captured["output_config"]["effort"] == "medium"
    assert captured["extra_body"] == {"fallbacks": "default"}
    assert captured["system"][0]["cache_control"] == {"type": "ephemeral"}
    # Haiku 4.5 does not accept effort; the adapter must omit it
    captured.clear()
    await provider.generate_structured("claude-haiku-4-5", req, CoursePlan)
    assert "effort" not in captured.get("output_config", {})


async def test_anthropic_refusal_raises(monkeypatch):
    from app.ai.anthropic_provider import AnthropicProvider
    from app.ai.base import AIRefusal

    msg = SimpleNamespace(stop_reason="refusal", stop_details=SimpleNamespace(category="bio"), content=[],
                          usage=SimpleNamespace(input_tokens=1, output_tokens=0))

    class FakeMessages:
        def stream(self, **params):
            return _FakeStream(msg)

    provider = AnthropicProvider.__new__(AnthropicProvider)
    provider._client = SimpleNamespace(messages=FakeMessages())
    with pytest.raises(AIRefusal):
        await provider.generate_text("claude-sonnet-5", AIRequest(task="t", system="s",
                                                                   messages=[ChatMessage("user", "x")]))


async def test_router_falls_back_to_next_provider(monkeypatch):
    import app.generation.offline  # noqa: F401

    svc = AIService()

    class Broken:
        name = "anthropic"

        def available(self):
            return True

        async def generate_structured(self, model, req, schema):
            raise AIError("boom", retryable=False, provider="anthropic")

    svc.providers["anthropic"] = Broken()
    monkeypatch.setattr(type(svc), "live_providers", property(lambda self: ["anthropic"]))

    async def routes(tier):
        return [Route("anthropic", "claude-sonnet-5"), Route("offline", "offline")]

    monkeypatch.setattr(svc, "routes", routes)
    plan = await svc.structured(task="course_plan", tier="planning", system="s", prompt="p", schema=CoursePlan,
                                offline_context={"topic": "volcanoes", "num_lectures": 2})
    assert len(plan.lectures) == 2


async def test_openverse_search_records_license():
    from app.services.assets import search_openverse

    png = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (800, 500), "green").save(png, "PNG")

    def handler(request: httpx.Request):
        if "api.openverse.org" in str(request.url):
            return httpx.Response(200, json={"results": [{
                "url": "https://images.example/leaf.png", "width": 800, "title": "Leaf", "creator": "Ana",
                "license": "by", "license_version": "4.0", "foreign_landing_url": "https://example/leaf"}]})
        return httpx.Response(200, content=png.getvalue())

    async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
        found = await search_openverse("leaf", client)
    assert found and found["license"] == "CC BY 4.0" and "Ana" in found["attribution"]
