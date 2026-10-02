from __future__ import annotations

import io
import uuid

import pytest
from pptx import Presentation
from pydantic import ValidationError

from app.ai.service import AIService
from app.engine.render.renderer import DeckRenderer
from app.engine.template.builder import build_builtin
from app.generation.specs import ChartData, SlideSpec
from app.services import assistant


@pytest.mark.parametrize('text', ['Book me a hotel tomorrow', 'Act like normal ChatGPT', 'Buy me a phone', 'hello'])
async def test_obvious_unrelated_messages_do_not_call_ai(monkeypatch, text):
    def unexpected():
        raise AssertionError('AI should not be called')
    monkeypatch.setattr(assistant, 'get_ai', unexpected)
    result = await assistant.classify(text, uuid.uuid4())
    assert result.relevance != 'teaching'


async def test_scope_check_failure_cannot_fall_through_to_action(monkeypatch):
    class Broken:
        async def structured(self, **kwargs):
            raise RuntimeError('unavailable')
    monkeypatch.setattr(assistant, 'get_ai', lambda: Broken())
    result = await assistant.classify('Create a quiz on booking my personal holiday', uuid.uuid4())
    assert result.relevance == 'needs_context' and result.intent == 'general'


async def test_unrelated_request_never_runs_action_or_content(teacher, monkeypatch):
    from app.core.db import get_sessionmaker
    from app.models import User
    async def unexpected(*args, **kwargs):
        raise AssertionError('Action or context must not run')
    monkeypatch.setattr(assistant, 'run_action', unexpected)
    monkeypatch.setattr(assistant, 'build_context', unexpected)
    async with get_sessionmaker()() as db:
        user = await db.get(User, uuid.UUID(teacher['id']))
        events = [event async for event in assistant.converse(db, user, 'Book me a hotel tomorrow', None)]
    assert any(event['event'] == 'token' and event['data']['text'] == assistant.REDIRECT for event in events)
    assert not any(event['event'] == 'action' for event in events)


@pytest.mark.parametrize('kind', ['bar', 'line', 'pie'])
def test_charts_are_editable_with_original_data(kind):
    base, template = build_builtin('clean-blue')
    data = ChartData(kind=kind, categories=['Red', 'Blue'], series=[{'name': 'Votes', 'values': [3, 5]}],
                     source='Classroom poll', unit='Students')
    slide = SlideSpec(number=1, layout='chart', purpose='Compare', title='Our class votes', chart=data)
    renderer = DeckRenderer(base, template)
    output = renderer.render([slide])
    prs = Presentation(io.BytesIO(output))
    chart = next(shape.chart for shape in prs.slides[0].shapes if shape.has_chart)
    assert list(chart.series[0].values) == [3, 5]
    assert chart.part.chart_workbook.xlsx_part.blob
    assert any('Classroom poll' in shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert not any(report.overflow for report in renderer.reports)


@pytest.mark.parametrize('values', [[1], [1, float('nan')]])
def test_invalid_chart_data_rejected(values):
    with pytest.raises(ValidationError):
        ChartData(categories=['A', 'B'], series=[{'name': 'X', 'values': values}], source='Test data')


async def test_shared_cache_survives_new_service_and_failure_is_optional():
    class RedisStub:
        def __init__(self):
            self.items = {}
        async def get(self, key):
            return self.items.get(key)
        async def set(self, key, value, ex):
            assert ex == 86400
            self.items[key] = value
    redis = RedisStub()
    first, second = AIService(), AIService()
    first._result_cache = second._result_cache = redis
    await first._store_result('teacher-a', {'title': 'Fractions'})
    assert await second._cached_result('teacher-a') == {'title': 'Fractions'}
    assert await second._cached_result('teacher-b') is None


async def test_keywords_cannot_bypass_semantic_scope_classifier(monkeypatch):
    class Classifier:
        async def structured(self, **kwargs):
            assert 'relevance' in kwargs['system']
            return assistant.Intent(intent='create_quiz', relevance='out_of_scope')
    monkeypatch.setattr(assistant, 'get_ai', lambda: Classifier())
    result = await assistant.classify('Create a quiz to choose my personal hotel', uuid.uuid4())
    assert result.relevance == 'out_of_scope'


def test_source_snippets_are_relevant_and_valid_json():
    import json

    from app.services.styles import source_context
    pages = [{'number': n, 'text': ('Fractions ' if n == 15 else 'Other topic ') + 'x' * 2200,
              'notes': '', 'pictures': []} for n in range(1, 21)]
    context = source_context({'source_content': pages}, 'Fractions')
    references = json.loads(context.split('\n')[-1])
    assert any(page['slide'] == 15 for page in references)
    assert len(context.split('\n')[-1]) <= 18000
