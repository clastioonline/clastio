"""Verify added browser objects survive as editable PowerPoint shapes."""
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.engine.render.manual import apply_manual, editable_objects
from app.generation.specs import ManualObjectEdit


def test_custom_objects_export_and_keep_stable_ids():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    edits = {
        'custom-title': ManualObjectEdit(object_type='text', text='Teacher example', x=.1, y=.1,
            width=.6, height=.15, font_size=28, color='#112233', bold=True, italic=True),
        'custom-shape': ManualObjectEdit(object_type='ellipse', x=.2, y=.4, width=.3, height=.3,
            fill='#aabbcc', rotation=15, layer='back'),
    }
    apply_manual(slide, edits, deck.slide_width, deck.slide_height)
    assert len(slide.shapes) == 2
    assert slide.shapes[0].shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    text = slide.shapes[1]
    assert text.text == 'Teacher example'
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 28
    assert text.text_frame.paragraphs[0].runs[0].font.italic
    output = io.BytesIO()
    deck.save(output)
    assert {item['id'] for item in editable_objects(output.getvalue())[0]} == {'custom-title', 'custom-shape'}
    restored = Presentation(io.BytesIO(output.getvalue()))
    apply_manual(restored.slides[0], edits, restored.slide_width, restored.slide_height)
    assert len(restored.slides[0].shapes) == 2  # no duplicates on the next rebuild
    edits['custom-shape'] = ManualObjectEdit(hidden=True)
    apply_manual(restored.slides[0], edits, restored.slide_width, restored.slide_height)
    assert len(restored.slides[0].shapes) == 1
