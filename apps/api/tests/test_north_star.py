"""The core product promise, end to end through the HTTP API:

upload an old deck -> style + template -> "Photosynthesis, Grade 8, 5 lectures x 10 slides" ->
5 connected lessons, 50 editable slides in the teacher's design, notes, previews, downloads, QC clean.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from tests.conftest import SAMPLES


async def upload_style(client, headers, sample: str) -> str:
    data = (SAMPLES / sample).read_bytes()
    r = await client.post("/api/v1/uploads", headers=headers, data={"kind": "style", "name": "My style"},
                          files={"file": (sample, data, "application/octet-stream")})
    assert r.status_code == 200, r.text
    f = r.json()["file"]
    assert f["status"] == "ready", f
    detail = (await client.get(f"/api/v1/uploads/{f['id']}", headers=headers)).json()
    assert detail["template_id"]
    return detail["template_id"]


async def test_style_upload_creates_template_with_previews(client, teacher):
    tid = await upload_style(client, teacher["headers"], "science_ms_sara.pptx")
    t = (await client.get(f"/api/v1/templates/{tid}", headers=teacher["headers"])).json()
    assert t["mode"] == "native"
    assert t["colors"]["primary"] == "#0F766E"
    assert t["fonts"]["heading"] == "Montserrat"
    assert len(t["previews"]) == 6
    assert t["analysis"]["stats"]["decoration_count"] >= 4
    prof = (await client.get("/api/v1/me/profile", headers=teacher["headers"])).json()
    assert prof["default_template_id"] == tid  # first upload becomes the default
    mem = (await client.get("/api/v1/memory", headers=teacher["headers"])).json()
    inferred = {p["key"]: p for p in mem["preferences"]}
    assert "words_per_bullet" in inferred and inferred["words_per_bullet"]["confirmed"] is False


async def test_pdf_upload_reconstructs_template(client, teacher):
    tid = await upload_style(client, teacher["headers"], "science_ms_sara.pdf")
    t = (await client.get(f"/api/v1/templates/{tid}", headers=teacher["headers"])).json()
    assert t["mode"] == "reconstructed"
    assert t["colors"]["primary"] == "#0F766E"


async def test_north_star_photosynthesis_course(client, teacher):
    h = teacher["headers"]
    tid = await upload_style(client, h, "science_ms_sara.pptx")
    r = await client.post("/api/v1/courses", headers=h, json={
        "topic": "Photosynthesis", "grade": "8", "subject": "Science", "num_lectures": 5, "slides_per_lecture": 10,
        "template_id": tid, "auto_generate": True})
    assert r.status_code == 200, r.text
    project_id = r.json()["course"]["project_id"]
    proj = (await client.get(f"/api/v1/projects/{project_id}", headers=h)).json()
    course = proj["course"]
    assert course["status"] == "ready", course
    plan = course["plan"]
    assert len(plan["lectures"]) == 5
    # Progression: no key concept is introduced twice.
    concepts = [c.lower() for lec in plan["lectures"] for c in lec["key_concepts"]]
    assert len(concepts) == len(set(concepts))
    assert len(proj["lessons"]) == 5 and all(les["has_pptx"] for les in proj["lessons"])

    total_slides = 0
    for les in proj["lessons"]:
        d = (await client.get(f"/api/v1/lessons/{les['id']}", headers=h)).json()
        assert len(d["slides"]) == 10
        assert d["slides"][0]["spec"]["layout"] == "cover"
        assert all(s["preview"] for s in d["slides"])
        assert not any(s["qc"]["overflow"] for s in d["slides"])
        assert d["lesson"]["plan"]["duration_minutes"] == 45
        assert d["lesson"]["plan"]["differentiation"]["eal"]
        # QC: renderer says everything fits, visual QC found no errors
        assert all(not r_["overflow"] for r_ in d["lesson"]["qc"]["render"])
        visual_errors = [i for issues in d["lesson"]["qc"]["visual"].values() for i in issues
                         if i["severity"] == "error"]
        assert visual_errors == []
        pptx = await client.get(d["downloads"]["pptx"])
        assert pptx.status_code == 200
        prs = Presentation(io.BytesIO(pptx.content))
        assert len(prs.slides) == 10
        total_slides += len(prs.slides)
        # Teacher's design system survives: the logo picture and header band are on content slides.
        content_slide = prs.slides[2]
        pics = [s for s in content_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        assert pics, "teacher logo missing"
        fills = {str(s.fill.fore_color.rgb) for s in content_slide.shapes
                 if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.fill.type == 1}
        assert "0F766E" in fills
        # Everything is editable text (no flattened slides) and notes exist.
        assert any(s.has_text_frame and s.text_frame.text.strip() for s in content_slide.shapes)
        assert content_slide.has_notes_slide and content_slide.notes_slide.notes_text_frame.text.strip()
        # Fonts are the teacher's.
        fonts = {r_.font.name for s in content_slide.shapes if s.has_text_frame
                 for p in s.text_frame.paragraphs for r_ in p.runs if r_.font.name}
        assert "Open Sans" in fonts or "Montserrat" in fonts
    assert total_slides == 50
    # Lesson summaries recorded in teacher memory
    mem = (await client.get("/api/v1/memory?kind=lesson_summary", headers=h)).json()
    assert len(mem["items"]) >= 5


async def test_slide_edit_and_regenerate(client, teacher):
    h = teacher["headers"]
    r = await client.post("/api/v1/courses", headers=h, json={
        "topic": "Magnets", "grade": "6", "subject": "Science", "num_lectures": 1, "slides_per_lecture": 8,
        "auto_generate": True})
    project_id = r.json()["course"]["project_id"]
    lesson_id = (await client.get(f"/api/v1/projects/{project_id}", headers=h)).json()["lessons"][0]["id"]
    before = (await client.get(f"/api/v1/lessons/{lesson_id}", headers=h)).json()
    slide3 = before["slides"][2]
    r = await client.patch(f"/api/v1/lessons/{lesson_id}/slides/3", headers=h,
                           json={"spec": {"title": "Magnets around us"}})
    assert r.status_code == 200
    after = (await client.get(f"/api/v1/lessons/{lesson_id}", headers=h)).json()
    assert after["slides"][2]["spec"]["title"] == "Magnets around us"
    assert after["lesson"]["version"] == before["lesson"]["version"] + 1
    r = await client.post(f"/api/v1/lessons/{lesson_id}/slides/3/regenerate", headers=h, json={"action": "simpler"})
    assert r.status_code == 200
    versions = (await client.get(f"/api/v1/lessons/{lesson_id}/slides/3/versions", headers=h)).json()["items"]
    assert len(versions) >= 3
    r = await client.post(f"/api/v1/lessons/{lesson_id}/slides/3/restore", headers=h,
                          json={"version": versions[-1]["version"]})
    assert r.status_code == 200
    restored = (await client.get(f"/api/v1/lessons/{lesson_id}", headers=h)).json()
    assert restored["slides"][2]["spec"]["title"] == slide3["spec"]["title"]


async def test_documents_worksheet_quiz_homework_lesson_plan(client, teacher):
    h = teacher["headers"]
    r = await client.post("/api/v1/courses", headers=h, json={
        "topic": "Fractions", "grade": "6", "subject": "Mathematics", "num_lectures": 1, "slides_per_lecture": 8,
        "auto_generate": True})
    project_id = r.json()["course"]["project_id"]
    lesson_id = (await client.get(f"/api/v1/projects/{project_id}", headers=h)).json()["lessons"][0]["id"]
    for kind in ("worksheet", "quiz", "homework", "lesson_plan", "teacher_guide"):
        r = await client.post("/api/v1/documents", headers=h, json={"kind": kind, "lesson_id": lesson_id,
                                                                     "num_questions": 10})
        assert r.status_code == 200, r.text
    docs = {d["kind"]: d for d in (await client.get(f"/api/v1/documents?lesson_id={lesson_id}", headers=h)
                                   ).json()["items"]}
    assert docs["worksheet"]["status"] == "ready"
    assert {"docx", "pdf", "key_docx", "key_pdf", "csv", "gift"} <= set(docs["worksheet"]["files"])
    assert {"docx", "key_docx"} <= set(docs["homework"]["files"])
    assert "docx" in docs["lesson_plan"]["files"] and "pdf" in docs["lesson_plan"]["files"]
    csv = await client.get(docs["quiz"]["files"]["csv"])
    assert csv.status_code == 200 and b"Question" in csv.content
    bank = (await client.get("/api/v1/questions", headers=h)).json()["items"]
    assert len(bank) >= 10


async def test_free_plan_limits(client):
    from tests.conftest import make_user

    u = await make_user(client, plan="free")
    r = await client.post("/api/v1/courses", headers=u["headers"], json={
        "topic": "Photosynthesis", "grade": "8", "subject": "Science", "num_lectures": 5, "slides_per_lecture": 10})
    assert r.status_code == 402
    assert r.json()["error"]["code"] == "limit_exceeded"
    # 3 lectures x 20 slides = 60 slide credits + plan cost exceeds the 60-credit free allowance
    r = await client.post("/api/v1/courses", headers=u["headers"], json={
        "topic": "Photosynthesis", "grade": "8", "subject": "Science", "num_lectures": 3, "slides_per_lecture": 20})
    assert r.status_code == 200  # planning is allowed
    course_id = r.json()["course"]["id"]
    r = await client.post(f"/api/v1/courses/{course_id}/generate", headers=u["headers"], json={})
    assert r.status_code == 402
    usage = (await client.get("/api/v1/me/usage", headers=u["headers"])).json()
    assert usage["plan"]["code"] == "free" and usage["usage"]["credits"]["used"] > 0
