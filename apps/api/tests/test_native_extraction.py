from __future__ import annotations

import copy
import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from app.engine.render.renderer import DeckRenderer
from app.engine.style.content import extract_content
from app.engine.style.pptx_analyzer import analyze_pptx
from app.engine.template.builder import build_native
from app.generation.pipeline import normalize_deck
from app.generation.specs import Bullet, LessonDeck, SlideSpec, Visual
from app.services.sources import extract_pages


def make_stage_deck(path):
    prs = Presentation()
    width, height = prs.slide_width, prs.slide_height
    master = prs.slide_master
    scratch_prs = Presentation()
    scratch = scratch_prs.slides.add_slide(scratch_prs.slide_layouts[6])
    # Different arbitrary layout names; the pointer, not the name, determines the stage.
    stages = ["Engage", "Objective", "Evaluate"]
    for index, stage in enumerate(stages):
        shape = scratch.shapes.add_textbox(0, int(height * (.1 + index * .1)), int(width * .1), int(height * .1))
        shape.text = stage
        master.shapes._spTree.insert_element_before(copy.deepcopy(shape._element), "p:extLst")
    for index, stage in enumerate(stages):
        layout = prs.slide_layouts[index]
        layout.name = f"Imported layout {index}"
        panel = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(width * .12), 0, int(width * .88), height)
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
        layout.shapes._spTree.insert_element_before(copy.deepcopy(panel._element), "p:extLst")
        pointer = scratch.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, int(width * .085),
                                          int(height * (.14 + index * .1)), int(width * .025), int(height * .02))
        layout.shapes._spTree.insert_element_before(copy.deepcopy(pointer._element), "p:extLst")
        slide = prs.slides.add_slide(layout)
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(width * .12), 0,
                                      int(width * .88), int(height * .1))
        band.fill.solid()
        band.fill.fore_color.rgb = RGBColor(230, 240 - index * 20, 220)
        band.text = "Date: old date"
        text = slide.shapes.add_textbox(int(width * .2), int(height * .3), int(width * .6), int(height * .3))
        text.text = f"{stage}: Count 3 + 2 = 5"
    prs.save(path)
    return prs


def test_stage_layouts_keep_headers_and_navigation_without_old_dates(tmp_path):
    path = tmp_path / "stages.pptx"
    original = make_stage_deck(path)
    analysis = analyze_pptx(path)
    base, spec = build_native(path, analysis)
    assert set(spec["stage_variants"]) == {"engage", "objective", "evaluate"}
    slides = [SlideSpec(number=i + 1, layout="concept", purpose="test", title=f"New {stage}",
                        teaching_stage=stage, bullets=[Bullet(text="3 + 2 = 5")])
              for i, stage in enumerate(["engage", "objective", "evaluate"])]
    renderer = DeckRenderer(base, spec)
    rendered = Presentation(io.BytesIO(renderer.render(slides)))
    for index, slide in enumerate(rendered.slides):
        assert slide.slide_layout.name == original.slides[index].slide_layout.name
        assert not any("old date" in sh.text for sh in slide.shapes if sh.has_text_frame)
        assert any(sh.has_text_frame and sh.text.startswith("New ") for sh in slide.shapes)
        assert any(sh.fill.type is not None and sh.fill.fore_color.type is not None
                   for sh in slide.shapes if sh.shape_type == MSO_SHAPE.RECTANGLE)
    assert not any(report.overflow for report in renderer.reports)


def test_extracts_groups_tables_notes_inherited_text_and_picture_bytes(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_textbox(100000, 100000, 3000000, 1000000)
    shape.text = "Put together 4 and 2"
    table = slide.shapes.add_table(2, 2, 400000, 1500000, 3000000, 1000000).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Part", "Whole"
    table.cell(1, 0).text, table.cell(1, 1).text = "4 + 2", "6"
    slide.notes_slide.notes_text_frame.text = "Expected answer: six."
    layout_text = slide.shapes.add_textbox(2000000, 4000000, 3000000, 500000)
    layout_text.text = "Explain with cubes"
    slide.slide_layout.shapes._spTree.insert_element_before(copy.deepcopy(layout_text._element), "p:extLst")
    slide.shapes._spTree.remove(layout_text._element)
    buffer = io.BytesIO()
    Image.new("RGB", (80, 60), "green").save(buffer, "PNG")
    slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), 4000000, 1000000)
    pages, images = extract_content(prs)
    assert "Put together 4 and 2" in pages[0]["text"]
    assert "Part | Whole" in pages[0]["text"] and "4 + 2 | 6" in pages[0]["text"]
    assert "Explain with cubes" in pages[0]["text"]
    assert "Expected answer: six." in pages[0]["notes"]
    assert images[pages[0]["pictures"][0]["image_key"]] == buffer.getvalue()
    output = io.BytesIO()
    prs.save(output)
    assert "Explain with cubes" in extract_pages(output.getvalue(), "pptx")[0][1]


def test_paragraph_font_is_resolved_instead_of_theme_fallback(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_textbox(2000000, 1000000, 5000000, 800000)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.name = "Times New Roman"
    paragraph.font.size = Pt(30)
    paragraph.add_run().text = "Addition"
    path = tmp_path / "font.pptx"
    prs.save(path)
    analysis = analyze_pptx(path)
    assert analysis["fonts"]["heading"]["family"] == "Times New Roman"


def test_uploaded_picture_is_contained_without_cropping():
    from app.engine.template.builder import build_builtin

    base, spec = build_builtin("clean-blue")
    buffer = io.BytesIO()
    Image.new("RGB", (900, 200), "red").save(buffer, "PNG")
    slide = SlideSpec(number=1, layout="image_text", purpose="addition", title="Count all the cubes",
                      visual=Visual(kind="image", source_image_key="original"), asset_id="picture")
    data = DeckRenderer(base, spec).render([slide], {"picture": buffer.getvalue()})
    rendered = Presentation(io.BytesIO(data))
    image = next(sh for sh in rendered.slides[0].shapes if hasattr(sh, "image"))
    assert image.crop_left == image.crop_right == image.crop_top == image.crop_bottom == 0
    assert image.width / image.height == pytest.approx(4.5)


def test_slide_timings_match_requested_duration():
    from app.generation.offline import lesson_deck

    req = {"topic": "Addition", "grade": "1", "subject": "Mathematics", "lecture_minutes": 45,
           "slides_per_lecture": 10, "lecture": {"number": 1, "title": "Addition", "key_concepts": ["Parts"]}}
    deck = lesson_deck(req, LessonDeck)
    for slide in deck.slides:
        slide.timing_minutes = 6
    deck = normalize_deck(deck, req=req, lecture_number=1, total=1, lecture_title="Addition")
    assert sum(slide.timing_minutes for slide in deck.slides) == pytest.approx(45)
    assert all(slide.timing_minutes >= 0 for slide in deck.slides)


def test_counting_diagram_contains_exact_number_of_editable_objects():
    from app.engine.template.builder import build_builtin
    from app.generation.specs import CountingGroup

    base, spec = build_builtin("clean-blue")
    visual = Visual(kind="diagram", show_total=False, counting_groups=[
        CountingGroup(count=3, label="Red cubes", color="red"),
        CountingGroup(count=2, label="Green cubes", color="green")])
    slide = SlideSpec(number=1, layout="image_text", title="Find the whole", purpose="check", visual=visual)
    data = DeckRenderer(base, spec).render([slide])
    rendered = Presentation(io.BytesIO(data))
    objects = [sh for sh in rendered.slides[0].shapes if sh.name.startswith("Counting group")]
    assert len(objects) == 5
    assert sum(sh.name.startswith("Counting group 1 ") for sh in objects) == 3
    assert sum(sh.name.startswith("Counting group 2 ") for sh in objects) == 2
    assert any(sh.has_text_frame and sh.text == "3 + 2 = ?" for sh in rendered.slides[0].shapes)


@pytest.mark.asyncio
async def test_original_assets_are_reused_and_other_teachers_assets_are_rejected(client):
    import uuid

    from app.core.db import get_sessionmaker
    from app.engine.style.content import extract_content
    from app.services.assets import resolve_images
    from app.services.styles import store_source_images
    from tests.conftest import make_user

    teacher = await make_user(client)
    other_teacher = await make_user(client)
    owner = uuid.UUID(teacher["id"])
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    buffer = io.BytesIO()
    Image.new("RGB", (90, 70), "green").save(buffer, "PNG")
    slide.shapes.add_picture(io.BytesIO(buffer.getvalue()), 2000000, 1000000)
    content, _ = extract_content(prs)
    spec = {"source_content": content}
    data = io.BytesIO()
    prs.save(data)
    async with get_sessionmaker()() as db:
        await store_source_images(db, owner, data.getvalue(), spec, "original.pptx")
        await db.commit()
        digest = content[0]["pictures"][0]["image_key"]
        page = SlideSpec(number=1, layout="image_text", title="Original", purpose="test",
                         visual=Visual(kind="image", source_image_key=digest))
        images, counts = await resolve_images(db, owner_id=owner, slides=[page], colors={}, openverse=False,
                                              source_images=spec["source_images"])
        assert counts["reused"] == 1 and counts["placeholder"] == 0
        assert images[page.asset_id] == buffer.getvalue()
        assert spec.get("source_reference_key")
        foreign_page = page.model_copy(deep=True)
        foreign_page.asset_id = None
        _, counts = await resolve_images(db, owner_id=uuid.UUID(other_teacher["id"]),
                                         slides=[foreign_page], colors={}, openverse=False,
                                         source_images=spec["source_images"])
        assert counts["reused"] == 0 and counts["placeholder"] == 1
        assert foreign_page.visual.source_image_key is None
