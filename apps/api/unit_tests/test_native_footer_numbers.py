import copy
import io
from pathlib import Path

import pytest
from pptx import Presentation

from app.engine.render.renderer import DeckRenderer
from app.engine.style.footer_numbers import is_footer_number_text
from app.engine.style.pptx_analyzer import ShapeInfo, analyze_pptx
from app.engine.template.builder import build_native
from app.generation.specs import Bullet, SlideSpec

SAMPLE = Path(__file__).resolve().parents[3] / 'samples' / 'science_ms_sara.pptx'


@pytest.mark.parametrize('legacy_spec', [False, True])
@pytest.mark.parametrize('enabled', [False, True])
def test_selected_native_page_does_not_repeat_old_numbers_and_preserves_teaching_digits(enabled, legacy_spec):
    analysis = analyze_pptx(SAMPLE)
    base, spec = build_native(SAMPLE, analysis)
    selected = next(page for page in spec['page_variants'] if page['number'] == 3)
    assert selected['slide_number_ids']
    assert set(selected['slide_number_ids']) <= set(selected['item_ids']), 'keep donor shape order for saved edits'
    if legacy_spec:
        spec = copy.deepcopy(spec)
        for page in spec['page_variants']:
            page.pop('slide_number_ids', None)
    slides = [SlideSpec(number=number, layout='concept', design_page_number=3, title=f'Energy transfer {number}',
                        purpose='Explain', bullets=[Bullet(text='3')]) for number in (1, 2)]
    deck = Presentation(io.BytesIO(DeckRenderer(base, spec).render(slides, core_props={
        'teacher_footer_text': 'Ms New Teacher', 'page_numbers': 'true' if enabled else 'false',
    })))
    for number, slide in enumerate(deck.slides, start=1):
        named = {shape.name: shape for shape in slide.shapes}
        assert 'Template page number' not in named
        assert ('Teacher page number' in named) is enabled
        if enabled:
            assert named['Teacher page number'].text == f'{number} / 2'
        digits = [shape for shape in slide.shapes if shape.has_text_frame and shape.text.strip() == '3']
        assert len(digits) == 1, 'the teaching numeral stays; the original page 3 is removed'
        assert digits[0].top < deck.slide_height * .85


def test_footer_number_recognition_keeps_equations_diagram_values_and_large_numerals():
    footer = (.85, .91, .07, .05)
    assert is_footer_number_text('3', footer, 11)
    assert not is_footer_number_text('Q = 3', footer, 11)
    assert not is_footer_number_text('3.14', footer, 11)
    assert not is_footer_number_text('3', (.4, .91, .07, .05), 11)
    assert not is_footer_number_text('3', (.85, .5, .07, .05), 11)
    assert not is_footer_number_text('3', footer, 32)
    content = ShapeInfo(shape_id=9, kind='text', bbox=(.4, .5, .1, .1), text='3', max_font_pt=24)
    assert content.signature()[0] == 'text'


def test_teacher_number_toggle_does_not_change_teaching_shape_ids():
    base, spec = build_native(SAMPLE, analyze_pptx(SAMPLE))
    slides = [SlideSpec(number=1, layout='concept', design_page_number=3, title='Example', purpose='Explain',
                        bullets=[Bullet(text='3')])]
    ids = []
    for enabled in (False, True):
        deck = Presentation(io.BytesIO(DeckRenderer(base, spec).render(slides,
            core_props={'page_numbers': 'true' if enabled else 'false'})))
        ids.append({shape.text: shape.shape_id for shape in deck.slides[0].shapes
                    if shape.has_text_frame and shape.text in {'Example', '3'}})
    assert ids[0] == ids[1]
