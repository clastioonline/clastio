import io

import pytest
from pptx import Presentation

from app.engine.render.renderer import DeckRenderer
from app.engine.template.builder import BUILTIN_STYLES, build_builtin
from app.generation.math_visuals import algebra_visual
from app.generation.specs import Bullet, SlideSpec, Visual


def slide(text):
    return SlideSpec(number=1, layout='image_text', purpose='Explain', title='Substitution',
                     bullets=[Bullet(text=text)], visual=Visual(kind='diagram'))


def test_expression_and_substitution_values_come_from_slide():
    result = algebra_visual(slide('m = 4x + 10, with x = 5.'))
    assert result['expression'] == '4x + 10'
    assert result['answer'] == 30


@pytest.mark.parametrize('text', ['-3x + 5', '0.3x + 5', '1/3x + 5', '3x + 5y', '3x + 5^2', 'Photosynthesis'])
def test_unsupported_expressions_are_not_misrepresented(text):
    assert algebra_visual(slide(text)) is None


def test_expression_visual_is_native_editable_and_has_correct_answer():
    base, template = build_builtin(next(iter(BUILTIN_STYLES)))
    renderer = DeckRenderer(base, template)
    data = renderer.render([slide('m = 4x + 10, with x = 5.')])
    shapes = Presentation(io.BytesIO(data)).slides[0].shapes
    text = '\n'.join(shape.text for shape in shapes if shape.has_text_frame)
    for label in ['4x + 10', 'Coefficient: 4', 'Variable: x', 'Constant: 10', 'x = 5 → 30']:
        assert label in text
    assert not any(shape.shape_type.name=='PICTURE' for shape in shapes)
    assert not any(report.overflow for report in renderer.reports)


async def test_unavailable_sources_do_not_publish_placeholder(monkeypatch):
    import uuid
    from types import SimpleNamespace
    from unittest.mock import AsyncMock, Mock

    from app.services import assets

    item=slide('A thermometer measures temperature.')
    item.visual.kind='image'
    db=AsyncMock()
    db.execute.return_value=SimpleNamespace(scalars=lambda:SimpleNamespace(first=lambda:None))
    storage=SimpleNamespace(put=AsyncMock())
    monkeypatch.setattr(assets,'get_storage',lambda:storage)
    monkeypatch.setattr(assets,'get_ai',lambda:SimpleNamespace(image=AsyncMock()))
    from app.services import settings
    monkeypatch.setattr(settings,'get_setting',AsyncMock(return_value={}))
    placeholder=Mock(side_effect=AssertionError('Placeholder must never be published'))
    monkeypatch.setattr(assets,'placeholder_illustration',placeholder)
    images,counters=await assets.resolve_images(db,owner_id=uuid.uuid4(),slides=[item],colors={},openverse=False)
    assert images=={}
    assert item.visual.kind=='none'
    assert item.layout=='concept'
    assert counters['placeholder']==0
    storage.put.assert_not_awaited()
    placeholder.assert_not_called()


async def test_stock_search_rejects_unrelated_metadata(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from app.services import assets

    response=SimpleNamespace(raise_for_status=lambda:None,json=lambda:{'results':[
        {'title':'Mountain landscape','license':'cc0','width':1000,'url':'https://example.com/mountain.jpg','tags':[]}]})
    download=AsyncMock()
    monkeypatch.setattr(assets,'download_public_image',download)
    found=await assets.search_openverse('laboratory thermometer',SimpleNamespace(get=AsyncMock(return_value=response)))
    assert found is None
    download.assert_not_awaited()
