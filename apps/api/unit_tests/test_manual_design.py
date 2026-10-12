import io

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.engine.render.manual import apply_manual, editable_objects
from app.engine.render.watermark import apply_watermark
from app.generation.specs import ManualObjectEdit


@pytest.mark.parametrize('position', ['footer', 'diagonal'])
def test_teacher_watermark_is_editable_and_applied_once_per_slide(position):
    deck = Presentation()
    for _ in range(2):
        deck.slides.add_slide(deck.slide_layouts[6])
    props = {'watermark_text': 'Ms Teacher · Science', 'watermark_position': position}
    apply_watermark(deck, props)
    apply_watermark(deck, props)
    output = io.BytesIO()
    deck.save(output)
    restored = Presentation(io.BytesIO(output.getvalue()))
    for slide in restored.slides:
        marks = [shape for shape in slide.shapes if shape.name == 'Teacher watermark']
        assert len(marks) == 1
        assert marks[0].text == props['watermark_text']
        assert marks[0].rotation == (330 if position == 'diagonal' else 0)


def test_watermark_is_opt_in():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    apply_watermark(deck, {})
    assert len(slide.shapes) == 0


def test_teacher_footer_and_numbering_roundtrip_without_duplicate_numbers():
    deck = Presentation()
    for _ in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        number = slide.shapes.add_textbox(0, 0, Inches(1), Inches(.2))
        number.name = 'Template page number'
        number.text = '99'
    props = {'teacher_footer_text': 'Ms Sara', 'page_numbers': 'true', 'watermark_text': 'School'}
    apply_watermark(deck, props)
    apply_watermark(deck, props)
    output = io.BytesIO()
    deck.save(output)
    restored = Presentation(io.BytesIO(output.getvalue()))
    for index, slide in enumerate(restored.slides, start=1):
        named = {shape.name: shape for shape in slide.shapes}
        assert len(slide.shapes) == 3
        assert 'Template page number' not in named
        assert named['Teacher name footer'].text == 'Ms Sara'
        assert named['Teacher page number'].text == f'{index} / 3'
        assert named['Teacher name footer'].left > named['Teacher watermark'].left
    apply_watermark(restored, {'page_numbers': 'false'})
    assert all(not list(slide.shapes) for slide in restored.slides)


def test_manual_text_and_geometry_survive_powerpoint_roundtrip():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    run = box.text_frame.paragraphs[0].add_run()
    run.text, run.font.name, run.font.size = "Original", "Carlito", Pt(24)
    key = str(box.shape_id)
    apply_manual(slide, {key: ManualObjectEdit(text="Worked example: 2 + 3 = 5", color="#1255AA",
        font_size=32, x=.2, y=.3, width=.5, height=.2, bold=True)}, deck.slide_width, deck.slide_height)
    buffer = io.BytesIO()
    deck.save(buffer)
    objects = editable_objects(buffer.getvalue())[0]
    assert objects[0]["text"] == "Worked example: 2 + 3 = 5"
    assert objects[0]["color"] == "#1255AA"
    assert objects[0]["font_size"] == 32
    assert objects[0]["font_family"] == "Carlito"
    assert objects[0]["x"] == pytest.approx(.2)
    assert objects[0]["width"] == pytest.approx(.5)
    restored = Presentation(io.BytesIO(buffer.getvalue()))
    assert restored.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold is True


def test_unspecified_formatting_is_preserved_and_missing_objects_ignored():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text, run.font.name, run.font.size = "Keep this", "Carlito", Pt(28)
    apply_manual(slide, {str(shape.shape_id): ManualObjectEdit(x=.1),
                         "9999": ManualObjectEdit(text="Missing")}, deck.slide_width, deck.slide_height)
    assert run.text == "Keep this"
    assert run.font.size.pt == 28
    assert run.font.name == "Carlito"


@pytest.mark.parametrize("edit", [{"color": "red"}, {"font_size": 0}, {"width": 0}, {"x": -1}])
def test_invalid_manual_formatting_is_rejected(edit):
    with pytest.raises(ValueError):
        ManualObjectEdit(**edit)


@pytest.mark.parametrize("model_name", ["LessonDeck", "SlideRewrite", "SlideSpec"])
def test_manual_objects_provider_schema_is_closed_and_keeps_runtime_limits(model_name):
    import json

    from app.ai.schema_utils import strict_schema
    from app.generation import specs

    model = getattr(specs, model_name)
    schema = strict_schema(model)
    encoded = json.dumps(schema)
    assert "maxProperties" not in encoded
    assert "minProperties" not in encoded
    slide = schema["properties"]["slides"]["items"] if model_name == "LessonDeck" else (
        schema["properties"]["slide"] if model_name == "SlideRewrite" else schema)
    assert slide["properties"]["manual_objects"] == {
        "type": "object", "properties": {}, "additionalProperties": False, "required": []}
    manual = {str(n): {"color": "#1255AA"} for n in range(151)}
    with pytest.raises(ValueError):
        specs.SlideSpec(number=1, layout="concept", title="Example", purpose="teach", manual_objects=manual)
