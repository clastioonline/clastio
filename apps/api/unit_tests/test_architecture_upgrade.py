"""Failure boundaries, visible word budgets, and provider-specific wire contracts."""
import io
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import httpx
import pytest
from pptx import Presentation
from pydantic import BaseModel, ValidationError

from app.ai.base import AIRequest
from app.ai.compatible_provider import CompatibleProvider
from app.engine.fonts import required_fonts
from app.engine.qc import visual
from app.generation import offline
from app.generation.specs import Bullet, CoursePlan, LessonDeck
from app.services.voice import trusted_media_url


def test_voice_media_boundary():
    assert trusted_media_url("https://lookaside.fbsbx.com/whatsapp_business/attachments/x")
    for url in ["http://lookaside.fbsbx.com/x", "https://fbsbx.com.evil.test/x", "https://127.0.0.1/x",
                "https://user@lookaside.fbsbx.com/x", "https://lookaside.fbsbx.com:8080/x"]:
        assert not trusted_media_url(url)


def test_fonts_read_actual_donor_xml():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Hello"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = "School Proprietary Font"
    buf = io.BytesIO()
    prs.save(buf)
    assert "School Proprietary Font" in required_fonts(buf.getvalue())


def test_sidecar_failure_never_launches_local_office(monkeypatch):
    monkeypatch.setattr(visual, "get_settings", lambda: SimpleNamespace(gotenberg_url="http://renderer:3000", render_timeout_s=2))
    def failure(*args, **kwargs):
        raise httpx.ConnectError("offline")
    monkeypatch.setattr(httpx.Client, "post", failure)
    monkeypatch.setattr(visual.subprocess, "run", lambda *a, **k: pytest.fail("Local Office fallback defeats memory isolation"))
    with pytest.raises(visual.RenderError):
        visual.pptx_to_pdf(b"deck")
    with pytest.raises(visual.RenderError):
        visual.docx_to_pdf(b"doc")


def test_visible_slide_limit_excludes_speaker_notes():
    req = {"topic": "Addition", "grade": "2", "subject": "Maths", "curriculum": "british",
           "num_lectures": 1, "slides_per_lecture": 6, "lecture_minutes": 40, "outcomes": []}
    course = offline.course_plan(req, CoursePlan)
    deck = offline.lesson_deck({**req, "course": course.model_dump(), "lecture": course.lectures[0].model_dump()}, LessonDeck)
    deck.slides[0].speaker_notes = "explanation " * 300
    LessonDeck.model_validate(deck.model_dump())
    deck.slides[0].bullets = [Bullet(text="word " * 121)]
    with pytest.raises(ValidationError, match="120-word"):
        LessonDeck.model_validate(deck.model_dump())


async def test_groq_json_mode_validates_schema():
    class Intent(BaseModel):
        subject: str
    provider = CompatibleProvider("groq", None, "https://api.groq.com/openai/v1")
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"subject":"Science"}', refusal=None), finish_reason="stop")], usage=None)
    provider._call = AsyncMock(return_value=response)
    result = await provider.generate_structured("llama-3.3-70b-versatile", AIRequest(task="intent", system="Route", messages=[]), Intent)
    assert result.data.subject == "Science"
    args = provider._call.call_args.kwargs
    assert args["response_format"] == {"type": "json_object"}
    assert "max_tokens" in args and "max_completion_tokens" not in args
    assert json.loads(args["messages"][0]["content"].split("schema: ")[1])["properties"]["subject"]


async def test_scanned_pdf_keeps_page_provenance(monkeypatch):
    import pymupdf

    from app.services import pdf_ingestion
    document = pymupdf.open()
    first = document.new_page()
    first.insert_text((50, 50), "Native source text")
    document.new_page()
    data = document.tobytes()
    document.close()
    ai = SimpleNamespace(mode="live", structured=AsyncMock(return_value=pdf_ingestion.PageText(text="Scanned page text")))
    monkeypatch.setattr(pdf_ingestion, "get_ai", lambda: ai)
    result = await pdf_ingestion.extract_scanned(data, [(1, "Native source text")], __import__("uuid").uuid4())
    assert result == [(1, "Native source text"), (2, "Scanned page text")]
    assert ai.structured.call_args.kwargs["tier"] == "ingestion"
    assert ai.structured.call_args.kwargs["prompt"] == "Transcribe source page 2."
