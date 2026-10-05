import io

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.engine.render.renderer import DeckRenderer
from app.engine.style.design_details import inspect_design
from app.engine.style.pptx_analyzer import analyze_pptx
from app.engine.template.builder import build_native
from app.generation.specs import ChartData, ChartSeries, SlideSpec, TableData


def uploaded_deck(path):
    prs = Presentation()
    for number in range(16):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(7), Inches(.8))
        title.text = f"Original lesson {number + 1}"
    slide = prs.slides[-1]
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(5), Inches(2)).table
    for row_number, row in enumerate(table.rows):
        for cell in row.cells:
            cell.text = "Original private data"
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("C0FFEE" if row_number == 0 else "FEEDAA")
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor.from_string("112233")
    data = CategoryChartData()
    data.categories = ["Original category"]
    data.add_series("Original private series", [91])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4),
                                   Inches(5), Inches(1), data).chart
    chart.chart_style = 12
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string("123ABC")
    prs.save(path)
    return prs


def test_inspection_covers_final_slide_and_excludes_source_data(tmp_path):
    result = inspect_design(uploaded_deck(tmp_path / "source.pptx"))
    assert result["slides_inspected"] == 16
    assert result["slides"][-1]["tables"] == 1
    assert result["slides"][-1]["charts"] == 1
    assert "Original private" not in str(result)
    assert result["charts"]["bar"]["style"] == 12


def test_native_generation_reuses_table_and_chart_formatting(tmp_path):
    path = tmp_path / "source.pptx"
    uploaded_deck(path)
    analysis = analyze_pptx(path)
    base, spec = build_native(path, analysis)
    assert spec["design_inspection"]["slides_inspected"] == 16
    assert spec["donor_count"] == 16
    assert len(Presentation(io.BytesIO(base)).slides) == 16
    assert spec["page_variants"][-1]["number"] == 16
    slides = [SlideSpec(number=1, layout="table", title="New table", purpose="teach",
                        table=TableData(headers=["New heading", "Value"], rows=[["New fact", "7"]])),
              SlideSpec(number=2, layout="chart", title="New chart", purpose="teach",
                        chart=ChartData(categories=["New category"], series=[ChartSeries(name="New data", values=[7])],
                                        source="Teacher supplied"))]
    output = Presentation(io.BytesIO(DeckRenderer(base, spec).render(slides)))
    table = next(shape.table for shape in output.slides[0].shapes if shape.has_table)
    assert table.cell(0, 0).text == "New heading"
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "C0FFEE"
    assert str(table.cell(1, 0).fill.fore_color.rgb) == "FEEDAA"
    assert str(table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == "112233"
    chart = next(shape.chart for shape in output.slides[1].shapes if shape.has_chart)
    assert chart.chart_style == 12
    assert str(chart.series[0].format.fill.fore_color.rgb) == "123ABC"
    assert tuple(chart.series[0].values) == (7.0,)


def test_plain_deck_has_no_invented_table_or_chart_style():
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    result = inspect_design(prs)
    assert result["table"] is None
    assert result["charts"] == {}
