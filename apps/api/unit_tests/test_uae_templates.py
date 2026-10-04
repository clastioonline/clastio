from __future__ import annotations

import io
import re
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock
from zipfile import ZipFile

import pytest
from pptx import Presentation

from app.engine.render.renderer import DeckRenderer
from app.engine.style.common import contrast_ratio
from app.engine.template.builder import BUILTIN_STYLES, build_builtin, builtin_analysis
from app.engine.template.catalog import UAE_STYLES
from app.engine.template.preview import PREVIEW_SLIDES
from app.generation.specs import Bullet, SlideSpec
from app.services import styles
from app.services.template_recommendations import grade_number, suggest_templates


def row(key, *, owner=None, org=None, shared=False):
    return SimpleNamespace(id=uuid.uuid4(), owner_id=owner, org_id=org, is_shared=shared,
                           status="ready", name=BUILTIN_STYLES[key]["name"],
                           spec={"catalog": {**builtin_catalog(key)}})


def builtin_catalog(key):
    from app.engine.template.catalog import catalog_metadata

    return catalog_metadata(key, BUILTIN_STYLES[key])


@pytest.mark.parametrize("key", list(BUILTIN_STYLES))
def test_builtin_renders_editable_native_deck_with_legible_type(key):
    base, spec = build_builtin(key)
    output = DeckRenderer(base, spec).render(PREVIEW_SLIDES, {}, core_props={"title": "UAE template check"})
    deck = Presentation(io.BytesIO(output))
    assert len(deck.slides) == 6
    assert spec["catalog"]["key"] == key
    assert spec["zones"]["body"][2] >= .8 and spec["zones"]["body"][3] >= .5
    assert contrast_ratio(spec["colors"]["text"], spec["colors"]["background"]) >= 4.5
    title_background = spec["colors"]["primary"] if BUILTIN_STYLES[key]["band"] else spec["colors"]["background"]
    assert contrast_ratio(spec["colors"]["title"], title_background) >= 4.5
    assert all(slide.shapes for slide in deck.slides)
    assert any("Water" in shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    assert not any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes), "Chrome is editable, not a raster screenshot"


def test_new_designs_have_distinct_geometry_and_typography():
    signatures = []
    for key in UAE_STYLES:
        analysis = builtin_analysis(key)
        geometry = str(analysis["decorations"])
        signatures.append((geometry, analysis["fonts"]["heading"]["family"]))
    assert len(set(signatures)) == 12
    assert len({style["layout"] for style in UAE_STYLES.values()}) >= 10


def test_arabic_template_uses_real_rtl_and_arabic_font_in_export():
    base, spec = build_builtin("uae-arabic-classroom")
    slides = [SlideSpec(number=1, layout="concept", purpose="Explain", title="دورة الماء",
                        bullets=[Bullet(text="تسخن الشمس الماء"), Bullet(text="يتحول الماء إلى بخار")])]
    output = DeckRenderer(base, spec, language="ar").render(slides, {})
    with ZipFile(io.BytesIO(output)) as archive:
        slide_part = next(name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
        xml = archive.read(slide_part).decode()
    assert 'rtl="1"' in xml and "Noto Sans Arabic" in xml
    assert "دورة الماء" in xml


@pytest.mark.parametrize(("subject", "curriculum", "grade", "language", "expected"), [
    ("Mathematics", "british", "8", "en", "uae-maths-grid"),
    ("Science", "cbse", "8", "en", "uae-cbse-concepts"),
    ("Arabic", "moe", "6", "ar", "uae-arabic-classroom"),
    ("English", "british", "7", "en", "uae-language-stories"),
    ("Science", "british", "KG", "en", "uae-early-years"),
    ("Social Studies & Moral Education", "moe", "5", "en", "uae-heritage"),
    ("Artificial Intelligence", "uae_ai", "9", "en", "uae-future-makers"),
    ("Art", "ib", "4", "en", "uae-ib-inquiry"),
])
def test_suggestions_match_explicit_context(subject, curriculum, grade, language, expected):
    candidates = [row(key) for key in BUILTIN_STYLES]
    suggestions = suggest_templates(candidates, owner_id=uuid.uuid4(), subjects=[subject], curriculum=curriculum,
                                    grades=[grade], language=language, country="AE")
    assert suggestions[0]["template_id"] == str(next(t.id for t in candidates if t.name == BUILTIN_STYLES[expected]["name"]))
    assert suggestions[0]["reason"] and len(suggestions) <= 3


def test_saved_default_and_own_uploaded_style_are_preserved_and_private_uploads_excluded():
    owner, other, org = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    saved, owned, leaked = row("warm-sand"), row("uae-maths-grid", owner=owner), row("uae-maths-grid", owner=other, shared=True)
    shared = row("uae-maths-grid", owner=other, org=org, shared=True)
    suggestions = suggest_templates([saved, owned, leaked, shared], owner_id=owner, default_id=saved.id,
                                    subjects=["Science"], grades=["8"], curriculum="british")
    ids = [item["template_id"] for item in suggestions]
    assert ids[0] == str(saved.id) and str(owned.id) in ids
    assert str(leaked.id) not in ids and str(shared.id) not in ids
    organisation = suggest_templates([shared], owner_id=owner, org_id=org)
    assert organisation[0]["template_id"] == str(shared.id)


@pytest.mark.parametrize(("value", "expected"), [("KG", 0), ("FS2", 0), ("Year 8", 8), ("Grade 4", 4),
                                                 ("12", 12), ("unrecognised", None), ("99", None)])
def test_grade_context_is_validated(value, expected):
    assert grade_number(value) == expected


async def test_builtin_seeding_updates_old_design_in_place_and_is_idempotent(monkeypatch):
    existing = row("warm-sand")
    existing.spec = {}  # Original catalogue has no revision marker.
    rows = {existing.name: existing}
    writes = AsyncMock()
    previews = AsyncMock(return_value=["preview.webp"])
    monkeypatch.setattr(styles, "get_storage", lambda: SimpleNamespace(put=writes))
    monkeypatch.setattr(styles, "render_previews", previews)
    db = AsyncMock()

    async def execute(query):
        name = next(value for value in query.compile().params.values() if value in {st["name"] for st in BUILTIN_STYLES.values()})
        return SimpleNamespace(scalars=lambda: SimpleNamespace(first=lambda: rows.get(name)))

    db.execute.side_effect = execute
    db.add = lambda template: rows.__setitem__(template.name, template)
    await styles.ensure_builtin_templates(db)
    assert rows[existing.name].id == existing.id
    assert len(rows) == len(BUILTIN_STYLES) == 15
    assert writes.await_count == 15 and previews.await_count == 15
    storage_key = existing.base_storage_key
    await styles.ensure_builtin_templates(db)
    assert writes.await_count == 15 and previews.await_count == 15
    assert existing.base_storage_key == storage_key
