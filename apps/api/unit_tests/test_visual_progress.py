import io
from types import SimpleNamespace

import pymupdf as fitz
import pytest
from pptx import Presentation

from app.engine.qc import visual
from app.services import courses


def test_visual_inspection_reports_every_slide(monkeypatch):
    presentation = Presentation()
    pdf = fitz.open()
    for number in range(3):
        presentation.slides.add_slide(presentation.slide_layouts[6])
        page = pdf.new_page(width=720, height=540)
        page.insert_text((40, 40), f"Slide {number + 1}")
    buffer = io.BytesIO()
    presentation.save(buffer)
    monkeypatch.setattr(visual, "pptx_to_pdf", lambda _: pdf.tobytes())
    stages = []
    report = visual.inspect(buffer.getvalue(), on_progress=stages.append)
    assert len(report.previews) == len(report.thumbnails) == 3
    assert stages[0] == "Converting PowerPoint to PDF"
    assert [stage for stage in stages if stage.startswith("Checking slide")] == [
        "Checking slide 1 of 3", "Checking slide 2 of 3", "Checking slide 3 of 3"]
    assert stages[-1] == "Visual quality check complete"
    pdf.close()


def test_visual_deadline_includes_conversion(monkeypatch):
    monkeypatch.setattr(visual, "get_settings", lambda: SimpleNamespace(render_timeout_s=10))
    clock = iter([0, 0, 0, 11])
    monkeypatch.setattr(visual.time, "monotonic", lambda: next(clock))
    monkeypatch.setattr(visual, "pptx_to_pdf", lambda _: b"pdf")
    with pytest.raises(visual.RenderError, match="time limit"):
        visual.inspect(b"pptx")


async def test_visual_progress_propagates_render_failure(monkeypatch):
    def fail(*args, **kwargs):
        raise visual.RenderError("LibreOffice timed out")
    monkeypatch.setattr(courses, "inspect", fail)
    with pytest.raises(visual.RenderError, match="timed out"):
        await courses.inspect_with_progress(SimpleNamespace(job_id="test-job"), b"pptx")
