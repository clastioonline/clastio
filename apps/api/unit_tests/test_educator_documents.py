"""Printable exports preserve pedagogical evidence, answers, provenance and RTL."""

import io
import subprocess
from unittest.mock import Mock

import pytest
from docx import Document
from docx.oxml.ns import qn
from pydantic import ValidationError
from test_educator_schema import demo_dict, plan_dict, resource_dict

from app.engine.exports import educator_documents as exports
from app.engine.qc.visual import RenderError
from app.generation.educator_specs import CopilotTask, UAEPlanContent


def read(data):
    doc = Document(io.BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    tables = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    return doc, text + "\n" + tables


def source():
    return {"title": "Verified teaching source", "url_or_file_id": "https://publisher.example/chapter",
            "page_or_section": "Page 9", "relevance": "Food-chain explanation"}


def test_plan_export_keeps_backward_design_inclusion_expected_answers_and_sources():
    value = plan_dict()
    value["source_references"] = [source()]
    value["review_checks"]["framework_references"] = [source()]
    value["uae_context"]["arabic_culture"]["relevant"] = False
    value["uae_context"]["arabic_culture"]["connection_or_reason"] = "No direct cultural connection for this calculation."
    doc, text = read(exports.build_plan_docx(UAEPlanContent.model_validate(value), title="Teacher planning draft"))
    assert doc.paragraphs[0].text == "Teacher planning draft"
    for expected in [value["learning_objectives"][0]["success_measure"],
                     value["success_criteria"][0]["observable_evidence"],
                     value["differentiation"]["sen"][0]["assessment_accommodation"],
                     value["differentiation"]["ell"][0]["scaffold"],
                     value["formative_assessment"][0]["expected_response"],
                     value["formative_assessment"][0]["teaching_response"],
                     value["lesson_sequence"][0]["teacher_actions"],
                     value["lesson_sequence"][0]["student_actions"],
                     value["uae_context"]["arabic_culture"]["connection_or_reason"],
                     value["review_checks"]["evidence_gaps"][0], "Objectives: o1", "Assessments: a1",
                     "https://publisher.example/chapter", "Page 9", "Framework references"]:
        assert expected in text
    assert exports.DRAFT_LABEL in doc.sections[0].footer.paragraphs[0].text
    assert "Relevant to this lesson: No" in text
    assert "KHDA-approved" not in text


def test_plan_builder_rejects_invalid_or_unaligned_draft_instead_of_exporting_it():
    value = plan_dict()
    value["formative_assessment"][0]["objective_ids"] = ["missing-objective"]
    with pytest.raises(ValidationError):
        exports.build_plan_docx(value)


def test_demo_export_contains_full_script_timeline_and_the_matching_slide_answers():
    value = demo_dict()
    value["slide_payload"]["slides"][0]["quiz"] = {
        "question": "Which way does the energy arrow point?", "options": ["To the consumer", "To the sun"],
        "answer_index": 0, "explanation": "The consumer obtains energy from the organism it eats.",
    }
    value["slide_payload"]["slides"][0]["speaker_notes"] = "Pause and check each group's arrow direction."
    _, text = read(exports.build_demo_docx(value))
    for expected in ["00:00–05:00", "05:00–10:00", "10:00–15:00", value["hook"]["teacher_says"],
                     value["teacher_script"][0]["teacher_says"], value["teacher_script"][0]["teacher_does"],
                     value["teacher_script"][0]["expected_student_response"],
                     value["teacher_script"][0]["formative_check"], value["teacher_script"][0]["transition"],
                     "Expected answer: To the consumer", "Pause and check each group's arrow direction."]:
        assert expected in text
    for slide in value["slide_payload"]["slides"]:
        assert f"Slide {slide['number']}: {slide['title']}" in text


def test_resource_export_keeps_all_chapters_and_separates_teacher_answers():
    value = resource_dict()
    second = {**value["chapters"][0], "number": 2, "title": "Local habitat study",
              "explanation": "Observe a mangrove habitat without disturbing organisms.",
              "practice_questions": ["Name one fieldwork precaution."],
              "expected_answers": ["Observe from the walkway without collecting organisms."]}
    value["chapters"].append(second)
    value["chapters"][0]["source_references"] = [source()]
    _, text = read(exports.build_resource_docx(value))
    student, key = text.split("Teacher answer key", 1)
    for chapter in value["chapters"]:
        for field in ["title", "explanation", "uae_case_study", "summary"]:
            assert chapter[field] in student
        for field in ["learning_goals", "definitions", "worked_examples", "practice_questions"]:
            assert all(item in student for item in chapter[field])
        assert all(item in key for item in chapter["expected_answers"])
    assert second["expected_answers"][0] not in student
    assert "not an official or publisher textbook" in text
    assert "https://publisher.example/chapter" in text


def task_dict():
    return {"title": "Mangrove reading task", "topic": "Ecosystems", "subject": "Science", "grade": "6",
            "education_stage": "secondary", "curriculum": "british", "language": "en",
            "instructions": "Read the text and explain the relationships.",
            "reading_text": "Mangroves provide shelter for juvenile fish along UAE coasts.",
            "vocabulary": [{"term": "Habitat", "meaning": "The place where an organism lives."}],
            "support": {"sen": "Use short chunks with labelled images.", "ell": "Use the frame: The fish needs __.",
                        "core": "Explain a habitat relationship.", "extension": "Predict a loss of shelter."},
            "questions": [{"question": "Why are roots useful for juvenile fish?",
                           "expected_answer": "Roots provide shelter from predators."}],
            "assessment": "Check that explanations connect shelter and survival.",
            "uae_link": "Compare the example with UAE mangrove conservation.",
            "source_references": [source()], "review_checks": plan_dict()["review_checks"]}


def test_copilot_task_export_has_printable_reading_support_and_a_separate_answer_key():
    value = task_dict()
    _, text = read(exports.build_task_docx(value))
    student, key = text.split("Teacher answer key and assessment", 1)
    for field in ["instructions", "reading_text", "uae_link"]:
        assert value[field] in student
    assert all(item in student for item in value["support"].values())
    assert value["vocabulary"][0]["meaning"] in student
    assert value["questions"][0]["question"] in student
    assert value["questions"][0]["expected_answer"] not in student
    assert value["questions"][0]["expected_answer"] in key
    assert value["assessment"] in key
    assert value["review_checks"]["evidence_gaps"][0] in key


def test_copilot_schema_exports_structured_sen_and_ell_strategies():
    value = task_dict()
    value["support"] = plan_dict()["differentiation"]
    value["questions"].append({"question": "What is a habitat?",
                               "expected_answer": "A place where an organism lives."})
    _, text = read(exports.build_task_docx(CopilotTask.model_validate(value)))
    for key in ["sen", "ell"]:
        for field in ["learning_barrier", "scaffold", "assessment_accommodation"]:
            assert value["support"][key][0][field] in text
    assert "Learning Barrier:" in text
    assert "Assessment Accommodation:" in text


def test_arabic_plan_preserves_text_and_sets_direction_on_paragraphs_tables_and_footer():
    value = plan_dict()
    value["language"] = "ar"
    value["title"] = "النظم البيئية لغابات القرم"
    value["lesson_sequence"][0]["teacher_actions"] = "اشرح انتقال الطاقة باستخدام رسم واضح."
    doc, text = read(exports.build_plan_docx(value))
    assert value["title"] in text
    assert value["lesson_sequence"][0]["teacher_actions"] in text
    paragraphs = [*doc.paragraphs, *doc.tables[0].cell(1, 1).paragraphs,
                  *doc.sections[0].footer.paragraphs]
    for paragraph in paragraphs:
        assert paragraph._p.pPr.find(qn("w:bidi")) is not None
        for run in paragraph.runs:
            assert run._r.rPr.find(qn("w:rtl")) is not None


async def test_pdf_export_uses_existing_renderer_without_paid_calls(monkeypatch):
    renderer = Mock(return_value=b"%PDF-1.7\nvalidated rendering fixture")
    monkeypatch.setattr(exports, "docx_to_pdf", renderer)
    docx = exports.build_task_docx(task_dict())
    assert (await exports.educator_docx_to_pdf(docx)).startswith(b"%PDF-")
    renderer.assert_called_once_with(docx)


@pytest.mark.parametrize("error", [FileNotFoundError("renderer not installed"),
                                   subprocess.TimeoutExpired("soffice", 10),
                                   RenderError("Sidecar conversion failed")])
async def test_pdf_export_reports_missing_renderer_and_conversion_failures(monkeypatch, error):
    monkeypatch.setattr(exports, "docx_to_pdf", Mock(side_effect=error))
    with pytest.raises(RenderError):
        await exports.educator_docx_to_pdf(b"docx fixture")


async def test_pdf_export_rejects_a_non_pdf_response(monkeypatch):
    monkeypatch.setattr(exports, "docx_to_pdf", Mock(return_value=b"server error"))
    with pytest.raises(RenderError, match="invalid PDF"):
        await exports.educator_docx_to_pdf(b"docx fixture")
