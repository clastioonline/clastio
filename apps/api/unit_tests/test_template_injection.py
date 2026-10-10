import io
from copy import deepcopy

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from app.engine.render.renderer import DeckRenderer
from app.engine.render.template_injection import manifest, replace_text
from app.engine.template.builder import BUILTIN_STYLES, build_builtin
from app.generation.specs import Bullet, SlideSpec


def xml_signature(element):
    return (element.tag, sorted(element.attrib.items()), element.text, tuple(xml_signature(child) for child in element))


def template():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = 'Original title'
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.name = 'Arial'
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    body = slide.placeholders[1]
    body.text = 'Original explanation'
    run = body.text_frame.paragraphs[0].runs[0]
    run.font.name, run.font.size = 'Arial', Pt(20)
    run.font.color.rgb = RGBColor.from_string('123456')
    decoration = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(.3), Inches(5))
    decoration.name = 'Keep decoration'
    output = io.BytesIO()
    prs.save(output)
    _, spec = build_builtin(next(iter(BUILTIN_STYLES)))
    spec.update(preserve_slots=True,template_slots=manifest(output.getvalue()))
    return output.getvalue(),spec


def test_injection_preserves_theme_geometry_decorations_and_run_style():
    data,spec = template()
    source = Presentation(io.BytesIO(data))
    original_decoration = xml_signature(source.slides[0].shapes[-1]._element)
    original_rpr = xml_signature(source.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]._r.rPr)
    slides = [SlideSpec(number=i,layout='cover' if i==1 else 'concept',purpose='Explain',title=f'Lesson {i}',
                        bullets=[Bullet(text='A variable represents an unknown value.')]) for i in range(1,5)]
    renderer = DeckRenderer(data,spec)
    output = renderer.render(slides)
    result = Presentation(io.BytesIO(output))
    assert len(result.slides)==4
    for page in result.slides:
        assert xml_signature(page.shapes[-1]._element)==original_decoration
        body = page.placeholders[1]
        assert body.text=='A variable represents an unknown value.'
        assert xml_signature(body.text_frame.paragraphs[0].runs[0]._r.rPr)==original_rpr
        assert body.left==source.slides[0].placeholders[1].left
        assert body.width==source.slides[0].placeholders[1].width
    assert all(not report.overflow for report in renderer.reports)
    assert result.slide_width==source.slide_width
    assert source.slides[0].shapes.title.text=='Original title'


def test_markdown_replacement_does_not_reformat_original_run():
    data,_ = template()
    body = Presentation(io.BytesIO(data)).slides[0].placeholders[1]
    formatting = deepcopy(body.text_frame.paragraphs[0].runs[0]._r.rPr)
    replace_text(body,['**Definition:** an unknown value','A second explanation.'])
    assert body.text=='Definition: an unknown value\nA second explanation.'
    for paragraph in body.text_frame.paragraphs:
        assert xml_signature(paragraph.runs[0]._r.rPr)==xml_signature(formatting)


def test_template_without_slots_is_rejected():
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    data = io.BytesIO()
    prs.save(data)
    with pytest.raises(ValueError,match='placeholder'):
        manifest(data.getvalue())


def test_large_text_is_reported_without_font_shrinking():
    data,spec = template()
    slide = SlideSpec(number=1,layout='concept',purpose='Explain',title='Long content',bullets=[Bullet(text='word '*1000)])
    renderer = DeckRenderer(data,spec)
    output = renderer.render([slide])
    assert renderer.reports[0].overflow
    body = Presentation(io.BytesIO(output)).slides[0].placeholders[1]
    assert body.text_frame.paragraphs[0].runs[0].font.size.pt==20


def test_named_picture_replacement_keeps_geometry_and_original_logo():
    data,spec = template()
    prs = Presentation(io.BytesIO(data))
    image = io.BytesIO()
    Image.new('RGB',(120,80),'red').save(image,'PNG')
    picture = prs.slides[0].shapes.add_picture(io.BytesIO(image.getvalue()), Inches(7), Inches(1), width=Inches(2),height=Inches(2))
    picture.name = 'clastio:image'
    before = (picture.left,picture.top,picture.width,picture.height)
    original = io.BytesIO()
    prs.save(original)
    spec['template_slots'] = manifest(original.getvalue())
    replacement = io.BytesIO()
    Image.new('RGB',(80,80),'blue').save(replacement,'PNG')
    slide = SlideSpec(number=1,layout='concept',purpose='Explain',title='Relevant picture',bullets=[Bullet(text='A useful definition.')],asset_id='photo')
    output = DeckRenderer(original.getvalue(),spec).render([slide],{'photo':replacement.getvalue()})
    result = Presentation(io.BytesIO(output))
    picture = next(shape for shape in result.slides[0].shapes if shape.name=='clastio:image')
    assert picture.image.blob==replacement.getvalue()
    assert (picture.left,picture.top,picture.width,picture.height)==before
    assert next(shape for shape in result.slides[0].shapes if shape.name=='Keep decoration')


async def test_multipart_endpoint_maps_payload_and_stores_unchanged_original(monkeypatch):
    import json
    import uuid
    from types import SimpleNamespace
    from unittest.mock import AsyncMock, Mock

    from fastapi import UploadFile

    from app.api.routes import content
    from app.services import usage

    original,_ = template()
    user = SimpleNamespace(id=uuid.uuid4())
    db = AsyncMock()
    db.add = Mock()
    async def flush():
        for call in db.add.call_args_list:
            if getattr(call.args[0], 'id', None) is None:
                call.args[0].id = uuid.uuid4()
    db.flush.side_effect = flush
    db.execute.return_value = SimpleNamespace(scalars=lambda:SimpleNamespace(first=lambda:None))
    monkeypatch.setattr(content,'upload_limit_mb',AsyncMock(return_value=10))
    monkeypatch.setattr(usage,'check_count_limit',AsyncMock())
    monkeypatch.setattr(usage,'check_storage',AsyncMock())
    uploaded = SimpleNamespace(id=uuid.uuid4(),storage_key='immutable-template.pptx')
    store = AsyncMock(return_value=(uploaded,True))
    monkeypatch.setattr(content,'store_upload',store)
    create = AsyncMock(return_value={'course':{'project_id':'project'}})
    monkeypatch.setattr(content,'create_course',create)
    payload = {'topic':'Algebra','grade':'10','subject':'Mathematics','num_lectures':1,'slides_per_lecture':6}
    result = await content.create_course_from_template(user,db,UploadFile(io.BytesIO(original),filename='base.pptx'),json.dumps(payload))
    assert result['course']['project_id']=='project'
    assert store.call_args.kwargs['data']==original
    assert create.call_args.args[0].template_id is not None
    from app.models import Template
    saved = next(call.args[0] for call in db.add.call_args_list if isinstance(call.args[0], Template))
    assert saved.spec['preserve_slots']
    assert saved.base_storage_key==uploaded.storage_key


def test_template_model_override_is_scoped_and_strict_schema_is_enforced():
    from types import SimpleNamespace

    from app.ai.service import AIService, template_content_model
    from app.generation.specs import LessonDeck

    service = object.__new__(AIService)
    service.settings = SimpleNamespace(model_content='openrouter:other/model')
    token = template_content_model.set('openrouter:openai/gpt-4o-mini')
    try:
        assert service._configured('content',{'model_routing':{'content':'other:model'}})=='openrouter:openai/gpt-4o-mini'
    finally:
        template_content_model.reset(token)
    assert service._configured('content',{})=='openrouter:other/model'
    # Structured requests use the inherited OpenAI strict-schema path (not free-text parsing).
    from app.ai.schema_utils import strict_schema
    schema = strict_schema(LessonDeck)
    assert schema['additionalProperties'] is False
    assert 'slides' in schema['required']


def test_theme_and_master_parts_are_not_rebuilt():
    import zipfile
    data,spec = template()
    slide = SlideSpec(number=1,layout='cover',purpose='Introduce',title='Algebra',bullets=[Bullet(text='A variable is an unknown value.')])
    result = DeckRenderer(data,spec).render([slide])
    with zipfile.ZipFile(io.BytesIO(data)) as before, zipfile.ZipFile(io.BytesIO(result)) as after:
        for name in before.namelist():
            if name.startswith(('ppt/theme/','ppt/slideMasters/','ppt/slideLayouts/')):
                assert before.read(name)==after.read(name), name


async def test_overflow_fails_instead_of_changing_template_fonts(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from test_content_quality import deck

    from app.generation import pipeline
    from app.generation.quality import ContentQualityError

    data,spec = template()
    lesson = deck()
    # A valid but physically tiny slot cannot fit even concise slide text.
    source = Presentation(io.BytesIO(data))
    source.slides[0].placeholders[1].height = Inches(.03)
    output = io.BytesIO()
    source.save(output)
    ai = SimpleNamespace(mode='offline',structured=AsyncMock())
    with pytest.raises(ContentQualityError,match='original font sizes'):
        await pipeline.render_with_qc(ai,base_pptx=output.getvalue(),template_spec=spec,deck=lesson,budgets={},
            images={},language='en',core_props={},context_text='',grade='2',max_rounds=0)


async def test_repeated_template_upload_reuses_template_without_storage_or_profile_charge(monkeypatch):
    import json
    import uuid
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from fastapi import UploadFile

    from app.api.routes import content
    from app.services import usage
    original,_ = template()
    uploaded = SimpleNamespace(storage_key='saved.pptx')
    existing = SimpleNamespace(id=uuid.uuid4())
    db = AsyncMock()
    db.execute.side_effect = [SimpleNamespace(scalars=lambda:SimpleNamespace(first=lambda:uploaded)),
                              SimpleNamespace(scalars=lambda:SimpleNamespace(first=lambda:existing))]
    count = AsyncMock()
    storage = AsyncMock()
    store = AsyncMock()
    monkeypatch.setattr(content,'upload_limit_mb',AsyncMock(return_value=10))
    monkeypatch.setattr(usage,'check_count_limit',count)
    monkeypatch.setattr(usage,'check_storage',storage)
    monkeypatch.setattr(content,'store_upload',store)
    create = AsyncMock(return_value={'course':{}})
    monkeypatch.setattr(content,'create_course',create)
    await content.create_course_from_template(SimpleNamespace(id=uuid.uuid4()),db,
        UploadFile(io.BytesIO(original),filename='base.pptx'),json.dumps({'topic':'Algebra','grade':'10','subject':'Mathematics','num_lectures':1,'slides_per_lecture':6}))
    assert create.call_args.args[0].template_id==existing.id
    count.assert_not_awaited()
    storage.assert_not_awaited()
    store.assert_not_awaited()
