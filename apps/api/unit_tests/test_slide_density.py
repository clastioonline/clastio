import io

import pytest
from pptx import Presentation

from app.engine.render.renderer import DeckRenderer
from app.engine.template.builder import BUILTIN_STYLES, build_builtin
from app.generation import pipeline
from app.generation.budgets import compute_budgets
from app.generation.specs import Bullet, LessonDeck, SlideSpec


def test_sparse_explanation_is_flagged_without_blocking_purposeful_whitespace():
    base, template = build_builtin(next(iter(BUILTIN_STYLES)))
    slide = SlideSpec(number=1, layout='concept', title='Fractions', purpose='Explain',
                      bullets=[Bullet(text='A fraction is part of a whole.')])
    deck = LessonDeck.model_construct(slides=[slide])
    assert any(issue.code == 'sparse_explanation' for issue in pipeline.content_qc(deck, compute_budgets(template)))
    for layout in ['cover', 'section', 'quiz', 'activity']:
        changed = slide.model_copy(update={'layout': layout})
        deck = LessonDeck.model_construct(slides=[changed])
        assert not any(issue.code == 'sparse_explanation' for issue in pipeline.content_qc(deck, compute_budgets(template)))


@pytest.mark.parametrize('use_placeholder', [False, True])
def test_explanation_cards_use_body_area_and_remain_editable(use_placeholder):
    base, template = build_builtin(next(iter(BUILTIN_STYLES)))
    template['content']['use_placeholders'] = use_placeholder
    renderer = DeckRenderer(base, template)
    points = ['A fraction describes equal parts of a whole.', 'The denominator counts all equal parts.',
              'The numerator counts the parts selected.']
    slide = SlideSpec(number=1, layout='concept', title='Fractions', purpose='Explain',
                      bullets=[Bullet(text=text) for text in points])
    data = renderer.render([slide])
    shapes = Presentation(io.BytesIO(data)).slides[0].shapes
    cards = [shape for shape in shapes if shape.has_text_frame and shape.text in points]
    assert len(cards) == 3
    body = renderer.zone('body')
    assert min(shape.top for shape in cards) == body.y
    assert abs(max(shape.top + shape.height for shape in cards) - body.bottom) < 5
    assert not any(report.overflow for report in renderer.reports)


def test_uploaded_template_placeholder_does_not_bypass_content_cards(tmp_path):
    from app.engine.style.pptx_analyzer import analyze_pptx
    from app.engine.template.builder import build_native

    source = tmp_path / 'density-native-source.pptx'
    source.write_bytes(build_builtin(next(iter(BUILTIN_STYLES)))[0])
    try:
        base, template = build_native(source, analyze_pptx(source))
    finally:
        source.unlink()
    template['content']['use_placeholders'] = True
    renderer = DeckRenderer(base, template)
    points = ['The denominator shows the number of equal parts.',
              'The numerator shows how many parts are selected.',
              'Three selected parts out of four makes three quarters.']
    data = renderer.render([SlideSpec(number=1, layout='concept', title='Reading fractions',
                                      purpose='Explain', bullets=[Bullet(text=text) for text in points])])
    shapes = Presentation(io.BytesIO(data)).slides[0].shapes
    cards = [shape for shape in shapes if shape.has_text_frame and shape.text in points]
    assert len(cards) == 3
    assert all(not shape.is_placeholder for shape in cards)
    assert not any(report.overflow for report in renderer.reports)


@pytest.mark.parametrize('layout', ['objectives', 'worked_example'])
def test_short_teaching_sequences_use_available_body_height(layout):
    from app.generation.specs import Step

    base, template = build_builtin(next(iter(BUILTIN_STYLES)))
    # A stage catalogue must not force objectives back into a sparse plain list.
    template['stage_variants'] = {'unused_stage': {}}
    renderer = DeckRenderer(base, template)
    points = ['Identify the variable in 3x + 5.', 'Substitute x = 2 into the expression.',
              'Calculate 3 times 2 plus 5 to get 11.']
    slide = SlideSpec(number=1, layout=layout, title='Working with expressions', purpose='Explain',
                      bullets=[Bullet(text=text) for text in points] if layout=='objectives' else [],
                      steps=[Step(label=f'Step {index+1}', detail=text) for index,text in enumerate(points)]
                      if layout=='worked_example' else [])
    data=renderer.render([slide])
    shapes=Presentation(io.BytesIO(data)).slides[0].shapes
    body=renderer.zone('body')
    bodies=[shape for shape in shapes if shape.has_text_frame and any(text in shape.text for text in points)]
    assert len(bodies)==3
    assert max(shape.top+shape.height for shape in bodies)>=body.bottom-5
    assert not any(report.overflow for report in renderer.reports)
