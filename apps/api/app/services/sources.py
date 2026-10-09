"""Grounding sources: teacher-uploaded textbook chapters / syllabi, chunked + embedded for retrieval."""

from __future__ import annotations

import asyncio
import io
import re
import uuid
from typing import Any

import pymupdf as fitz
from docx import Document as DocxDocument
from pptx import Presentation
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.ai.service import get_ai
from app.core.db import get_sessionmaker
from app.core.storage import get_storage
from app.engine.style.content import extract_content
from app.jobs.queue import PermanentJobError
from app.models import SourceChunk, UploadedFile

CHUNK_CHARS = 900


def extract_pages(data: bytes, ext: str) -> list[tuple[int, str]]:
    pages: list[tuple[int, str]] = []
    if ext == "pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        for i, p in enumerate(doc, start=1):
            pages.append((i, p.get_text("text")))
    elif ext == "pptx":
        prs = Presentation(io.BytesIO(data))
        content, _ = extract_content(prs)
        pages = [(p["number"], "\n".join([p["text"], p["notes"]])) for p in content]
    elif ext == "docx":
        doc = DocxDocument(io.BytesIO(data))
        text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
        pages = [(i + 1, text[i * 3000:(i + 1) * 3000]) for i in range(max(1, len(text) // 3000 + 1))]
    elif ext == "txt":
        text = data.decode("utf-8", "ignore")
        pages = [(i + 1, text[i * 3000:(i + 1) * 3000]) for i in range(max(1, len(text) // 3000 + 1))]
    return [(n, re.sub(r"[ \t]+", " ", t).strip()) for n, t in pages if t and t.strip()]


def chunk(pages: list[tuple[int, str]]) -> list[tuple[int, str]]:
    out: list[tuple[int, str]] = []
    for page, text in pages:
        paras = [p.strip() for p in re.split(r"\n{2,}|\n(?=[A-Z•\-])", text) if p.strip()]
        buf = ""
        for para in paras:
            if len(buf) + len(para) > CHUNK_CHARS and buf:
                out.append((page, buf.strip()))
                buf = ""
            buf += para + "\n"
        if buf.strip():
            out.append((page, buf.strip()))
    return [(page, text[start:start + CHUNK_CHARS]) for page, text in out
            for start in range(0, len(text), CHUNK_CHARS)]


async def index_source(file_id: uuid.UUID) -> dict[str, Any]:
    async with get_sessionmaker()() as s:
        f = await s.get(UploadedFile, file_id)
        if f is None:
            raise ValueError("source not found")
        if f.status == "ready" and (await s.execute(select(SourceChunk.id).where(SourceChunk.file_id == file_id).limit(1))).first():
            return {"cached": True}
        f.status, f.stage = "processing", "Reading document"
        await s.commit()
        owner_id, key = f.owner_id, f.storage_key
    ext = key.rsplit(".", 1)[-1]
    data = await get_storage().get(key)
    if ext == "ppt":
        from app.services.styles import convert_ppt_to_pptx
        data = await asyncio.to_thread(convert_ppt_to_pptx, data)
        ext = "pptx"
    total_pages, chapters = None, []
    if ext == "pdf":
        with fitz.open(stream=data, filetype="pdf") as document:
            total_pages = len(document)
            toc = document.get_toc()
            for index, (level, title, start) in enumerate(toc):
                if level == 1 and start > 0:
                    following = next((page for depth, _, page in toc[index + 1:] if depth == 1 and page > start), total_pages + 1)
                    chapters.append({"title": title, "start": start, "end": following - 1})
    pages = await asyncio.to_thread(extract_pages, data, ext)
    native_page_numbers = {number for number, _ in pages}
    if ext == "pdf":
        from app.services.pdf_ingestion import extract_scanned
        pages = await extract_scanned(data, pages, owner_id)
    chunks = chunk(pages)
    if not chunks:
        raise PermanentJobError("No readable text found. For a scanned book, upload an OCR/text PDF or typed notes.")
    vectors = []
    for offset in range(0, len(chunks), 32):
        vectors.extend(await get_ai().embed([c[1] for c in chunks[offset:offset + 32]], owner_id=owner_id))
    async with get_sessionmaker()() as s:
        await s.execute(delete(SourceChunk).where(SourceChunk.file_id == file_id))
        for (page, text), vec in zip(chunks, vectors, strict=False):
            s.add(SourceChunk(file_id=file_id, owner_id=owner_id, page=page, text=text, embedding=vec))
        f = await s.get(UploadedFile, file_id)
        f.status, f.stage, f.page_count = "ready", "Ready", total_pages or len({c[0] for c in chunks})
        f.meta = {**(f.meta or {}), "chapters": (f.meta or {}).get("chapters") or chapters, "ai_transcribed_pages": [number for number, _ in pages if number not in native_page_numbers]}
        await s.commit()
    return {"chunks": len(chunks)}


async def retrieve(db: AsyncSession, owner_id: uuid.UUID, query: str, k: int = 6,
                   file_ids: list[str] | None = None) -> list[dict[str, Any]]:
    filters = [SourceChunk.owner_id == owner_id]
    if file_ids is not None:
        if not file_ids:
            return []
        filters.append(SourceChunk.file_id.in_([uuid.UUID(str(fid)) for fid in file_ids]))
    has_any = (await db.execute(select(SourceChunk.id).where(*filters).limit(1))).first()
    if not has_any:
        return []
    vec = (await get_ai().embed([query], owner_id=owner_id))[0]
    rows = (await db.execute(
        select(SourceChunk, UploadedFile.filename).join(UploadedFile, UploadedFile.id == SourceChunk.file_id)
        .where(*filters, SourceChunk.embedding.is_not(None), UploadedFile.status == "ready")
        .order_by(SourceChunk.embedding.cosine_distance(vec)).limit(k))).all()
    return [{"file_id": str(c.file_id), "file": fn, "page": c.page, "text": c.text} for c, fn in rows]
