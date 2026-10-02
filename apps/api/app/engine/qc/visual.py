"""Post-render visual QC and previews.

PPTX -> PDF with LibreOffice (headless, isolated profile) -> PyMuPDF. We then check what a viewer would
actually see: text outside the slide, text that escaped its box (overflow), overlapping text lines and
low-contrast text. Page renders double as slide previews for the editor.
"""

from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pymupdf as fitz
from PIL import Image
from pptx import Presentation

from app.core.config import get_settings
from app.engine.style.common import contrast_ratio, rgb_to_hex


class RenderError(Exception):
    pass


def pptx_to_pdf(pptx: bytes, timeout: int | None = None) -> bytes:
    settings = get_settings()
    soffice = shutil.which(settings.soffice_path) or settings.soffice_path
    with tempfile.TemporaryDirectory(prefix="ata-lo-") as tmp:
        src = Path(tmp) / "deck.pptx"
        src.write_bytes(pptx)
        profile = Path(tmp) / "profile"
        cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
               "--convert-to", "pdf", "--outdir", tmp, str(src)]
        try:
            subprocess.run(cmd, capture_output=True, timeout=timeout or settings.render_timeout_s, check=False)
        except subprocess.TimeoutExpired as e:
            raise RenderError("LibreOffice timed out") from e
        out = Path(tmp) / "deck.pdf"
        if not out.exists():
            raise RenderError("LibreOffice could not render the presentation")
        return out.read_bytes()


def docx_to_pdf(docx: bytes, timeout: int | None = None) -> bytes:
    settings = get_settings()
    soffice = shutil.which(settings.soffice_path) or settings.soffice_path
    with tempfile.TemporaryDirectory(prefix="ata-lo-") as tmp:
        src = Path(tmp) / "doc.docx"
        src.write_bytes(docx)
        profile = Path(tmp) / "profile"
        subprocess.run([soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                        "--convert-to", "pdf", "--outdir", tmp, str(src)], capture_output=True,
                       timeout=timeout or settings.render_timeout_s, check=False)
        out = Path(tmp) / "doc.pdf"
        if not out.exists():
            raise RenderError("LibreOffice could not convert the document")
        return out.read_bytes()


@dataclass
class VisualReport:
    pdf: bytes
    previews: list[bytes]
    thumbnails: list[bytes]
    issues: dict[int, list[dict[str, Any]]] = field(default_factory=dict)

    @property
    def failing_slides(self) -> list[int]:
        return sorted(n for n, iss in self.issues.items() if any(i["severity"] == "error" for i in iss))


def _text_boxes(pptx: bytes) -> list[list[tuple[float, float, float, float]]]:
    """Per slide: bounding boxes (points) of every shape that can hold text."""
    prs = Presentation(io.BytesIO(pptx))
    out = []
    for slide in prs.slides:
        boxes = []
        inherited = []
        if slide._element.get("showMasterSp", "1") not in ("0", "false"):
            inherited.extend(sh for sh in slide.slide_layout.shapes if not sh.is_placeholder)
            if slide.slide_layout._element.get("showMasterSp", "1") not in ("0", "false"):
                inherited.extend(sh for sh in slide.slide_layout.slide_master.shapes if not sh.is_placeholder)
        for sh in [*slide.shapes, *inherited]:
            if sh.left is None or sh.width is None:
                continue
            holds_text = (getattr(sh, "has_text_frame", False) and sh.has_text_frame) or \
                (getattr(sh, "has_table", False) and sh.has_table)
            if holds_text:
                boxes.append((sh.left / 12700, sh.top / 12700, (sh.left + sh.width) / 12700,
                              (sh.top + sh.height) / 12700))
        out.append(boxes)
    return out


def _bg_color(pix: fitz.Pixmap, rect: fitz.Rect, scale: float) -> str | None:
    x0, y0 = max(0, int(rect.x0 * scale)), max(0, int(rect.y0 * scale))
    x1, y1 = min(pix.width - 1, int(rect.x1 * scale)), min(pix.height - 1, int(rect.y1 * scale))
    if x1 <= x0 or y1 <= y0:
        return None
    counts: Counter = Counter()
    step = max(1, (x1 - x0) // 12)
    for yy in (y0, (y0 + y1) // 2, y1):
        for xx in range(x0, x1 + 1, step):
            r, g, b = pix.pixel(xx, yy)[:3]
            counts[(r // 8 * 8, g // 8 * 8, b // 8 * 8)] += 1
    (r, g, b), _ = counts.most_common(1)[0]
    return rgb_to_hex((r, g, b))


def inspect(pptx: bytes, *, preview_dpi: int = 72, thumb_width: int = 320) -> VisualReport:
    pdf = pptx_to_pdf(pptx)
    doc = fitz.open(stream=pdf, filetype="pdf")
    boxes = _text_boxes(pptx)
    previews, thumbs = [], []
    issues: dict[int, list[dict[str, Any]]] = {}
    for i, page in enumerate(doc):
        n = i + 1
        pix = page.get_pixmap(dpi=preview_dpi)
        png = pix.tobytes("png")
        previews.append(png)
        im = Image.open(io.BytesIO(png))
        im.thumbnail((thumb_width, thumb_width))
        tb = io.BytesIO()
        im.convert("RGB").save(tb, "WEBP", quality=80)
        thumbs.append(tb.getvalue())
        scale = preview_dpi / 72
        W, H = page.rect.width, page.rect.height
        slide_boxes = boxes[i] if i < len(boxes) else []
        found: list[dict[str, Any]] = []
        spans = []
        for bi, b in enumerate(page.get_text("dict")["blocks"]):
            if b.get("type") != 0:
                continue
            for li, line in enumerate(b["lines"]):
                for sp in line["spans"]:
                    if sp["text"].strip():
                        sp["_line"] = (bi, li)
                        spans.append(sp)
        for sp in spans:
            r = fitz.Rect(sp["bbox"])
            if r.x0 < -2 or r.y0 < -2 or r.x1 > W + 2 or r.y1 > H + 2:
                found.append({"code": "off_slide", "severity": "error", "text": sp["text"][:60]})
                continue
            cx, cy = (r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2
            inside = any(bx0 - 6 <= cx <= bx1 + 6 and by0 - 6 <= cy <= by1 + 6 for bx0, by0, bx1, by1 in slide_boxes)
            if slide_boxes and not inside:
                found.append({"code": "text_overflow", "severity": "error", "text": sp["text"][:60]})
            color = "#{:06X}".format(sp.get("color", 0))
            bg = _bg_color(pix, r, scale)
            decorative = len(sp["text"].strip()) <= 1 and not sp["text"].strip().isalnum()  # bullet glyphs
            if bg and sp["size"] >= 9 and not decorative:
                ratio = contrast_ratio(color, bg)
                if ratio < 2.2:
                    found.append({"code": "low_contrast", "severity": "error" if ratio < 1.6 else "warning",
                                  "text": sp["text"][:60], "ratio": round(ratio, 2)})
        # overlapping lines of text (different spans occupying the same area)
        rects = [fitz.Rect(sp["bbox"]) for sp in spans]
        for a in range(len(rects)):
            for b in range(a + 1, len(rects)):
                if spans[a]["_line"] == spans[b]["_line"]:
                    continue  # same line (e.g. sub/superscripts, font fallback runs)
                if min(len(spans[a]["text"].strip()), len(spans[b]["text"].strip())) <= 2:
                    continue
                inter = rects[a] & rects[b]
                if inter.is_empty:
                    continue
                small = min(rects[a].get_area(), rects[b].get_area()) or 1
                if inter.get_area() / small > 0.35:
                    found.append({"code": "overlapping_text", "severity": "error",
                                  "text": spans[a]["text"][:40] + " / " + spans[b]["text"][:40]})
        if found:
            # de-duplicate by code+text
            seen, uniq = set(), []
            for f in found:
                k = (f["code"], f["text"])
                if k not in seen:
                    seen.add(k)
                    uniq.append(f)
            issues[n] = uniq
    return VisualReport(pdf=pdf, previews=previews, thumbnails=thumbs, issues=issues)
