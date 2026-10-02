"""Low-level PPTX (OOXML) helpers used by the analyzer, template builder and renderer."""

from __future__ import annotations

import copy
from typing import Any

from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationT
from pptx.util import Emu

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
R_ATTRS = (qn("r:embed"), qn("r:link"), qn("r:id"), qn("r:pict"))


def all_slide_layouts(prs: PresentationT) -> list:
    """Stable layout indices across every master, retaining first-master indices."""
    return [layout for master in prs.slide_masters for layout in master.slide_layouts]

SCHEME_SLOTS = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                "hlink", "folHlink"]


def emu_to_pt(v: int | float) -> float:
    return float(v) / 12700.0


def pt_to_emu(v: float) -> Emu:
    return Emu(int(round(v * 12700)))


# --------------------------------------------------------------------------- theme


def theme_part(prs: PresentationT, master=None):
    return (master or prs.slide_master).part.part_related_by(RT.THEME)


def read_theme(prs: PresentationT, master=None) -> dict[str, Any]:
    """Return {'colors': {slot: '#RRGGBB'}, 'fonts': {'major': str, 'minor': str, 'major_cs', 'minor_cs'}}."""
    root = etree.fromstring(theme_part(prs, master).blob)
    colors: dict[str, str] = {}
    scheme = root.find(".//a:clrScheme", NS)
    if scheme is not None:
        for slot in SCHEME_SLOTS:
            el = scheme.find(f"a:{slot}", NS)
            if el is None or len(el) == 0:
                continue
            c = el[0]
            val = c.get("val") if c.tag == qn("a:srgbClr") else c.get("lastClr")
            if val:
                colors[slot] = "#" + val.upper()
    fonts: dict[str, str] = {}
    for kind in ("major", "minor"):
        f = root.find(f".//a:fontScheme/a:{kind}Font", NS)
        if f is not None:
            latin = f.find("a:latin", NS)
            cs = f.find("a:cs", NS)
            if latin is not None and latin.get("typeface"):
                fonts[kind] = latin.get("typeface")
            if cs is not None and cs.get("typeface"):
                fonts[f"{kind}_cs"] = cs.get("typeface")
    return {"colors": colors, "fonts": fonts}


def write_theme(prs: PresentationT, colors: dict[str, str] | None = None, fonts: dict[str, str] | None = None) -> None:
    part = theme_part(prs)
    root = etree.fromstring(part.blob)
    if colors:
        scheme = root.find(".//a:clrScheme", NS)
        for slot, hexval in colors.items():
            el = scheme.find(f"a:{slot}", NS)
            if el is None:
                continue
            for child in list(el):
                el.remove(child)
            srgb = etree.SubElement(el, qn("a:srgbClr"))
            srgb.set("val", hexval.lstrip("#").upper())
    if fonts:
        for kind in ("major", "minor"):
            f = root.find(f".//a:fontScheme/a:{kind}Font", NS)
            if f is None:
                continue
            if fonts.get(kind):
                f.find("a:latin", NS).set("typeface", fonts[kind])
            if fonts.get(f"{kind}_cs"):
                cs = f.find("a:cs", NS)
                if cs is not None:
                    cs.set("typeface", fonts[f"{kind}_cs"])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def resolve_scheme_color(name: str, theme_colors: dict[str, str]) -> str | None:
    mapping = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
    return theme_colors.get(mapping.get(name, name))


# --------------------------------------------------------------------------- slides


def delete_slide(prs: PresentationT, index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst
    sld_id = sld_id_lst[index]
    rid = sld_id.get(qn("r:id"))
    removed_part = prs.part.related_part(rid)
    sld_id_lst.remove(sld_id)
    _drop_slide_references(prs, removed_part)


def _drop_slide_references(prs, removed_part) -> None:
    # Masters can contain navigation links to slides. A dangling relationship keeps a
    # deleted slide in the package, colliding with the names of newly generated slides.
    parts = list(prs.part.package.iter_parts())
    for part in parts:
        for rel in list(part.rels.values()):
            if rel.is_external or rel.target_part is not removed_part:
                continue
            root = getattr(part, "_element", None)
            if root is not None:
                for node in list(root.iter()):
                    for attr in R_ATTRS:
                        if node.get(attr) != rel.rId:
                            continue
                        if node.tag in (qn("a:hlinkClick"), qn("a:hlinkHover"), qn("p:sld")):
                            parent = node.getparent()
                            if parent is not None:
                                parent.remove(node)
                        else:
                            node.attrib.pop(attr, None)
            part.drop_rel(rel.rId)


def prune_unlisted_slide_links(prs: PresentationT) -> None:
    """Repair older cached templates with slides retained only by master navigation."""
    active = {slide.part for slide in prs.slides}
    targets = {rel.target_part for part in prs.part.package.iter_parts() for rel in part.rels.values()
               if not rel.is_external and rel.reltype == RT.SLIDE and rel.target_part not in active}
    for target in targets:
        _drop_slide_references(prs, target)


def delete_all_slides_except(prs: PresentationT, keep: list[int]) -> None:
    for i in reversed(range(len(prs.slides))):
        if i not in keep:
            delete_slide(prs, i)


def move_slide(prs: PresentationT, old_index: int, new_index: int) -> None:
    lst = prs.slides._sldIdLst
    el = lst[old_index]
    lst.remove(el)
    lst.insert(new_index, el)


def copy_element_with_rels(element, src_part, dst_part):
    """Deep-copy an XML element from one slide to another, re-creating relationships (images, links)."""
    new_el = copy.deepcopy(element)
    for node in new_el.iter():
        for attr in R_ATTRS:
            rid = node.get(attr)
            if not rid:
                continue
            try:
                rel = src_part.rels[rid]
            except KeyError:
                node.attrib.pop(attr, None)
                continue
            if rel.is_external:
                new_rid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, new_rid)
    return new_el


def copy_background(src_slide, dst_slide) -> bool:
    bg = src_slide._element.cSld.find(qn("p:bg"))
    if bg is None:
        return False
    new_bg = copy_element_with_rels(bg, src_slide.part, dst_slide.part)
    csld = dst_slide._element.cSld
    existing = csld.find(qn("p:bg"))
    if existing is not None:
        csld.remove(existing)
    csld.insert(0, new_bg)
    return True


# --------------------------------------------------------------------------- text helpers


def set_bullet(paragraph, char: str | None = "•", color_hex: str | None = None, indent_emu: int = 285750,
               level: int = 0) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if char is None:
        etree.SubElement(pPr, qn("a:buNone"))
        pPr.set("marL", str(int(level * indent_emu)))
        pPr.set("indent", "0")
        return
    left = indent_emu * (level + 1)
    pPr.set("marL", str(int(left)))
    pPr.set("indent", str(-int(indent_emu * 0.9)))
    if color_hex:
        clr = etree.SubElement(pPr, qn("a:buClr"))
        s = etree.SubElement(clr, qn("a:srgbClr"))
        s.set("val", color_hex.lstrip("#").upper())
    font = etree.SubElement(pPr, qn("a:buFont"))
    font.set("typeface", "Arial")
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def set_rtl(paragraph, rtl: bool) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    if rtl:
        pPr.set("rtl", "1")
    elif "rtl" in pPr.attrib:
        del pPr.attrib["rtl"]


def set_line_spacing(paragraph, pct: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for el in pPr.findall(qn("a:lnSpc")):
        pPr.remove(el)
    ln = etree.Element(qn("a:lnSpc"))
    pct_el = etree.SubElement(ln, qn("a:spcPct"))
    pct_el.set("val", str(int(pct * 100000)))
    pPr.insert(0, ln)


def set_space_after(paragraph, pt: float) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for el in pPr.findall(qn("a:spcAft")):
        pPr.remove(el)
    sa = etree.SubElement(pPr, qn("a:spcAft"))
    pts = etree.SubElement(sa, qn("a:spcPts"))
    pts.set("val", str(int(pt * 100)))


def set_cs_font(run, typeface: str) -> None:
    """Complex-script (Arabic etc.) font on a run."""
    rPr = run._r.get_or_add_rPr()
    for el in rPr.findall(qn("a:cs")):
        rPr.remove(el)
    cs = etree.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", typeface)


def set_shape_alt_text(shape, text: str) -> None:
    cNvPr = shape._element.find(".//" + qn("p:cNvPr"))
    if cNvPr is not None:
        cNvPr.set("descr", text[:1000])
