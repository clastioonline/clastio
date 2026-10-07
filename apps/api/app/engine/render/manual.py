"""Editable slide-object formatting; coordinates use fractions of the slide canvas."""
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt


def apply_manual(slide, edits, width, height):
    shapes = {(shape.name.removeprefix("ClastioCustom:") if shape.name.startswith("ClastioCustom:") else str(shape.shape_id)): shape for shape in slide.shapes}
    for key, edit in edits.items():
        shape = shapes.get(key)
        if shape is None and key.startswith("custom-") and edit.object_type:
            left, top = int((edit.x or 0.1) * width), int((edit.y or 0.1) * height)
            w, h = int((edit.width or 0.4) * width), int((edit.height or 0.15) * height)
            shape = (slide.shapes.add_textbox(left, top, w, h) if edit.object_type == "text" else
                     slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if edit.object_type == "rectangle" else MSO_SHAPE.OVAL, left, top, w, h))
            shape.name = "ClastioCustom:" + key
        if shape is None:
            continue
        if edit.hidden:
            shape._element.getparent().remove(shape._element)
            continue
        if edit.rotation is not None:
            shape.rotation = edit.rotation
        if edit.fill is not None and hasattr(shape, "fill"):
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(edit.fill.lstrip("#"))
        if edit.layer:
            parent = shape._element.getparent()
            parent.remove(shape._element)
            parent.insert(2 if edit.layer == "back" else len(parent), shape._element)
        for name, scale in (("x", width), ("y", height), ("width", width), ("height", height)):
            value = getattr(edit, name)
            if value is not None:
                setattr(shape, {"x": "left", "y": "top"}.get(name, name), int(value * scale))
        shape.left = max(0, min(shape.left, width - shape.width))
        shape.top = max(0, min(shape.top, height - shape.height))
        frames = [shape.text_frame] if shape.has_text_frame else []
        if shape.has_table:
            frames = [cell.text_frame for row in shape.table.rows for cell in row.cells]
        if edit.text is not None and shape.has_text_frame:
            # Keep the first paragraph/run properties when changing text.
            paragraph = shape.text_frame.paragraphs[0]
            first = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            first.text = edit.text
            for run in list(paragraph.runs)[1:]:
                paragraph._p.remove(run._r)
            for other in list(shape.text_frame.paragraphs)[1:]:
                shape.text_frame._txBody.remove(other._p)
        for frame in frames:
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    if edit.font_size is not None:
                        run.font.size = Pt(edit.font_size)
                    if edit.color is not None:
                        run.font.color.rgb = RGBColor.from_string(edit.color.lstrip("#"))
                    if edit.font_family is not None:
                        run.font.name = edit.font_family
                    if edit.bold is not None:
                        run.font.bold = edit.bold
                    if edit.italic is not None:
                        run.font.italic = edit.italic


def editable_objects(pptx):
    presentation = Presentation(io.BytesIO(pptx))
    slides = []
    for slide in presentation.slides:
        objects = []
        for shape in slide.shapes:
            if shape.width <= 0 or shape.height <= 0:
                continue
            run = next((r for p in shape.text_frame.paragraphs for r in p.runs), None) if shape.has_text_frame else None
            color = None
            if run:
                try:
                    color = "#" + str(run.font.color.rgb) if run.font.color.rgb else None
                except (AttributeError, ValueError):
                    pass
            objects.append({"id": shape.name.removeprefix("ClastioCustom:") if shape.name.startswith("ClastioCustom:") else str(shape.shape_id), "name": shape.name,
                "text": shape.text if shape.has_text_frame else None,
                "kind": "text" if shape.has_text_frame else "table" if shape.has_table else "object",
                "x": shape.left / presentation.slide_width, "y": shape.top / presentation.slide_height,
                "width": shape.width / presentation.slide_width, "height": shape.height / presentation.slide_height,
                "font_size": run.font.size.pt if run and run.font.size else None,
                "font_family": run.font.name if run else None, "color": color,
                "bold": run.font.bold if run else None,
                "italic": run.font.italic if run else None, "rotation": shape.rotation,
                "aspect_ratio": presentation.slide_width / presentation.slide_height})
        slides.append(objects)
    return slides
