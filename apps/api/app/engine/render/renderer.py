"""SlideSpec[] + template -> editable PPTX.

Every slide is built from the teacher's own layouts (native mode) with their decorations cloned in,
then content is laid out in the template's body zone using native, editable shapes. Text is
pre-fitted with real font metrics, so the renderer knows (and reports) whether anything overflows.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field
from typing import Any

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.engine.pptx_xml import (
    all_slide_layouts,
    copy_background,
    copy_element_with_rels,
    delete_slide,
    set_bullet,
    set_cs_font,
    set_line_spacing,
    set_rtl,
    set_shape_alt_text,
)
from app.engine.render.textfit import FitResult, Para, fit, is_arabic
from app.engine.style.common import mix, readable_text_on
from app.generation.specs import SlideSpec

EMU_PER_PT = 12700
RTL_LANGS = {"ar", "ur", "fa", "he"}


@dataclass
class Box:
    x: int
    y: int
    w: int
    h: int

    @property
    def w_pt(self) -> float:
        return self.w / EMU_PER_PT

    @property
    def h_pt(self) -> float:
        return self.h / EMU_PER_PT

    @property
    def right(self) -> int:
        return self.x + self.w

    @property
    def bottom(self) -> int:
        return self.y + self.h

    def inset(self, dx: int, dy: int | None = None) -> Box:
        dy = dx if dy is None else dy
        return Box(self.x + dx, self.y + dy, max(1, self.w - 2 * dx), max(1, self.h - 2 * dy))

    def split_h(self, frac: float, gap: int) -> tuple[Box, Box]:
        w1 = int((self.w - gap) * frac)
        return Box(self.x, self.y, w1, self.h), Box(self.x + w1 + gap, self.y, self.w - w1 - gap, self.h)

    def split_v(self, frac: float, gap: int) -> tuple[Box, Box]:
        h1 = int((self.h - gap) * frac)
        return Box(self.x, self.y, self.w, h1), Box(self.x, self.y + h1 + gap, self.w, self.h - h1 - gap)

    def emu(self) -> tuple[Emu, Emu, Emu, Emu]:
        return Emu(self.x), Emu(self.y), Emu(self.w), Emu(self.h)


@dataclass
class P:
    text: str
    level: int = 0
    bold: bool = False
    color: str | None = None
    scale: float = 1.0
    bullet: bool | None = None


@dataclass
class TextReport:
    role: str
    size_pt: float
    fits: bool
    box: list[int]


@dataclass
class SlideReport:
    number: int
    layout: str
    texts: list[TextReport] = field(default_factory=list)
    issues: list[dict[str, Any]] = field(default_factory=list)

    @property
    def overflow(self) -> bool:
        return any(not t.fits for t in self.texts)

    def to_dict(self) -> dict[str, Any]:
        return {"number": self.number, "layout": self.layout, "overflow": self.overflow,
                "min_font_pt": min((t.size_pt for t in self.texts), default=None),
                "texts": [t.__dict__ for t in self.texts], "issues": self.issues}


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper()[:6])


def no_shadow(shape) -> None:
    """Remove inherited theme effects (shadows) - both the explicit override and the style reference."""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        ref = style.find(qn("a:effectRef"))
        if ref is not None:
            ref.set("idx", "0")


class DeckRenderer:
    def __init__(self, base_pptx: bytes, spec: dict[str, Any], *, language: str = "en",
                 min_font_pt: float | None = None):
        self.prs = Presentation(io.BytesIO(base_pptx))
        self.spec = spec
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height
        self.c = spec["colors"]
        self.f = spec["fonts"]
        self.t = spec["typography"]
        self.rtl = language in RTL_LANGS
        self.language = language
        self.min_pt = float(min_font_pt or self.t.get("min_pt", 16))
        self.rounded = spec.get("corner_radius", "rounded") == "rounded"
        self.donor_count = spec.get("donor_count", 0)
        self.donors = [self.prs.slides[i] for i in range(self.donor_count)]
        self.gap = int(self.W * 0.018)
        self.reports: list[SlideReport] = []
        self._report: SlideReport | None = None

    # ================================================================== geometry helpers

    def zone(self, key: str) -> Box:
        z = self.spec["zones"][key]
        return Box(int(z[0] * self.W), int(z[1] * self.H), int(z[2] * self.W), int(z[3] * self.H))

    def mirror(self, b: Box) -> Box:
        if not self.rtl:
            return b
        return Box(self.W - b.x - b.w, b.y, b.w, b.h)

    # ================================================================== chrome (decorations)

    def _max_shape_id(self, slide) -> int:
        ids = [int(el.get("id")) for el in slide._element.iter(qn("p:cNvPr")) if el.get("id", "").isdigit()]
        return max(ids or [1])

    def _copy_items(self, slide, donor, ids: list[int], number: int | None = None, is_number: bool = False) -> None:
        if donor is None or not ids:
            return
        tree = slide.shapes._spTree
        next_id = self._max_shape_id(slide) + 1
        insert_at = 2
        for el in list(donor.shapes._spTree):
            c = el.find(".//" + qn("p:cNvPr"))
            if c is None or not c.get("id", "").isdigit() or int(c.get("id")) not in ids:
                continue
            new_el = copy_element_with_rels(el, donor.part, slide.part)
            for c2 in new_el.iter(qn("p:cNvPr")):
                c2.set("id", str(next_id))
                next_id += 1
            if is_number and number is not None:
                runs = list(new_el.iter(qn("a:t")))
                if runs:
                    runs[0].text = str(number)
                    for extra in runs[1:]:
                        extra.text = ""
            if is_number:
                tree.append(new_el)
            else:
                tree.insert(insert_at, new_el)
                insert_at += 1

    def _new_slide(self, role: str, number: int):
        cfg = self.spec.get(role) or self.spec["content"]
        layout_index = cfg.get("layout_index")
        layouts = all_slide_layouts(self.prs)
        if layout_index is None or layout_index >= len(layouts):
            layout_index = self.spec["content"]["layout_index"]
        slide = self.prs.slides.add_slide(layouts[layout_index])
        donor_idx = cfg.get("donor_index")
        donor = self.donors[donor_idx] if donor_idx is not None and donor_idx < len(self.donors) else None
        if donor is not None and cfg.get("copy_background"):
            copy_background(donor, slide)
        self._copy_items(slide, donor, cfg.get("item_ids", []))
        self._copy_items(slide, donor, cfg.get("slide_number_ids", []), number=number, is_number=True)
        return slide

    def _placeholder(self, slide, kind: str):
        for ph in slide.placeholders:
            t = ph.placeholder_format.type
            name = str(t).split(".")[-1].split(" ")[0]
            if kind == "title" and name in ("TITLE", "CENTER_TITLE"):
                return ph
            if kind == "body" and name in ("BODY", "OBJECT", "SUBTITLE"):
                return ph
        return None

    @staticmethod
    def _remove_shape(shape) -> None:
        el = shape._element
        el.getparent().remove(el)

    def _cleanup_placeholders(self, slide) -> None:
        for ph in list(slide.placeholders):
            if not ph.has_text_frame or not ph.text_frame.text.strip():
                self._remove_shape(ph)

    # ================================================================== text primitives

    def _style_frame(self, tf, inset_pt: float, anchor: str) -> None:
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        m = Pt(inset_pt)
        tf.margin_left = tf.margin_right = m
        tf.margin_top = tf.margin_bottom = Pt(inset_pt * 0.6)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                              "bottom": MSO_ANCHOR.BOTTOM}[anchor]

    def _fill_frame(self, tf, paras: list[P], size: float, *, family: str, color: str, bold: bool, align: str,
                    bullets: bool, bullet_color: str | None, spacing: float, space_em: float,
                    keep_font: bool = False) -> None:
        tf.clear()
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            arabic = is_arabic(para.text) or self.rtl
            run = p.add_run()
            run.text = para.text
            f = run.font
            f.size = Pt(round(size * para.scale, 1))
            if not keep_font:
                f.name = self.f.get("arabic") if arabic and is_arabic(para.text) else family
            f.bold = para.bold or bold
            if not keep_font or para.color:
                f.color.rgb = _rgb(para.color or color)
            if arabic:
                set_cs_font(run, self.f.get("arabic") or "Noto Sans Arabic")
            use_bullet = bullets if para.bullet is None else para.bullet
            if use_bullet:
                set_bullet(p, "•", bullet_color or self.c["primary"], indent_emu=int(Pt(size * 0.9)),
                           level=para.level)
            elif not keep_font:
                set_bullet(p, None)
            set_rtl(p, arabic)
            al = align
            if arabic and align == "left":
                al = "right"
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[al]
            set_line_spacing(p, spacing)
            p.space_after = Pt(round(size * space_em, 1))

    def text(self, slide, box: Box, paras: list[P], *, role: str = "body", family: str | None = None,
             max_pt: float | None = None, min_pt: float | None = None, color: str | None = None, bold: bool = False,
             align: str = "left", anchor: str = "top", bullets: bool = False, bullet_color: str | None = None,
             spacing: float = 1.0, space_em: float = 0.35, inset_pt: float = 5.0, shape=None,
             keep_font: bool = False) -> tuple[Any, FitResult]:
        family = family or self.f["body"]
        max_pt = max_pt or self.t["body_pt"]
        min_pt = min(min_pt or self.min_pt, max_pt)
        paras = [p for p in paras if p.text and p.text.strip()]
        fit_paras = [Para(p.text, bold=p.bold or bold, size_scale=p.scale,
                          indent_pt=((max_pt * 0.9) * (p.level + 1)) if (bullets if p.bullet is None else p.bullet)
                          else 0, space_after_em=space_em,
                          family=self.f.get("arabic") if is_arabic(p.text) else None) for p in paras]
        result = fit(fit_paras, family, box.w_pt - 2 * inset_pt, box.h_pt - 1.2 * inset_pt, max_pt, min_pt,
                     spacing=spacing)
        if shape is None:
            shape = slide.shapes.add_textbox(*box.emu())
        tf = shape.text_frame
        self._style_frame(tf, inset_pt, anchor)
        self._fill_frame(tf, paras, result.size_pt, family=family, color=color or self.c["text"], bold=bold,
                         align=align, bullets=bullets, bullet_color=bullet_color, spacing=spacing, space_em=space_em,
                         keep_font=keep_font)
        if self._report is not None:
            self._report.texts.append(TextReport(role, result.size_pt, result.fits, [box.x, box.y, box.w, box.h]))
        return shape, result

    # ================================================================== shape primitives

    def card(self, slide, box: Box, *, fill: str | None = None, line: str | None = None, radius: float = 0.12,
             shape_type=None):
        st = shape_type or (MSO_SHAPE.ROUNDED_RECTANGLE if self.rounded else MSO_SHAPE.RECTANGLE)
        sh = slide.shapes.add_shape(st, *box.emu())
        if st == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                sh.adjustments[0] = radius
            except IndexError:
                pass
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(1.25)
        else:
            sh.line.fill.background()
        no_shadow(sh)
        return sh

    def card_text(self, slide, box: Box, paras: list[P], *, fill: str, line: str | None = None, role: str = "card",
                  max_pt: float | None = None, align: str = "left", anchor: str = "middle", bold: bool = False,
                  shape_type=None, color: str | None = None, bullets: bool = False) -> FitResult:
        sh = self.card(slide, box, fill=fill, line=line, shape_type=shape_type)
        _, r = self.text(slide, box, paras, role=role, max_pt=max_pt, color=color or readable_text_on(fill),
                         align=align, anchor=anchor, bold=bold, shape=sh, bullets=bullets, inset_pt=9)
        return r

    def number_badge(self, slide, box: Box, label: str, fill: str | None = None) -> None:
        fill = fill or self.c["primary"]
        sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, *box.emu())
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
        sh.line.fill.background()
        no_shadow(sh)
        self.text(slide, box, [P(label, bold=True)], role="badge", shape=sh, max_pt=min(28, box.h_pt * 0.5),
                  min_pt=10, color=readable_text_on(fill), align="center", anchor="middle", inset_pt=1,
                  family=self.f["heading"])

    def arrow(self, slide, x1: int, y1: int, x2: int, y2: int, color: str | None = None, width_pt: float = 2.25):
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2))
        conn.line.color.rgb = _rgb(color or self.c["secondary"])
        conn.line.width = Pt(width_pt)
        ln = conn.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
        return conn

    def picture(self, slide, box: Box, data: bytes, alt: str = "") -> None:
        with Image.open(io.BytesIO(data)) as im:
            iw, ih = im.size
        pic = slide.shapes.add_picture(io.BytesIO(data), *box.emu())
        img_ratio, box_ratio = iw / ih, box.w / box.h
        if img_ratio > box_ratio:  # too wide -> crop left/right
            excess = 1 - box_ratio / img_ratio
            pic.crop_left = pic.crop_right = excess / 2
        elif img_ratio < box_ratio:
            excess = 1 - img_ratio / box_ratio
            pic.crop_top = pic.crop_bottom = excess / 2
        if alt:
            set_shape_alt_text(pic, alt)

    # ================================================================== titles

    def title(self, slide, spec: SlideSpec) -> None:
        box = self.zone("title")
        ph = self._placeholder(slide, "title") if self.spec["content"].get("use_placeholders") else None
        max_pt = self.t["title_pt"]
        min_pt = max(18.0, min(max_pt * 0.6, 24.0))
        if ph is not None:
            self.text(slide, box, [P(spec.title)], role="title", shape=ph, max_pt=max_pt, min_pt=min_pt,
                      family=self.f["heading"], keep_font=True, anchor="middle", space_em=0,
                      align=self.t.get("title_align", "left") if not self.rtl else "right",
                      color=self.c["title"])
            return
        self.text(slide, box, [P(spec.title)], role="title", family=self.f["heading"], max_pt=max_pt,
                  min_pt=min_pt, color=self.c["title"], bold=self.t.get("title_bold", True), anchor="middle",
                  align=self.t.get("title_align", "left") if not self.rtl else "right", space_em=0)

    # ================================================================== slide kinds

    def render_cover(self, slide, spec: SlideSpec) -> None:
        cz = (self.spec.get("cover") or {}).get("zone")
        if cz and cz.get("title"):
            tb = Box(int(cz["title"][0] * self.W), int(cz["title"][1] * self.H),
                     int(cz["title"][2] * self.W), int(cz["title"][3] * self.H))
            tb = Box(tb.x, tb.y, max(tb.w, int(self.W * 0.55)), max(tb.h, int(self.H * 0.16)))
            if tb.right > self.W * 0.96:
                tb.w = int(self.W * 0.96) - tb.x
            align = cz.get("align", "left")
            self.text(slide, tb, [P(spec.title)], role="title", family=self.f["heading"],
                      max_pt=min(float(cz.get("title_pt") or 44), 54), min_pt=26,
                      color=cz.get("title_color") or self.c["title"], bold=True, align=align, anchor="bottom",
                      space_em=0)
            if spec.subtitle:
                sz = cz.get("subtitle")
                sb = Box(tb.x, tb.bottom + int(self.H * 0.02), tb.w, int(self.H * 0.1)) if not sz else \
                    Box(int(sz[0] * self.W), max(int(sz[1] * self.H), tb.bottom + int(self.H * 0.015)),
                        max(int(sz[2] * self.W), tb.w), max(int(sz[3] * self.H), int(self.H * 0.08)))
                self.text(slide, sb, [P(spec.subtitle)], role="subtitle", max_pt=float(cz.get("subtitle_pt") or 22),
                          min_pt=14, color=cz.get("subtitle_color") or self.c["secondary"], align=align, space_em=0)
            return
        body = self.zone("body")
        tb = Box(body.x, body.y + int(body.h * 0.18), body.w, int(body.h * 0.4))
        self.text(slide, tb, [P(spec.title)], role="title", family=self.f["heading"], max_pt=48, min_pt=26,
                  color=self.c["title"] if not self.spec["content"].get("use_placeholders") else self.c["text"],
                  bold=True, anchor="bottom", align="center")
        if spec.subtitle:
            self.text(slide, Box(body.x, tb.bottom + int(self.H * 0.02), body.w, int(body.h * 0.2)),
                      [P(spec.subtitle)], role="subtitle", max_pt=24, min_pt=14, color=self.c["primary"],
                      align="center")

    def render_section(self, slide, spec: SlideSpec) -> None:
        body = self.zone("body")
        band = Box(body.x, body.y + int(body.h * 0.25), body.w, int(body.h * 0.42))
        self.card(slide, Box(band.x, band.y, int(self.W * 0.012), band.h), fill=self.c["secondary"],
                  shape_type=MSO_SHAPE.RECTANGLE)
        inner = Box(band.x + int(self.W * 0.03), band.y, band.w - int(self.W * 0.03), band.h)
        top, bottom = inner.split_v(0.62, int(self.H * 0.01))
        self.text(slide, top, [P(spec.title)], role="title", family=self.f["heading"], max_pt=44, min_pt=24,
                  color=self.c["heading_accent"], bold=True, anchor="bottom", space_em=0)
        if spec.subtitle or spec.bullets:
            self.text(slide, bottom, [P(spec.subtitle or spec.bullets[0].text)], role="subtitle", max_pt=24,
                      min_pt=14, color=self.c["muted"])

    def _bullets(self, slide, box: Box, spec: SlideSpec, *, max_pt: float | None = None) -> None:
        paras = [P(b.text, level=b.level) for b in spec.bullets] or [P(spec.question or spec.purpose)]
        if max_pt is None and len(paras) <= 4:
            max_pt = self.t["body_pt"] * 1.15
        ph = self._placeholder(slide, "body") if self.spec["content"].get("use_placeholders") else None
        if ph is not None and not self.rtl:
            self.text(slide, box, paras, role="body", shape=ph, keep_font=True, max_pt=max_pt or self.t["body_pt"],
                      bullets=True, color=self.c["text"])
            return
        self.text(slide, box, paras, role="body", max_pt=max_pt or self.t["body_pt"], bullets=True,
                  bullet_color=self.c["primary"])

    def render_concept(self, slide, spec: SlideSpec, image: bytes | None) -> None:
        body = self.zone("body")
        if image is not None or spec.visual.kind in ("image", "diagram"):
            return self.render_image_text(slide, spec, image)
        self._bullets(slide, body, spec)

    def render_image_text(self, slide, spec: SlideSpec, image: bytes | None) -> None:
        body = self.zone("body")
        left, right = body.split_h(0.56, self.gap * 2)
        text_box, img_box = (left, right) if self.spec["zones"].get("image_side", "right") == "right" else (right, left)
        if self.rtl:
            text_box, img_box = img_box, text_box
        self._remove_body_placeholder(slide)
        self.text(slide, text_box, [P(b.text, level=b.level) for b in spec.bullets] or [P(spec.purpose)],
                  role="body", bullets=True, bullet_color=self.c["primary"])
        img_box = Box(img_box.x, img_box.y + int(img_box.h * 0.04), img_box.w, int(img_box.h * 0.92))
        if image is not None:
            self.picture(slide, img_box, image, alt=spec.visual.alt_text or spec.visual.description)
        else:
            self.card_text(slide, img_box, [P(spec.visual.description or spec.title, scale=1.0)],
                           fill=self.c["card_bg"], role="visual", max_pt=18, align="center")

    def _remove_body_placeholder(self, slide) -> None:
        ph = self._placeholder(slide, "body")
        if ph is not None:
            self._remove_shape(ph)

    def render_columns(self, slide, spec: SlideSpec, comparison: bool) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        cols = spec.columns[:3] or [type("C", (), {"heading": "", "bullets": [b.text for b in spec.bullets]})()]
        n = len(cols)
        order = list(range(n))[::-1] if self.rtl else list(range(n))
        col_w = int((body.w - self.gap * 2 * (n - 1)) / n)
        head_h = int(body.h * 0.16)
        for pos, i in enumerate(order):
            col = cols[i]
            x = body.x + pos * (col_w + self.gap * 2)
            head = Box(x, body.y, col_w, head_h)
            rest = Box(x, body.y + head_h + int(self.gap * 0.6), col_w, body.h - head_h - int(self.gap * 0.6))
            fill = self.c["primary"] if (comparison and i == 0) else (self.c["secondary"] if comparison else
                                                                      self.c["primary"])
            self.card_text(slide, head, [P(col.heading or f"Part {i + 1}", bold=True)], fill=fill, role="heading",
                           max_pt=self.t["body_pt"] + 2, align="center", shape_type=None)
            if comparison:
                self.card(slide, rest, fill=mix(fill, self.c["background"], 0.9), line=None)
            self.text(slide, rest.inset(int(self.W * 0.008)), [P(b) for b in col.bullets], role="body",
                      bullets=True, bullet_color=fill, max_pt=self.t["body_pt"])

    def render_numbered_cards(self, slide, spec: SlideSpec, items: list[str] | None = None) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        items = items or [b.text for b in spec.bullets] or [spec.question or spec.purpose]
        items = items[:6]
        n = len(items)
        if n > 4:
            # two columns of cards
            half = math.ceil(n / 2)
            left, right = body.split_h(0.5, self.gap * 2)
            self._card_stack(slide, left, items[:half], start=1)
            self._card_stack(slide, right, items[half:], start=half + 1)
        else:
            self._card_stack(slide, body, items, start=1)

    def _card_stack(self, slide, box: Box, items: list[str], start: int) -> None:
        n = max(1, len(items))
        gap = int(self.H * 0.022)
        h = int((box.h - gap * (n - 1)) / n)
        h = min(h, int(self.H * 0.17))
        total = h * n + gap * (n - 1)
        y = box.y + max(0, (box.h - total) // 2) if n <= 2 else box.y
        badge = int(min(h * 0.62, self.H * 0.085))
        for i, item in enumerate(items):
            row = Box(box.x, y + i * (h + gap), box.w, h)
            self.card(slide, row, fill=self.c["card_bg"])
            bx = row.x + int(self.W * 0.012) if not self.rtl else row.right - int(self.W * 0.012) - badge
            self.number_badge(slide, Box(bx, row.y + (h - badge) // 2, badge, badge), str(start + i))
            tx = row.x + badge + int(self.W * 0.024) if not self.rtl else row.x + int(self.W * 0.012)
            self.text(slide, Box(tx, row.y, row.w - badge - int(self.W * 0.036), h), [P(item)], role="body",
                      color=self.c["card_text"], anchor="middle", max_pt=self.t["body_pt"], space_em=0)

    def render_process(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        steps = spec.steps[:6] or [type("S", (), {"label": b.text, "detail": ""})() for b in spec.bullets[:5]]
        n = max(1, len(steps))
        if n > 5:
            return self.render_worked_example(slide, spec)
        arrow_w = int(self.W * 0.028)
        box_w = int((body.w - arrow_w * (n - 1) - self.gap * 2 * (n - 1)) / n)
        head_h = int(body.h * 0.34)
        top = body.y + int(body.h * 0.06)
        order = list(range(n))[::-1] if self.rtl else list(range(n))
        for pos, i in enumerate(order):
            x = body.x + pos * (box_w + arrow_w + self.gap * 2)
            step = steps[i]
            head = Box(x, top, box_w, head_h)
            fill = self.c["primary"]
            self.card_text(slide, head, [P(step.label, bold=True)], fill=fill, role="step", align="center",
                           max_pt=self.t["body_pt"])
            self.number_badge(slide, Box(x + int(box_w / 2) - int(self.H * 0.03), top - int(self.H * 0.03),
                                         int(self.H * 0.06), int(self.H * 0.06)), str(i + 1), fill=self.c["secondary"])
            if step.detail:
                self.text(slide, Box(x, head.bottom + int(self.H * 0.02), box_w, body.bottom - head.bottom -
                                     int(self.H * 0.02)), [P(step.detail)], role="detail", align="center",
                          max_pt=max(self.min_pt, self.t["body_pt"] - 4))
            if pos < n - 1:
                ax = x + box_w + self.gap
                ay = top + head_h // 2
                arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW if not self.rtl else MSO_SHAPE.LEFT_ARROW,
                                             Emu(ax), Emu(ay - int(self.H * 0.025)), Emu(arrow_w),
                                             Emu(int(self.H * 0.05)))
                arr.fill.solid()
                arr.fill.fore_color.rgb = _rgb(self.c["secondary"])
                arr.line.fill.background()
                no_shadow(arr)

    def render_cycle(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        steps = spec.steps[:6] or [type("S", (), {"label": b.text, "detail": ""})() for b in spec.bullets[:6]]
        n = max(3, len(steps)) if steps else 3
        steps = list(steps) + [type("S", (), {"label": "…", "detail": ""})()] * (n - len(steps))
        cx, cy = body.x + body.w // 2, body.y + body.h // 2
        node_w = int(min(body.w / (2.6 if n <= 4 else 3.2), self.W * 0.26))
        node_h = int(min(body.h / (3.2 if n <= 4 else 3.6), self.H * 0.2))
        rx = (body.w - node_w) / 2
        ry = (body.h - node_h) / 2
        centers = []
        for i in range(n):
            ang = -math.pi / 2 + 2 * math.pi * i / n * (-1 if self.rtl else 1)
            centers.append((cx + rx * math.cos(ang), cy + ry * math.sin(ang)))
        for i, (x, y) in enumerate(centers):
            b = Box(int(x - node_w / 2), int(y - node_h / 2), node_w, node_h)
            st = steps[i]
            paras = [P(st.label, bold=True)] + ([P(st.detail, scale=0.8)] if st.detail else [])
            self.card_text(slide, b, paras, fill=self.c["primary"] if i % 2 == 0 else self.c["card_bg"],
                           role="node", align="center", max_pt=self.t["body_pt"] - 2)
        for i in range(n):
            (x1, y1), (x2, y2) = centers[i], centers[(i + 1) % n]
            dx, dy = x2 - x1, y2 - y1
            dist = math.hypot(dx, dy) or 1
            # shorten arrows so they sit between the nodes
            shrink_x = node_w * 0.55 * abs(dx) / dist
            shrink_y = node_h * 0.55 * abs(dy) / dist
            s = max(shrink_x, shrink_y)
            ux, uy = dx / dist, dy / dist
            self.arrow(slide, int(x1 + ux * s), int(y1 + uy * s), int(x2 - ux * s), int(y2 - uy * s))

    def render_timeline(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        steps = spec.steps[:6] or [type("S", (), {"label": b.text, "detail": ""})() for b in spec.bullets[:6]]
        n = max(1, len(steps))
        line_y = body.y + int(body.h * 0.42)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(body.x), Emu(line_y - int(self.H * 0.004)),
                                      Emu(body.w), Emu(int(self.H * 0.008)))
        line.fill.solid()
        line.fill.fore_color.rgb = _rgb(self.c["primary"])
        line.line.fill.background()
        no_shadow(line)
        seg = body.w / n
        dot = int(self.H * 0.045)
        order = list(range(n))[::-1] if self.rtl else list(range(n))
        for pos, i in enumerate(order):
            st = steps[i]
            cx = int(body.x + seg * pos + seg / 2)
            self.number_badge(slide, Box(cx - dot // 2, line_y - dot // 2, dot, dot), str(i + 1),
                              fill=self.c["secondary"])
            lab = Box(int(cx - seg / 2 + self.gap / 2), body.y, int(seg - self.gap), line_y - dot - body.y)
            self.text(slide, lab, [P(st.label, bold=True)], role="label", align="center", anchor="bottom",
                      color=self.c["heading_accent"], max_pt=self.t["body_pt"])
            if st.detail:
                det = Box(lab.x, line_y + dot, lab.w, body.bottom - line_y - dot)
                self.text(slide, det, [P(st.detail)], role="detail", align="center",
                          max_pt=max(self.min_pt, self.t["body_pt"] - 4))

    def render_worked_example(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        steps = spec.steps[:6]
        if not steps:
            return self._bullets(slide, body, spec)
        paras = []
        for i, st in enumerate(steps):
            paras.append(P(f"{i + 1}. {st.label}", bold=True, color=self.c["heading_accent"], bullet=False))
            if st.detail:
                paras.append(P(st.detail, bullet=False))
        if spec.bullets:
            left, right = body.split_h(0.62, self.gap * 2)
            self.text(slide, left, paras, role="body", space_em=0.25)
            self.card_text(slide, right, [P(b.text) for b in spec.bullets], fill=self.c["card_bg"], role="note",
                           anchor="top", bullets=True)
        else:
            self.text(slide, body, paras, role="body", space_em=0.25)

    def render_table(self, slide, spec: SlideSpec, rows: list[list[str]], headers: list[str]) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        cols = max(len(headers), max((len(r) for r in rows), default=1))
        rows = [list(r) + [""] * (cols - len(r)) for r in rows[:8]]
        headers = list(headers) + [""] * (cols - len(headers))
        if self.rtl:
            headers = headers[::-1]
            rows = [r[::-1] for r in rows]
        n_rows = len(rows) + 1
        col_w = body.w / cols
        # Fit one font size for the whole table so rows never grow past the zone.
        def table_height(sz: float) -> float:
            return sum(max(fit([Para(c or " ")], self.f["body"], col_w / EMU_PER_PT - 14, 10_000, sz, sz).height_pt
                           for c in r) + 10 for r in [headers] + rows)

        size = self.t["body_pt"]
        while size > self.min_pt and table_height(size) > body.h_pt:
            size -= 1
        fits = table_height(size) <= body.h_pt
        shape = slide.shapes.add_table(n_rows, cols, *body.emu())
        table = shape.table
        tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
        tbl_pr.set("bandRow", "0")
        for ci in range(cols):
            table.columns[ci].width = Emu(int(col_w))
        row_h = int(body.h / n_rows)
        for ri in range(n_rows):
            table.rows[ri].height = Emu(row_h)
            for ci in range(cols):
                cell = table.cell(ri, ci)
                val = headers[ci] if ri == 0 else rows[ri - 1][ci]
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(self.c["primary"] if ri == 0 else (
                    self.c["card_bg"] if ri % 2 == 0 else mix(self.c["card_bg"], self.c["background"], 0.6)))
                tf = cell.text_frame
                tf.word_wrap = True
                cell.margin_left = cell.margin_right = Pt(6)
                cell.margin_top = cell.margin_bottom = Pt(3)
                self._fill_frame(tf, [P(val or "", bold=ri == 0 or (ci == 0 and spec.layout == "key_vocabulary"))],
                                 size, family=self.f["body"], color=self.c["on_primary"] if ri == 0
                                 else self.c["card_text"], bold=False, align="left" if not self.rtl else "right",
                                 bullets=False, bullet_color=None, spacing=1.0, space_em=0)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        if self._report is not None:
            self._report.texts.append(TextReport("table", size, fits, [body.x, body.y, body.w, body.h]))

    def render_quiz(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        q = spec.quiz
        question = q.question if q else (spec.question or spec.title)
        options = (q.options if q else [b.text for b in spec.bullets])[:4]
        qbox, obox = body.split_v(0.3, int(self.H * 0.02))
        self.text(slide, qbox, [P(question, bold=True)], role="question", max_pt=self.t["body_pt"] + 4,
                  anchor="middle", color=self.c["text"])
        letters = "ABCD"
        if not options:
            return
        cols = 2 if len(options) > 2 else 1
        rows = math.ceil(len(options) / cols)
        cw = int((obox.w - self.gap * 2 * (cols - 1)) / cols)
        rh = int((obox.h - self.gap * (rows - 1)) / rows)
        rh = min(rh, int(self.H * 0.16))
        for i, opt in enumerate(options):
            r, c = divmod(i, cols)
            if self.rtl:
                c = cols - 1 - c
            b = Box(obox.x + c * (cw + self.gap * 2), obox.y + r * (rh + self.gap), cw, rh)
            self.card(slide, b, fill=self.c["card_bg"], line=self.c["primary"])
            badge = int(min(rh * 0.6, self.H * 0.08))
            bx = b.x + int(self.W * 0.01) if not self.rtl else b.right - int(self.W * 0.01) - badge
            self.number_badge(slide, Box(bx, b.y + (rh - badge) // 2, badge, badge), letters[i])
            tx = b.x + badge + int(self.W * 0.02) if not self.rtl else b.x + int(self.W * 0.01)
            self.text(slide, Box(tx, b.y, b.w - badge - int(self.W * 0.03), rh), [P(opt)], role="option",
                      color=self.c["card_text"], anchor="middle", space_em=0)

    def render_discussion(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        question = spec.question or (spec.bullets[0].text if spec.bullets else spec.title)
        prompts = [b.text for b in spec.bullets if b.text != question][:3]
        qh = 0.55 if prompts else 0.8
        qbox = Box(body.x + int(body.w * 0.06), body.y + int(body.h * 0.04), int(body.w * 0.88), int(body.h * qh))
        self.card_text(slide, qbox, [P(question, bold=True)], fill=self.c["primary"], role="question",
                       max_pt=min(36, self.t["body_pt"] + 10), align="center")
        if prompts:
            pb = Box(qbox.x, qbox.bottom + int(self.H * 0.03), qbox.w, body.bottom - qbox.bottom - int(self.H * 0.03))
            self.text(slide, pb, [P(p) for p in prompts], role="prompts", bullets=True,
                      bullet_color=self.c["secondary"], max_pt=self.t["body_pt"] - 2)

    def render_activity(self, slide, spec: SlideSpec) -> None:
        self._remove_body_placeholder(slide)
        body = self.zone("body")
        side, main = body.split_h(0.27, self.gap * 2)
        if self.rtl:
            side, main = Box(body.right - side.w, side.y, side.w, side.h), Box(body.x, main.y, main.w, main.h)
        mins = int(round(spec.timing_minutes)) if spec.timing_minutes else 10
        info = [P(f"{mins} min", bold=True, scale=1.4)]
        if spec.subtitle:
            info.append(P(spec.subtitle))
        self.card_text(slide, Box(side.x, side.y, side.w, int(side.h * 0.55)), info, fill=self.c["secondary"],
                       role="activity-info", align="center", max_pt=self.t["body_pt"])
        steps = [s.label + (f": {s.detail}" if s.detail else "") for s in spec.steps] or [b.text for b in spec.bullets]
        self._card_stack(slide, main, steps[:5] or [spec.purpose], start=1)

    # ================================================================== notes

    @staticmethod
    def notes_text(spec: SlideSpec) -> str:
        parts = []
        if spec.speaker_notes:
            parts.append(spec.speaker_notes.strip())
        if spec.teacher_instruction:
            parts.append(f"Teacher: {spec.teacher_instruction.strip()}")
        if spec.question_to_ask:
            parts.append(f"Ask: {spec.question_to_ask.strip()}")
        if spec.quiz:
            letter = "ABCD"[spec.quiz.answer_index] if 0 <= spec.quiz.answer_index < 4 else "?"
            parts.append(f"Answer: {letter}. {spec.quiz.explanation}".strip())
        if spec.differentiation.support:
            parts.append(f"Support: {spec.differentiation.support}")
        if spec.differentiation.extension:
            parts.append(f"Stretch: {spec.differentiation.extension}")
        if spec.visual.kind != "none" and spec.visual.description and not spec.asset_id:
            parts.append(f"Visual suggestion: {spec.visual.description}")
        if spec.timing_minutes:
            parts.append(f"Timing: about {spec.timing_minutes:g} min")
        credits = [src.get("attribution") for src in spec.sources if src.get("type") == "image" and
                   src.get("attribution")]
        refs = [f"{src['file']} p.{src['page']}" for src in spec.sources if src.get("file")]
        if credits:
            parts.append("Image credit: " + "; ".join(credits))
        if refs:
            parts.append("Sources: " + "; ".join(refs))
        return "\n\n".join(parts)

    # ================================================================== main

    def _section_role(self) -> str:
        sec = self.spec.get("section") or {}
        if sec.get("layout_index") is not None and "title" in (self._layout_types(sec["layout_index"])):
            return "section"
        return "cover"

    def _layout_types(self, index: int) -> list[str]:
        for m in self.spec.get("layouts", []):
            if m["index"] == index:
                return m.get("placeholder_types", [])
        return []

    def render_section_slide(self, spec: SlideSpec, number: int):
        role = self._section_role()
        if role == "section":
            cfg = self.spec["section"]
            slide = self.prs.slides.add_slide(all_slide_layouts(self.prs)[cfg["layout_index"]])
            ph = self._placeholder(slide, "title")
            if ph is not None:
                self.text(slide, Box(ph.left, ph.top, ph.width, ph.height), [P(spec.title)], role="title", shape=ph,
                          keep_font=True, max_pt=40, min_pt=22, color=self.c["title"], anchor="bottom")
            sub = self._placeholder(slide, "body")
            if sub is not None and spec.subtitle:
                self.text(slide, Box(sub.left, sub.top, sub.width, sub.height), [P(spec.subtitle)], role="subtitle",
                          shape=sub, keep_font=True, max_pt=22, min_pt=14, color=self.c["muted"])
            self._cleanup_placeholders(slide)
            return slide
        slide = self._new_slide("cover", number)
        self._cleanup_placeholders(slide)
        self.render_cover(slide, spec)
        return slide

    def render_slide(self, spec: SlideSpec, image: bytes | None, number: int) -> None:
        self._report = SlideReport(number=number, layout=spec.layout)
        if spec.layout == "section":
            slide = self.render_section_slide(spec, number)
            notes = self.notes_text(spec)
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            self.reports.append(self._report)
            self._report = None
            return
        role = "cover" if spec.layout == "cover" else "content"
        slide = self._new_slide(role, number)
        if role == "cover":
            if not self.spec.get("cover", {}).get("use_placeholders"):
                self._cleanup_placeholders(slide)
                self.render_cover(slide, spec)
            else:
                ph = self._placeholder(slide, "title")
                if ph is not None:
                    self.text(slide, Box(ph.left, ph.top, ph.width, ph.height), [P(spec.title)], role="title",
                              shape=ph, keep_font=True, max_pt=44, min_pt=24, color=self.c["title"], anchor="middle")
                sub = self._placeholder(slide, "body")
                if sub is not None and spec.subtitle:
                    self.text(slide, Box(sub.left, sub.top, sub.width, sub.height), [P(spec.subtitle)],
                              role="subtitle", shape=sub, keep_font=True, max_pt=24, min_pt=14, color=self.c["text"])
                self._cleanup_placeholders(slide)
        else:
            if spec.layout != "section":
                self.title(slide, spec)
            kind = spec.layout
            if kind == "section":
                self._remove_body_placeholder(slide)
                self.render_section(slide, spec)
            elif kind in ("objectives", "summary", "exit_ticket", "homework"):
                items = [b.text for b in spec.bullets]
                if kind == "exit_ticket" and spec.question and spec.question not in items:
                    items = [spec.question] + items
                if len(items) > 6 or sum(len(i) for i in items) > 420:
                    self._bullets(slide, self.zone("body"), spec)
                else:
                    self.render_numbered_cards(slide, spec, items)
            elif kind == "concept":
                self.render_concept(slide, spec, image)
            elif kind == "image_text":
                self.render_image_text(slide, spec, image)
            elif kind in ("two_column", "comparison"):
                self.render_columns(slide, spec, comparison=kind == "comparison")
            elif kind == "process":
                self.render_process(slide, spec)
            elif kind == "cycle":
                self.render_cycle(slide, spec)
            elif kind == "timeline":
                self.render_timeline(slide, spec)
            elif kind == "table":
                t = spec.table
                self.render_table(slide, spec, t.rows if t else [], t.headers if t else ["", ""])
            elif kind == "key_vocabulary":
                has_tr = any(t.translation for t in spec.terms)
                headers = ["Term", "Meaning"] + (["Translation"] if has_tr else [])
                rows = [[t.term, t.meaning] + ([t.translation or ""] if has_tr else []) for t in spec.terms]
                if rows:
                    self.render_table(slide, spec, rows, headers)
                else:
                    self._bullets(slide, self.zone("body"), spec)
            elif kind == "quiz":
                self.render_quiz(slide, spec)
            elif kind == "discussion":
                self.render_discussion(slide, spec)
            elif kind == "activity":
                self.render_activity(slide, spec)
            elif kind == "worked_example":
                self.render_worked_example(slide, spec)
            else:
                self._bullets(slide, self.zone("body"), spec)
            self._cleanup_placeholders(slide)
        notes = self.notes_text(spec)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        self.reports.append(self._report)
        self._report = None

    def render(self, slides: list[SlideSpec], images: dict[str, bytes] | None = None,
               core_props: dict[str, str] | None = None) -> bytes:
        images = images or {}
        for i, spec in enumerate(slides, start=1):
            self.render_slide(spec, images.get(spec.asset_id) if spec.asset_id else None, i)
        # Remove donor slides (they sit at the start of the template base).
        for _ in range(self.donor_count):
            delete_slide(self.prs, 0)
        cp = self.prs.core_properties
        props = core_props or {}
        cp.title = props.get("title", "")
        cp.author = props.get("author", "")
        cp.last_modified_by = props.get("author", "")
        cp.subject = props.get("subject", "")
        cp.keywords = props.get("keywords", "")
        cp.comments = props.get("comments", "")
        cp.category = ""
        out = io.BytesIO()
        self.prs.save(out)
        return out.getvalue()
