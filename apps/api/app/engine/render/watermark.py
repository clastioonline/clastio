"""Teacher-selected attribution as an editable native PowerPoint text box."""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


def apply_watermark(presentation, props):
    text = str(props.get('watermark_text') or '').strip()[:120]
    teacher = str(props.get('teacher_footer_text') or '').strip()[:80]
    numbers = props.get('page_numbers') == 'true'
    width, height = presentation.slide_width, presentation.slide_height
    diagonal = props.get('watermark_position') == 'diagonal'
    for number, slide in enumerate(presentation.slides, start=1):
        for shape in list(slide.shapes):
            page_number = (shape.name in {'Teacher page number', 'Template page number'} or
                           shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)
            if shape.name in {'Teacher watermark', 'Teacher name footer'} or ('page_numbers' in props and page_number):
                shape._element.getparent().remove(shape._element)
        if teacher:
            _footer(slide, width, height, 'Teacher name footer', teacher, .58, .34, PP_ALIGN.RIGHT)
        if numbers:
            _footer(slide, width, height, 'Teacher page number', f'{number} / {len(presentation.slides)}', .04, .12, PP_ALIGN.LEFT)
        if not text:
            continue
        box = slide.shapes.add_textbox(int(width * .08), int(height * (.42 if diagonal else .94)),
                                      int(width * (.84 if diagonal or not teacher and not numbers else .40)), int(height * (.16 if diagonal else .05)))
        if not diagonal and (teacher or numbers):
            box.left = int(width * .18)
        box.name = 'Teacher watermark'
        box.rotation = 330 if diagonal else 0
        frame = box.text_frame
        frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER if diagonal else PP_ALIGN.RIGHT
        run = paragraph.add_run()
        run.text = text
        run.font.name = 'Arial'
        run.font.size = Pt(28 if diagonal else 10)
        run.font.color.rgb = RGBColor.from_string('64748B')
        color = run._r.get_or_add_rPr().find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '25000' if diagonal else '70000')
        color.append(alpha)


def _footer(slide, width, height, name, text, x, box_width, alignment):
    box = slide.shapes.add_textbox(int(width * x), int(height * .94), int(width * box_width), int(height * .05))
    box.name = name
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = alignment
    run = paragraph.add_run()
    run.text = text
    run.font.name = 'Arial'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string('334155')
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    fill_color = box.fill._xPr.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
    alpha = OxmlElement('a:alpha')
    alpha.set('val', '85000')
    fill_color.append(alpha)
