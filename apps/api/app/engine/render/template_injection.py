"""Populate real template slots while preserving original shape XML and theme parts."""
from __future__ import annotations

import copy
import io

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from app.engine.pptx_xml import copy_element_with_rels, delete_slide
from app.engine.render.textfit import Para, fit
from app.engine.render.vendor.presenton import infographic_markdown_to_plain_text, repeated_child_source_index


def role(shape):
    name = shape.name.lower().strip().removeprefix('clastio:')
    text = shape.text.strip().lower() if shape.has_text_frame else ''
    for key in ['title', 'subtitle', 'body', 'image']:
        if shape.name.lower().startswith('clastio:') and name == key or text == '{{'+key+'}}':
            return key
    if shape.is_placeholder:
        kind = shape.placeholder_format.type
        if kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            return 'title'
        if kind == PP_PLACEHOLDER.SUBTITLE:
            return 'subtitle'
        if kind in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT) and shape.has_text_frame:
            return 'body'
        if kind == PP_PLACEHOLDER.PICTURE:
            return 'image'
    return None


def manifest(data):
    prs = Presentation(io.BytesIO(data))
    result = []
    if not 1 <= len(prs.slides) <= 30:
        raise ValueError('Use a template with 1–30 slides.')
    for index, slide in enumerate(prs.slides):
        slots = [{'shape_id': shape.shape_id, 'role': role(shape), 'width_pt': shape.width/12700,
                  'height_pt': shape.height/12700} for shape in slide.shapes if role(shape)]
        if len(slots) > 32:
            raise ValueError('Use at most 32 content slots per template slide.')
        if any(slot['role'] == 'title' for slot in slots):
            result.append({'index': index, 'slots': slots})
    if not result or not any(any(slot['role'] == 'body' for slot in page['slots']) for page in result):
        raise ValueError('The template needs a title and body placeholder. You may also name text boxes clastio:title and clastio:body; decorative shapes are preserved.')
    return result


def replace_text(shape, lines):
    """Clone existing paragraph/run styling; only replace the text nodes."""
    frame = shape.text_frame
    originals = [copy.deepcopy(paragraph._p) for paragraph in frame.paragraphs]
    for child in list(frame._txBody):
        if child.tag == qn('a:p'):
            frame._txBody.remove(child)
    for index, line in enumerate(lines or ['']):
        paragraph = copy.deepcopy(originals[min(index, len(originals)-1)])
        nodes = paragraph.findall('.//'+qn('a:t'))
        if not nodes:
            from pptx.oxml.xmlchemy import OxmlElement
            run = OxmlElement('a:r')
            node = OxmlElement('a:t')
            run.append(node)
            paragraph.append(run)
            nodes = [node]
        # Retain each run and its formatting; replace sample text without changing style XML.
        nodes[0].text = str(infographic_markdown_to_plain_text(line))
        for node in nodes[1:]:
            node.text = ''
        frame._txBody.append(paragraph)


def teaching_lines(slide):
    lines = [bullet.text for bullet in slide.bullets]
    lines += [f'{term.term}: {term.meaning}' for term in slide.terms]
    lines += [f'{step.label}: {step.detail}' for step in slide.steps]
    for column in slide.columns:
        lines += [column.heading] + list(column.bullets)
    if slide.table:
        lines += [' | '.join(slide.table.headers)] + [' | '.join(map(str,row)) for row in slide.table.rows]
    if slide.quiz:
        lines += [slide.quiz.question] + list(slide.quiz.options)
    if slide.question:
        lines.append(slide.question)
    return lines


def inject(prs, slots, slides, images, spec, core_props):
    from app.engine.render.manual import apply_manual
    from app.engine.render.renderer import SlideReport, TextReport

    donors = list(prs.slides)
    reports = []
    body_pages = [page for page in slots if any(slot['role'] == 'body' for slot in page['slots'])]
    for index, content in enumerate(slides):
        page = slots[0] if index == 0 else body_pages[repeated_child_source_index(index-1,
            template_count=len(body_pages), content_count=len(slides)-1, center_when_reduced=False)]
        donor = donors[page['index']]
        slide = prs.slides.add_slide(donor.slide_layout)
        cloned = copy_element_with_rels(donor._element, donor.part, slide.part)
        slide._element[:] = list(cloned)
        slide._element.attrib.update(cloned.attrib)
        # add_slide cached its initial placeholder tree; resolve shapes from the cloned tree.
        slide.__dict__.pop("shapes", None)
        slide.__dict__.pop("placeholders", None)
        report = SlideReport(content.number,content.layout)
        bodies = [slot for slot in page['slots'] if slot['role'] == 'body']
        lines = teaching_lines(content)
        body_index = 0
        for slot in page['slots']:
            shape = next(shape for shape in slide.shapes if shape.shape_id == slot['shape_id'])
            if slot['role'] == 'image':
                data = images.get(content.asset_id)
                if not data:
                    from app.generation.quality import ContentQualityError
                    raise ContentQualityError(f'Slide {content.number}: no relevant image was found for the template image slot. Choose another image source or provide a teacher image.')
                if data:
                    if shape.is_placeholder and hasattr(shape,'insert_picture'):
                        shape.insert_picture(io.BytesIO(data))
                    elif shape._element.tag == qn('p:pic'):
                        _, rid = slide.part.get_or_add_image_part(io.BytesIO(data))
                        shape._element.find('.//'+qn('a:blip')).set(qn('r:embed'),rid)
                        # Crop to the existing frame instead of stretching the fetched image.
                        with Image.open(io.BytesIO(data)) as image:
                            source_ratio = image.width / image.height
                        target_ratio = shape.width / shape.height
                        horizontal = max(0, (1-target_ratio/source_ratio)/2)
                        vertical = max(0, (1-source_ratio/target_ratio)/2)
                        shape.crop_left = shape.crop_right = horizontal
                        shape.crop_top = shape.crop_bottom = vertical
                continue
            if not shape.has_text_frame:
                continue
            if slot['role'] == 'title':
                values = [content.title]
            elif slot['role'] == 'subtitle':
                values = [content.subtitle or (content.purpose if index == 0 else '')]
            else:
                start = len(lines)*body_index//len(bodies)
                end = len(lines)*(body_index+1)//len(bodies)
                values = lines[start:end]
                body_index += 1
            paragraph = shape.text_frame.paragraphs[0]
            font = paragraph.runs[0].font if paragraph.runs else paragraph.font
            size = font.size.pt if font.size else spec['typography'].get('body_pt',24)
            family = font.name or spec['fonts']['body']
            measurement = fit([Para(text=str(line)) for line in values],family,
                max(1,(shape.width-shape.text_frame.margin_left-shape.text_frame.margin_right)/12700),
                max(1,(shape.height-shape.text_frame.margin_top-shape.text_frame.margin_bottom)/12700),size,size)
            replace_text(shape,values)
            report.texts.append(TextReport(slot['role'],size,measurement.fits,[shape.left,shape.top,shape.width,shape.height]))
        slide.notes_slide.notes_text_frame.text = content.speaker_notes
        apply_manual(slide,content.manual_objects,prs.slide_width,prs.slide_height)
        reports.append(report)
    for _ in donors:
        delete_slide(prs,0)
    for key in ['title','author','subject','keywords','comments']:
        if key in core_props:
            setattr(prs.core_properties,key,core_props[key])
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue(),reports
