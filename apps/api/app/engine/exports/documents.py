"""DOCX builders (python-docx) for worksheets, quizzes, tests, homework, lesson plans and teacher guides,
plus quiz exports (Kahoot/Quizizz-style CSV, Moodle GIFT). PDFs are produced by LibreOffice."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.generation.specs import AssessmentDoc, HomeworkDoc, LessonPlan, Question, SlideSpec

LETTERS = "ABCDEFGH"


def _rgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _base(accent: str, font: str, rtl: bool = False) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(11.5)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:cs"), "Noto Sans Arabic")
    for lvl, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[lvl]
        st.font.name = font
        st.font.size = Pt(size)
        st.font.color.rgb = _rgb(accent)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)
        section.top_margin = section.bottom_margin = Pt(50)
        if rtl:
            section._sectPr.append(OxmlElement("w:bidi"))
    return doc


def _header_block(doc: Document, title: str, subtitle: str, accent: str, student_fields: bool) -> None:
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph(subtitle)
    p.runs[0].font.color.rgb = _rgb("#6B7280")
    if student_fields:
        t = doc.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, label in zip(t.rows[0].cells, ("Name: ____________________", "Class: ________",
                                                  "Date: ____________"), strict=False):
            cell.text = label
        doc.add_paragraph()


def _answer_lines(doc: Document, n: int) -> None:
    for _ in range(n):
        doc.add_paragraph("_" * 86)


def _question(doc: Document, i: int, q: Question, *, key: bool) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{i}. ")
    r.bold = True
    p.add_run(q.stem)
    marks = p.add_run(f"   [{q.marks} mark{'s' if q.marks != 1 else ''}]")
    marks.font.size = Pt(9)
    marks.font.color.rgb = _rgb("#6B7280")
    if q.qtype in ("mcq", "true_false") and q.options:
        for j, opt in enumerate(option for option in q.options if option.strip()):
            op = doc.add_paragraph(f"      {LETTERS[j]})  {opt}")
            op.paragraph_format.space_after = Pt(0)
            if key and opt.strip().lower() == q.answer.strip().lower():
                op.runs[0].bold = True
    if key:
        a = doc.add_paragraph()
        a.add_run("Answer: ").bold = True
        a.add_run(q.answer)
        if q.explanation:
            e = doc.add_paragraph(q.explanation)
            e.runs[0].italic = True
            e.runs[0].font.color.rgb = _rgb("#4B5563")
    elif q.qtype not in ("mcq", "true_false"):
        _answer_lines(doc, {"short_answer": 2, "fill_blank": 1, "long_answer": 6, "application": 4,
                            "case_based": 4, "matching": 1}.get(q.qtype, 2))
    doc.add_paragraph()


def assessment_docx(a: AssessmentDoc, *, subtitle: str, accent: str, font: str, key: bool, rtl: bool = False) -> bytes:
    doc = _base(accent, font, rtl)
    _header_block(doc, a.title + (" — Answer key" if key else ""), subtitle, accent, student_fields=not key)
    if a.instructions:
        doc.add_paragraph(a.instructions).runs[0].italic = True
    total = sum(q.marks for s in a.sections for q in s.questions)
    doc.add_paragraph(f"Total: {total} marks")
    if a.case_study:
        doc.add_heading("Case study", level=2)
        doc.add_paragraph(a.case_study)
    n = 1
    for section in a.sections:
        doc.add_heading(section.title, level=2)
        if section.instructions:
            doc.add_paragraph(section.instructions).runs[0].italic = True
        for q in section.questions:
            _question(doc, n, q, key=key)
            n += 1
    if key and a.rubric:
        doc.add_heading("Marking guidance", level=2)
        for r in a.rubric:
            doc.add_paragraph(r, style="List Bullet")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def homework_docx(h: HomeworkDoc, *, subtitle: str, accent: str, font: str, key: bool, rtl: bool = False) -> bytes:
    doc = _base(accent, font, rtl)
    _header_block(doc, h.title + (" — Answer key" if key else ""), subtitle + f" • About {h.estimated_minutes} min",
                  accent, student_fields=not key)
    doc.add_paragraph(h.instructions).runs[0].italic = True
    if h.tasks:
        doc.add_heading("Tasks", level=2)
        for t in h.tasks:
            doc.add_paragraph(t, style="List Number")
    if h.questions:
        doc.add_heading("Questions", level=2)
        for i, q in enumerate(h.questions, start=1):
            _question(doc, i, q, key=key)
    if h.support_hint:
        doc.add_heading("Need a hint?", level=3)
        doc.add_paragraph(h.support_hint)
    if h.extension_task:
        doc.add_heading("Challenge", level=3)
        doc.add_paragraph(h.extension_task)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _shade(cell, hexcolor: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor.lstrip("#"))
    tcPr.append(shd)


def lesson_plan_docx(plan: LessonPlan, *, meta: dict[str, Any], accent: str, font: str,
                     outcomes: list[str] | None = None) -> bytes:
    """Inspection-friendly lesson plan (objectives, success criteria, timings, differentiation, AfL)."""
    doc = _base(accent, font)
    doc.add_heading(f"Lesson plan: {plan.title}", level=1)
    cells = [("Teacher", meta.get("teacher", "")), ("Class", meta.get("class", "")),
             ("Subject", meta.get("subject", "")), ("Grade", meta.get("grade", "")),
             ("Date", meta.get("date", "")), ("Duration", f"{plan.duration_minutes} min"),
             ("Curriculum", meta.get("curriculum", "")), ("Lesson", f"{plan.lecture_number} of {meta.get('total', '')}")]
    info = doc.add_table(rows=(len(cells) + 1) // 2, cols=4)
    info.style = "Table Grid"
    for idx, (k, v) in enumerate(cells):
        r, c = divmod(idx, 2)
        info.cell(r, c * 2).text = k
        _shade(info.cell(r, c * 2), "F3F4F6")
        info.cell(r, c * 2 + 1).text = str(v)
    doc.add_heading("WALT — We Are Learning To", level=2)
    for o in plan.objectives:
        doc.add_paragraph(o, style="List Bullet")
    doc.add_heading("WILF — What I’m Looking For", level=2)
    for o in plan.success_criteria:
        doc.add_paragraph(o, style="List Bullet")
    if outcomes:
        doc.add_heading("Curriculum outcomes", level=2)
        for o in outcomes:
            doc.add_paragraph(o, style="List Bullet")
    doc.add_heading("Lesson sequence", level=2)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for cell, label in zip(t.rows[0].cells, ("Phase / time", "Teacher", "Students", "Notes"), strict=False):
        cell.text = label
        _shade(cell, accent)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].runs[0].bold = True
    for ph in plan.phases:
        row = t.add_row().cells
        row[0].text = f"{ph.name.title()} ({ph.minutes} min)"
        row[1].text = ph.teacher_actions
        row[2].text = ph.student_actions
        row[3].text = ph.description
    doc.add_heading("Differentiation", level=2)
    d = plan.differentiation
    dt = doc.add_table(rows=5, cols=2)
    dt.style = "Table Grid"
    for i, (k, v) in enumerate((("Support", d.support), ("Core", d.core), ("Extension / G&T", d.extension),
                                ("EAL learners", d.eal), ("Students of determination", d.send))):
        dt.cell(i, 0).text = k
        _shade(dt.cell(i, 0), "F3F4F6")
        dt.cell(i, 1).text = v
    doc.add_heading("Assessment for learning", level=2)
    for q in plan.questions_to_ask:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(q.question).bold = True
        p.add_run(f"  ({q.bloom_level}) — expected: {q.expected_response}")
    if plan.exit_ticket:
        doc.add_paragraph("Exit ticket: " + " | ".join(plan.exit_ticket))
    if plan.misconceptions:
        doc.add_heading("Common misconceptions", level=2)
        for m in plan.misconceptions:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(m.misconception).bold = True
            p.add_run(f" → {m.correction}")
    doc.add_heading("Resources", level=2)
    doc.add_paragraph(", ".join(plan.resources) or "—")
    if plan.safety_notes:
        doc.add_heading("Health & safety", level=2)
        doc.add_paragraph(plan.safety_notes)
    doc.add_heading("Homework", level=2)
    doc.add_paragraph(f"{plan.homework.task} (about {plan.homework.estimated_minutes} min)")
    if meta.get("local_links"):
        doc.add_heading("Links to UAE national identity / moral education", level=2)
        for link in meta["local_links"]:
            doc.add_paragraph(link, style="List Bullet")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def teacher_guide_docx(slides: list[SlideSpec], plan: LessonPlan | None, *, title: str, accent: str,
                       font: str) -> bytes:
    doc = _base(accent, font)
    doc.add_heading(f"Teacher guide: {title}", level=1)
    if plan:
        doc.add_paragraph(f"{plan.duration_minutes} minutes • {len(slides)} slides")
    for s in slides:
        doc.add_heading(f"Slide {s.number}: {s.title}", level=2)
        meta = doc.add_paragraph()
        meta.add_run(f"{s.layout.replace('_', ' ').title()} • about {s.timing_minutes:g} min").italic = True
        if s.speaker_notes:
            doc.add_paragraph(s.speaker_notes)
        if s.teacher_instruction:
            p = doc.add_paragraph()
            p.add_run("Do: ").bold = True
            p.add_run(s.teacher_instruction)
        if s.question_to_ask:
            p = doc.add_paragraph()
            p.add_run("Ask: ").bold = True
            p.add_run(s.question_to_ask)
        if s.quiz:
            p = doc.add_paragraph()
            p.add_run("Answer: ").bold = True
            letter = LETTERS[s.quiz.answer_index] if 0 <= s.quiz.answer_index < len(s.quiz.options) else "?"
            p.add_run(f"{letter} — {s.quiz.explanation}")
        if s.differentiation.support or s.differentiation.extension:
            p = doc.add_paragraph()
            p.add_run("Differentiate: ").bold = True
            p.add_run(" / ".join(x for x in (s.differentiation.support, s.differentiation.extension) if x))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def all_questions(doc: AssessmentDoc | HomeworkDoc) -> list[Question]:
    if isinstance(doc, HomeworkDoc):
        return list(doc.questions)
    return [q for s in doc.sections for q in s.questions]


def kahoot_csv(questions: list[Question], time_limit: int = 20) -> bytes:
    """Kahoot/Quizizz-style spreadsheet: Question, Answer 1-4, Time limit, Correct answer(s)."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Question", "Answer 1", "Answer 2", "Answer 3", "Answer 4", "Time limit (sec)", "Correct answer(s)"])
    for q in questions:
        if q.qtype not in ("mcq", "true_false") or len(q.options) < 2:
            continue
        opts = (q.options + ["", "", "", ""])[:4]
        correct = next((i + 1 for i, o in enumerate(opts) if o.strip().lower() == q.answer.strip().lower()), 1)
        w.writerow([q.stem[:120], *[o[:75] for o in opts], time_limit, correct])
    return buf.getvalue().encode("utf-8-sig")


def _gift_escape(s: str) -> str:
    return re.sub(r"([~=#{}:])", r"\\\1", s)


def moodle_gift(questions: list[Question], title: str) -> bytes:
    lines = [f"// {title}", ""]
    for i, q in enumerate(questions, start=1):
        stem = _gift_escape(q.stem)
        lines.append(f"::Q{i}:: {stem} {{")
        if q.qtype == "true_false":
            lines.append("  TRUE" if q.answer.strip().lower().startswith("t") else "  FALSE")
        elif q.qtype == "mcq" and q.options:
            for o in q.options:
                mark = "=" if o.strip().lower() == q.answer.strip().lower() else "~"
                lines.append(f"  {mark}{_gift_escape(o)}")
        elif q.qtype == "fill_blank":
            lines.append(f"  ={_gift_escape(q.answer)}")
        else:
            lines.append("")  # essay
        lines.append("}")
        lines.append("")
    return "\n".join(lines).encode("utf-8")
