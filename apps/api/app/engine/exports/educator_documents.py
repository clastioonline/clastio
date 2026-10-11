"""Printable, editable educator drafts; PDF conversion uses the configured renderer.

These builders never assert regulatory approval and never fetch model-provided URLs.
They keep source details, evidence gaps and teacher answer keys in the export.
"""

from __future__ import annotations

import asyncio
import io
import subprocess
from collections.abc import Mapping
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pydantic import BaseModel

from app.engine.exports.documents import _base, _shade
from app.engine.qc.visual import RenderError, docx_to_pdf
from app.generation.educator_specs import (
    DemoClassContent,
    GeneratedResourceContent,
    UAEPlanContent,
)

ACCENT = "#0F766E"
DRAFT_LABEL = "AI-generated teaching draft — teacher review required."


def _data(value: BaseModel | Mapping[str, Any]) -> dict[str, Any]:
    return value.model_dump(mode="json") if isinstance(value, BaseModel) else dict(value)


def _new_document(language: str) -> Document:
    rtl = language.split("-", 1)[0].lower() in {"ar", "ur", "fa", "he"}
    doc = _base(ACCENT, "Noto Sans Arabic" if rtl else "Arial", rtl=rtl)
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = DRAFT_LABEL
        for run in footer.runs:
            run.font.size = Pt(8)
    return doc


def _paragraph(doc: Document, label: str, value: Any) -> None:
    if value is None or value == "" or value == []:
        return
    paragraph = doc.add_paragraph()
    if label:
        paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(str(value))


def _list(doc: Document, values: list[Any], *, numbered: bool = False) -> None:
    for value in values:
        doc.add_paragraph(str(value), style="List Number" if numbered else "List Bullet")


def _table(doc: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
        _shade(cell, "E6F4F1")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            cell.text = str(value)


def _metadata(doc: Document, value: dict[str, Any], *, include_framework: bool = False) -> None:
    fields = [("Topic", "topic"), ("Subject", "subject"), ("Grade / year", "grade"),
              ("Curriculum", "curriculum"), ("Education stage", "education_stage"),
              ("Language", "language")]
    if include_framework:
        fields.append(("Requested alignment framework", "framework"))
    for label, key in fields:
        if value.get(key):
            _paragraph(doc, label, str(value[key]).replace("_", " "))
    if value.get("duration_minutes"):
        _paragraph(doc, "Duration", f"{value['duration_minutes']} minutes")
    doc.add_paragraph(DRAFT_LABEL)


def _sources(doc: Document, sources: list[dict[str, Any]], *, title: str = "Source references") -> None:
    doc.add_heading(title, level=2)
    if not sources:
        doc.add_paragraph("No source references supplied. Verify the material before classroom use.")
    for source in sources:
        _paragraph(doc, "Source", source.get("title", ""))
        _paragraph(doc, "URL / saved file ID", source.get("url_or_file_id", ""))
        _paragraph(doc, "Page / section", source.get("page_or_section", ""))
        _paragraph(doc, "Relevance", source.get("relevance", ""))


def _review(doc: Document, review: dict[str, Any]) -> None:
    doc.add_heading("Teacher review", level=2)
    doc.add_paragraph(DRAFT_LABEL)
    if review.get("evidence_gaps"):
        doc.add_heading("Evidence gaps", level=3)
        _list(doc, review["evidence_gaps"])
    doc.add_heading("Checks before use", level=3)
    _list(doc, review.get("teacher_checks", []))
    if review.get("framework_references"):
        _sources(doc, review["framework_references"], title="Framework references")


def _append_plan(doc: Document, plan: dict[str, Any]) -> None:
    doc.add_heading("Learning objectives and measurable outcomes", level=2)
    for objective in plan["learning_objectives"]:
        _paragraph(doc, objective["id"], objective["description"])
        _paragraph(doc, "Observable action", objective["observable_action"])
        _paragraph(doc, "Success measure", objective["success_measure"])
        _paragraph(doc, "Bloom level", objective["bloom_level"])
    doc.add_heading("Success criteria", level=2)
    for criterion in plan["success_criteria"]:
        _paragraph(doc, criterion["id"], criterion["description"])
        _paragraph(doc, "Objectives", ", ".join(criterion["objective_ids"]))
        _paragraph(doc, "Observable evidence", criterion["observable_evidence"])
    doc.add_heading("Prerequisite knowledge", level=2)
    _list(doc, plan["prerequisites"])
    doc.add_heading("Cross-curricular links", level=2)
    for link in plan["cross_curricular_links"]:
        _paragraph(doc, "Subject", link["subject"])
        _paragraph(doc, "Activity", link["activity"])
        _paragraph(doc, "Evidence of learning", link["evidence_of_learning"])
    doc.add_heading("UAE context and classroom connections", level=2)
    for key, label in [("national_identity", "National identity"), ("moral_education", "Moral education"),
                       ("arabic_culture", "Arabic culture")]:
        connection = plan["uae_context"][key]
        doc.add_heading(label, level=3)
        _paragraph(doc, "Relevant to this lesson", "Yes" if connection["relevant"] else "No")
        _paragraph(doc, "Connection / reason", connection["connection_or_reason"])
        _paragraph(doc, "Classroom application", connection["classroom_application"])
    doc.add_heading("Differentiation and inclusion", level=2)
    for key, label in [("sen", "SEN / students of determination"), ("ell", "ELL / EAL language support")]:
        doc.add_heading(label, level=3)
        for strategy in plan["differentiation"][key]:
            _paragraph(doc, "Learning barrier", strategy["learning_barrier"])
            _paragraph(doc, "Scaffold", strategy["scaffold"])
            _paragraph(doc, "Assessment accommodation", strategy["assessment_accommodation"])
    _paragraph(doc, "Core task", plan["differentiation"]["core"])
    _paragraph(doc, "Extension task", plan["differentiation"]["extension"])
    doc.add_heading("Formative assessments and expected responses", level=2)
    for assessment in plan["formative_assessment"]:
        doc.add_heading(f"{assessment['id']} — {assessment['tool']}", level=3)
        _paragraph(doc, "Objectives", ", ".join(assessment["objective_ids"]))
        for key, label in [("question_or_task", "Question / task"), ("expected_response", "Expected response"),
                           ("evidence_to_record", "Evidence to record"), ("teaching_response", "Teaching response")]:
            _paragraph(doc, label, assessment[key])
    doc.add_heading("Timed lesson sequence", level=2)
    _table(doc, ["Phase / minutes", "Teacher actions", "Student actions", "Alignment"], [
        [f"{phase['name']} ({phase['duration_minutes']:g} min)", phase["teacher_actions"],
         phase["student_actions"], "Objectives: " + ", ".join(phase["objective_ids"]) +
         "\nAssessments: " + (", ".join(phase["assessment_ids"]) or "No assessment in this phase")]
        for phase in plan["lesson_sequence"]
    ])
    doc.add_heading("Resources", level=2)
    _list(doc, plan["resources"])
    doc.add_heading("Homework", level=2)
    doc.add_paragraph(plan["homework"])
    _sources(doc, plan["source_references"])
    _review(doc, plan["review_checks"])


def _all_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in doc.sections:
        yield from section.footer.paragraphs


def _save(doc: Document, language: str) -> bytes:
    if language.split("-", 1)[0].lower() in {"ar", "ur", "fa", "he"}:
        for paragraph in _all_paragraphs(doc):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph_properties = paragraph._p.get_or_add_pPr()
            if paragraph_properties.find(qn("w:bidi")) is None:
                paragraph_properties.append(OxmlElement("w:bidi"))
            for run in paragraph.runs:
                properties = run._r.get_or_add_rPr()
                if properties.find(qn("w:rtl")) is None:
                    properties.append(OxmlElement("w:rtl"))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def build_plan_docx(plan: UAEPlanContent | dict[str, Any], title: str | None = None) -> bytes:
    value = UAEPlanContent.model_validate(_data(plan)).model_dump(mode="json")
    doc = _new_document(value["language"])
    doc.add_heading(title or value["title"], level=1)
    _metadata(doc, value, include_framework=True)
    _append_plan(doc, value)
    return _save(doc, value["language"])


def _clock(seconds: int) -> str:
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def _append_slide(doc: Document, slide: dict[str, Any]) -> None:
    doc.add_heading(f"Slide {slide['number']}: {slide['title']}", level=3)
    _paragraph(doc, "Time", f"{slide['timing_minutes']:g} minutes")
    _paragraph(doc, "Purpose", slide["purpose"])
    _paragraph(doc, "Subtitle", slide.get("subtitle"))
    _list(doc, [bullet["text"] for bullet in slide["bullets"]])
    for column in slide["columns"]:
        _paragraph(doc, "Column", column["heading"])
        _list(doc, column["bullets"])
    for step in slide["steps"]:
        _paragraph(doc, step["label"], step["detail"])
    for term in slide["terms"]:
        _paragraph(doc, term["term"], term["meaning"])
        _paragraph(doc, "Translation", term.get("translation"))
    if slide.get("table"):
        _table(doc, slide["table"]["headers"], slide["table"]["rows"])
    if slide.get("chart"):
        chart = slide["chart"]
        _table(doc, ["Category", *[series["name"] for series in chart["series"]]], [
            [category, *[series["values"][index] for series in chart["series"]]]
            for index, category in enumerate(chart["categories"])
        ])
        _paragraph(doc, "Data source", chart["source"])
        _paragraph(doc, "Units", chart.get("unit"))
    for key, label in [("speaker_notes", "Teacher notes"), ("teacher_instruction", "Teacher instruction"),
                       ("question", "Student question"), ("question_to_ask", "Question to ask")]:
        _paragraph(doc, label, slide.get(key))
    if slide.get("visual"):
        _paragraph(doc, "Visual description", slide["visual"].get("description"))
        _paragraph(doc, "Visual alternative text", slide["visual"].get("alt_text"))
    if slide.get("quiz"):
        quiz = slide["quiz"]
        _paragraph(doc, "Quiz question", quiz["question"])
        _list(doc, quiz["options"], numbered=True)
        index = quiz["answer_index"]
        answer = quiz["options"][index] if 0 <= index < len(quiz["options"]) else "Verify the answer index."
        _paragraph(doc, "Expected answer", answer)
        _paragraph(doc, "Explanation", quiz["explanation"])
    for key, scaffold in slide["differentiation"].items():
        _paragraph(doc, key.replace("_", " ").title(), scaffold)


def build_demo_docx(demo: DemoClassContent | dict[str, Any]) -> bytes:
    value = DemoClassContent.model_validate(_data(demo)).model_dump(mode="json")
    plan = value["lesson_plan"]
    doc = _new_document(plan["language"])
    doc.add_heading(f"15-minute demo class: {plan['title']}", level=1)
    _metadata(doc, plan, include_framework=True)
    doc.add_heading("Interactive hook", level=2)
    _paragraph(doc, "Teacher says", value["hook"]["teacher_says"])
    _paragraph(doc, "Learner interaction", value["hook"]["learner_interaction"])
    _list(doc, value["hook"]["materials"])
    doc.add_heading("Step-by-step teacher script", level=2)
    for step in value["teacher_script"]:
        doc.add_heading(f"{step['order']}. {_clock(step['start_seconds'])}–{_clock(step['end_seconds'])}", level=3)
        _paragraph(doc, "Slides", ", ".join(str(number) for number in step["slide_numbers"]))
        for key, label in [("teacher_says", "Teacher says"), ("teacher_does", "Teacher does"),
                           ("expected_student_response", "Expected student response"),
                           ("formative_check", "Formative check"), ("transition", "Transition")]:
            _paragraph(doc, label, step[key])
    doc.add_page_break()
    doc.add_heading("Demo lesson plan", level=1)
    _append_plan(doc, plan)
    doc.add_page_break()
    doc.add_heading("Matching slide teaching notes", level=1)
    for slide in value["slide_payload"]["slides"]:
        _append_slide(doc, slide)
    return _save(doc, plan["language"])


def build_resource_docx(resource: GeneratedResourceContent | dict[str, Any]) -> bytes:
    value = GeneratedResourceContent.model_validate(_data(resource)).model_dump(mode="json")
    doc = _new_document(value["language"])
    doc.add_heading(value["title"], level=1)
    _metadata(doc, value)
    _paragraph(doc, "Resource type", value["kind"].replace("_", " "))
    doc.add_paragraph("AI-generated learning resource; this is not an official or publisher textbook.")
    doc.add_heading("Overview", level=2)
    doc.add_paragraph(value["overview"])
    doc.add_heading("Contents", level=2)
    _list(doc, [f"{chapter['number']}. {chapter['title']}" for chapter in value["chapters"]])
    for chapter in value["chapters"]:
        doc.add_page_break()
        doc.add_heading(f"Chapter {chapter['number']}: {chapter['title']}", level=1)
        doc.add_heading("Learning goals", level=2)
        _list(doc, chapter["learning_goals"])
        doc.add_heading("Explanation", level=2)
        doc.add_paragraph(chapter["explanation"])
        doc.add_heading("Key definitions", level=2)
        _list(doc, chapter["definitions"])
        doc.add_heading("Worked examples", level=2)
        _list(doc, chapter["worked_examples"], numbered=True)
        doc.add_heading("UAE case study", level=2)
        doc.add_paragraph(chapter["uae_case_study"])
        doc.add_heading("Summary", level=2)
        doc.add_paragraph(chapter["summary"])
        doc.add_heading("Practice questions", level=2)
        _list(doc, chapter["practice_questions"], numbered=True)
        _sources(doc, chapter["source_references"])
    doc.add_page_break()
    doc.add_heading("Teacher answer key", level=1)
    for chapter in value["chapters"]:
        doc.add_heading(f"Chapter {chapter['number']}: {chapter['title']}", level=2)
        for index, (question, answer) in enumerate(zip(chapter["practice_questions"], chapter["expected_answers"], strict=True), 1):
            _paragraph(doc, f"{index}. Question", question)
            _paragraph(doc, "Expected answer", answer)
    _sources(doc, value["source_references"])
    _review(doc, value["review_checks"])
    return _save(doc, value["language"])


def _structured_section(doc: Document, label: str, content: Any, *, level: int = 2) -> None:
    if content is None or content == "" or content == [] or content == {}:
        return
    doc.add_heading(label, level=level)
    if isinstance(content, Mapping):
        for key, value in content.items():
            readable = str(key).replace("_", " ").title()
            if isinstance(value, (Mapping, list)):
                _structured_section(doc, readable, value, level=min(level + 1, 3))
            else:
                _paragraph(doc, readable, value)
    elif isinstance(content, list):
        for item in content:
            if isinstance(item, Mapping):
                for key, value in item.items():
                    _paragraph(doc, str(key).replace("_", " ").title(), value)
            else:
                doc.add_paragraph(str(item), style="List Bullet")
    else:
        doc.add_paragraph(str(content))


def build_task_docx(task: BaseModel | dict[str, Any]) -> bytes:
    """Export a copilot learning task with a separate teacher answer-key section."""
    value = _data(task)
    language = value.get("language") or "en"
    doc = _new_document(language)
    doc.add_heading(value["title"], level=1)
    _metadata(doc, value)
    for label, key in [("Instructions", "instructions"), ("Reading text", "reading_text"),
                       ("Vocabulary", "vocabulary"), ("Differentiation and support", "support"),
                       ("UAE connection", "uae_link")]:
        _structured_section(doc, label, value.get(key))
    doc.add_heading("Questions", level=2)
    for index, question in enumerate(value.get("questions", []), 1):
        _paragraph(doc, str(index), question["question"])
        doc.add_paragraph("________________________________________________________________")
        doc.add_paragraph("________________________________________________________________")
    doc.add_page_break()
    doc.add_heading("Teacher answer key and assessment", level=1)
    for index, question in enumerate(value.get("questions", []), 1):
        _paragraph(doc, f"{index}. Question", question["question"])
        _paragraph(doc, "Expected answer", question["expected_answer"])
    _structured_section(doc, "Assessment guidance", value.get("assessment"))
    _sources(doc, value.get("source_references", []))
    review = value.get("review_checks") or {}
    if isinstance(review, Mapping):
        _review(doc, dict(review))
    else:
        _structured_section(doc, "Teacher review", review)
    return _save(doc, language)


async def educator_docx_to_pdf(data: bytes) -> bytes:
    """Do not block the event loop or silently claim PDF success on render failure."""
    try:
        pdf = await asyncio.to_thread(docx_to_pdf, data)
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise RenderError("The document renderer is unavailable or timed out.") from exc
    if not pdf.startswith(b"%PDF-"):
        raise RenderError("The document renderer returned an invalid PDF.")
    return pdf
