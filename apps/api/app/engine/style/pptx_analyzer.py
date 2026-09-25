"""Exact, deterministic style extraction from PPTX by reading its OOXML.

No AI is needed here: themes, masters, layouts, placeholders, fonts, colours and geometry are
all explicit in the file. The output feeds both the TeacherStyleProfile and the template builder.
"""

from __future__ import annotations

import hashlib
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from app.engine.pptx_xml import NS, emu_to_pt, read_theme
from app.engine.style.common import (
    ColorTally,
    ContentStats,
    classify_slide,
    clean_font_name,
    fallback_for,
    is_neutral,
    luminance,
)

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.SUBTITLE}
FOOTER_TYPES = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE}

THEME_COLOR_SLOTS = {
    "ACCENT_1": "accent1", "ACCENT_2": "accent2", "ACCENT_3": "accent3", "ACCENT_4": "accent4",
    "ACCENT_5": "accent5", "ACCENT_6": "accent6", "TEXT_1": "dk1", "TEXT_2": "dk2", "BACKGROUND_1": "lt1",
    "BACKGROUND_2": "lt2", "DARK_1": "dk1", "DARK_2": "dk2", "LIGHT_1": "lt1", "LIGHT_2": "lt2",
    "HYPERLINK": "hlink",
}


@dataclass
class ShapeInfo:
    shape_id: int
    kind: str  # text | picture | autoshape | table | chart | group | placeholder
    bbox: tuple[float, float, float, float]  # normalised x, y, w, h
    text: str = ""
    paragraphs: list[str] = field(default_factory=list)
    max_font_pt: float = 0.0
    fonts: Counter = field(default_factory=Counter)
    text_colors: Counter = field(default_factory=Counter)
    fill: str | None = None
    image_hash: str | None = None
    placeholder_type: Any = None
    autoshape: str | None = None
    bold_ratio: float = 0.0
    alignment: str | None = None

    def signature(self) -> tuple:
        r = tuple(round(v * 50) / 50 for v in self.bbox)  # 2% grid
        if self.kind == "picture":
            return ("picture", r, self.image_hash)
        if self.kind == "autoshape" and not self.text.strip():
            return ("autoshape", r, self.fill, self.autoshape)
        if self.text.strip():
            t = self.text.strip()
            if t.isdigit():
                return ("slide_number", r)
            if len(t) <= 60:
                return ("text", r, t)
        return ("unique", self.shape_id, id(self))


@dataclass
class SlideInfo:
    index: int
    layout_index: int
    layout_name: str
    shapes: list[ShapeInfo]
    background: dict[str, Any]
    title_shape: ShapeInfo | None = None
    label: str = "concept"


class PptxAnalyzer:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.prs = Presentation(str(path))
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height
        self.theme = read_theme(self.prs)
        self.theme_colors: dict[str, str] = self.theme["colors"]
        self.theme_fonts: dict[str, str] = self.theme["fonts"]
        self.layouts = list(self.prs.slide_layouts)

    # ------------------------------------------------------------------ colour/font resolution

    def _color_of(self, color_format) -> str | None:
        try:
            ctype = color_format.type
        except Exception:
            return None
        if ctype == MSO_COLOR_TYPE.RGB:
            return "#" + str(color_format.rgb)
        if ctype == MSO_COLOR_TYPE.SCHEME:
            name = getattr(color_format.theme_color, "name", "")
            slot = THEME_COLOR_SLOTS.get(name)
            return self.theme_colors.get(slot) if slot else None
        return None

    def _fill_of(self, shape) -> str | None:
        try:
            fill = shape.fill
            if fill.type == MSO_FILL.SOLID:
                return self._color_of(fill.fore_color)
        except Exception:
            return None
        return None

    def _default_size(self, shape) -> float:
        """Resolve the inherited font size for a shape's text (placeholder -> layout -> master)."""
        el = shape._element
        sz = el.find(".//" + qn("a:lstStyle") + "/" + qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
        if sz is not None and sz.get("sz"):
            return int(sz.get("sz")) / 100
        if shape.is_placeholder:
            ptype = shape.placeholder_format.type
            try:
                layout_ph = shape.part.slide.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
            except Exception:
                layout_ph = None
            if layout_ph is not None:
                d = layout_ph._element.find(".//" + qn("a:lstStyle") + "/" + qn("a:lvl1pPr") + "/" + qn("a:defRPr"))
                if d is not None and d.get("sz"):
                    return int(d.get("sz")) / 100
            style = "titleStyle" if ptype in TITLE_TYPES else "bodyStyle"
        else:
            style = "otherStyle"
        master = self.prs.slide_master._element
        d = master.find(f".//p:txStyles/p:{style}/a:lvl1pPr/a:defRPr", NS)
        if d is not None and d.get("sz"):
            return int(d.get("sz")) / 100
        return 18.0

    def _font_family(self, run, shape) -> str | None:
        name = run.font.name
        if name and not name.startswith("+"):
            return name
        is_title = shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES
        if name == "+mj-lt" or is_title:
            return self.theme_fonts.get("major")
        return self.theme_fonts.get("minor")

    # ------------------------------------------------------------------ shape extraction

    def _norm_bbox(self, shape) -> tuple[float, float, float, float]:
        try:
            left, top, width, height = shape.left, shape.top, shape.width, shape.height
            if shape.is_placeholder and (not width or not height):
                # Broken/partial xfrm: fall back to the master placeholder of the same type.
                ptype = shape.placeholder_format.type
                for mph in self.prs.slide_master.placeholders:
                    if mph.placeholder_format.type == ptype or (
                            ptype in TITLE_TYPES and mph.placeholder_format.type in TITLE_TYPES) or (
                            ptype in BODY_TYPES and mph.placeholder_format.type in BODY_TYPES):
                        left, top, width, height = mph.left, mph.top, mph.width, mph.height
                        break
            return (left / self.W, top / self.H, width / self.W, height / self.H)
        except Exception:
            return (0.0, 0.0, 0.0, 0.0)

    def _shape_info(self, shape) -> ShapeInfo | None:
        if shape.left is None or shape.width is None:
            return None
        info = ShapeInfo(shape_id=shape.shape_id, kind="autoshape", bbox=self._norm_bbox(shape))
        st = shape.shape_type
        if shape.is_placeholder:
            info.placeholder_type = shape.placeholder_format.type
        if st == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, "image")):
            info.kind = "picture"
            try:
                info.image_hash = hashlib.sha1(shape.image.blob).hexdigest()
            except Exception:
                info.image_hash = None
            return info
        if st == MSO_SHAPE_TYPE.GROUP:
            info.kind = "group"
            return info
        if getattr(shape, "has_table", False) and shape.has_table:
            info.kind = "table"
            return info
        if getattr(shape, "has_chart", False) and shape.has_chart:
            info.kind = "chart"
            return info
        if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                info.autoshape = str(shape.auto_shape_type).split(".")[-1].split(" ")[0]
            except Exception:
                info.autoshape = None
            info.fill = self._fill_of(shape)
        if shape.has_text_frame:
            if st != MSO_SHAPE_TYPE.AUTO_SHAPE:
                info.kind = "text"
            default_size = self._default_size(shape)
            paras, bold, total = [], 0, 0
            for p in shape.text_frame.paragraphs:
                ptxt = "".join(r.text for r in p.runs).strip()
                if ptxt:
                    paras.append(ptxt)
                if p.alignment is not None and info.alignment is None:
                    info.alignment = str(p.alignment).split(".")[-1].split(" ")[0].lower()
                for r in p.runs:
                    n = len(r.text.strip())
                    if not n:
                        continue
                    size = r.font.size.pt if r.font.size else (p.font.size.pt if p.font.size else default_size)
                    info.max_font_pt = max(info.max_font_pt, size)
                    fam = self._font_family(r, shape)
                    if fam:
                        info.fonts[(fam, size)] += n
                    col = self._color_of(r.font.color) if r.font.color and r.font.color.type else None
                    if col is None:
                        col = self.theme_colors.get("dk1", "#000000")
                    info.text_colors[col] += n
                    total += n
                    if r.font.bold:
                        bold += n
            info.paragraphs = paras
            info.text = "\n".join(paras)
            info.bold_ratio = bold / total if total else 0.0
        return info

    def _background(self, slide) -> dict[str, Any]:
        for src, el in (("slide", slide._element), ("layout", slide.slide_layout._element),
                        ("master", self.prs.slide_master._element)):
            bg = el.find(qn("p:cSld") + "/" + qn("p:bg"))
            if bg is None:
                continue
            solid = bg.find(".//" + qn("a:solidFill"))
            blip = bg.find(".//" + qn("a:blip"))
            if blip is not None:
                return {"kind": "image", "source": src}
            if solid is not None and len(solid):
                c = solid[0]
                if c.tag == qn("a:srgbClr"):
                    return {"kind": "solid", "color": "#" + c.get("val").upper(), "source": src}
                if c.tag == qn("a:schemeClr"):
                    slot = {"bg1": "lt1", "bg2": "lt2", "tx1": "dk1", "tx2": "dk2"}.get(c.get("val"), c.get("val"))
                    return {"kind": "solid", "color": self.theme_colors.get(slot, "#FFFFFF"), "source": src}
            bgref = bg.find(qn("p:bgRef"))
            if bgref is not None:
                return {"kind": "theme", "color": self.theme_colors.get("lt1", "#FFFFFF"), "source": src}
        return {"kind": "solid", "color": self.theme_colors.get("lt1", "#FFFFFF"), "source": "default"}

    def _slides(self) -> list[SlideInfo]:
        out = []
        for i, slide in enumerate(self.prs.slides):
            shapes = []
            for sh in slide.shapes:
                info = self._shape_info(sh)
                if info:
                    shapes.append(info)
            layout = slide.slide_layout
            out.append(SlideInfo(index=i, layout_index=self.layouts.index(layout), layout_name=layout.name,
                                 shapes=shapes, background=self._background(slide)))
        return out

    # ------------------------------------------------------------------ analysis

    @staticmethod
    def _pick_title(slide: SlideInfo, decoration_sigs: set) -> ShapeInfo | None:
        for s in slide.shapes:
            if s.placeholder_type in TITLE_TYPES and s.text.strip():
                return s
        candidates = [s for s in slide.shapes if s.text.strip() and s.signature() not in decoration_sigs
                      and s.bbox[1] < 0.35 and s.kind in ("text", "autoshape", "placeholder")]
        if not candidates:
            candidates = [s for s in slide.shapes if s.text.strip() and s.signature() not in decoration_sigs]
        if not candidates:
            return None
        return max(candidates, key=lambda s: (s.max_font_pt, -s.bbox[1]))

    def analyze(self) -> dict[str, Any]:
        slides = self._slides()
        n = len(slides)
        content_slides = slides[1:] if n > 1 else slides

        # ---- recurring decorations (logos, header bands, footers, background pictures)
        sig_counts: Counter = Counter()
        for s in content_slides:
            for sig in {sh.signature() for sh in s.shapes}:
                if sig[0] != "unique":
                    sig_counts[sig] += 1
        threshold = max(2, int(0.5 * len(content_slides) + 0.5))
        decoration_sigs = {sig for sig, c in sig_counts.items() if c >= threshold}

        # ---- per-slide title + label
        stats = ContentStats()
        for s in slides:
            s.title_shape = self._pick_title(s, decoration_sigs)
            body_shapes = [sh for sh in s.shapes if sh is not s.title_shape and sh.signature() not in decoration_sigs
                           and sh.text.strip()]
            texts = [p for sh in body_shapes for p in sh.paragraphs]
            left_cols = [sh for sh in body_shapes if sh.bbox[0] + sh.bbox[2] <= 0.55]
            right_cols = [sh for sh in body_shapes if sh.bbox[0] >= 0.45]
            s.label = classify_slide(
                index=s.index, title=s.title_shape.text if s.title_shape else "", texts=texts,
                n_text_blocks=len(body_shapes),
                n_pictures=sum(1 for sh in s.shapes if sh.kind == "picture" and sh.signature() not in decoration_sigs),
                n_tables=sum(1 for sh in s.shapes if sh.kind == "table"),
                n_charts=sum(1 for sh in s.shapes if sh.kind == "chart"),
                has_two_columns=bool(left_cols and right_cols and len(body_shapes) >= 2),
                is_first=s.index == 0 and n > 1,
            )
            if s.index > 0:
                if s.title_shape:
                    stats.titles.append(s.title_shape.text.replace("\n", " "))
                bullets = [t for t in texts if len(t.split()) >= 1]
                stats.bullets_per_slide.append(len(bullets))
                stats.words_per_bullet.extend(len(t.split()) for t in bullets)
                stats.body_text.extend(bullets)

        # ---- colours
        fills, text_cols, bgs = ColorTally(), ColorTally(), ColorTally()
        for s in slides:
            if s.background.get("color"):
                bgs.add(s.background["color"], 1)
            for sh in s.shapes:
                area = sh.bbox[2] * sh.bbox[3]
                if sh.fill:
                    fills.add(sh.fill, area)
                for col, cnt in sh.text_colors.items():
                    text_cols.add(col, cnt)
        background = (bgs.ranked() or [(self.theme_colors.get("lt1", "#FFFFFF"), 1.0)])[0][0]
        fill_rank = fills.ranked()
        brand = [c for c, _ in fill_rank if not is_neutral(c)]
        theme_accents = [self.theme_colors[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5",
                                                         "accent6") if k in self.theme_colors]
        primary = brand[0] if brand else (theme_accents[0] if theme_accents else "#1D4ED8")
        secondary = next((c for c in brand[1:] if c != primary),
                         next((c for c in theme_accents if c.upper() != primary.upper()), "#F59E0B"))

        # title / body text colours & fonts
        title_fonts: Counter = Counter()
        body_fonts: Counter = Counter()
        title_sizes, body_sizes, title_colors, body_colors = [], [], ColorTally(), ColorTally()
        title_bold = []
        for s in content_slides:
            for sh in s.shapes:
                if sh.signature() in decoration_sigs or not sh.text.strip():
                    continue
                target_fonts = title_fonts if sh is s.title_shape else body_fonts
                for (fam, size), cnt in sh.fonts.items():
                    target_fonts[fam] += cnt
                    (title_sizes if sh is s.title_shape else body_sizes).extend([size] * min(cnt, 50))
                for col, cnt in sh.text_colors.items():
                    (title_colors if sh is s.title_shape else body_colors).add(col, cnt)
                if sh is s.title_shape:
                    title_bold.append(sh.bold_ratio)
        heading_font = (title_fonts.most_common(1)[0][0] if title_fonts else self.theme_fonts.get("major")) or "Arial"
        body_font = (body_fonts.most_common(1)[0][0] if body_fonts else self.theme_fonts.get("minor")) or "Arial"
        title_pt = round(statistics.median(title_sizes), 1) if title_sizes else 36.0
        body_pt = round(statistics.median(body_sizes), 1) if body_sizes else 20.0
        text_color = (body_colors.ranked() or [(self.theme_colors.get("dk1", "#1F2937"), 1)])[0][0]
        title_color = (title_colors.ranked() or [(text_color, 1)])[0][0]

        # ---- zones (normalised) from content slides
        def median_box(boxes: list[tuple[float, float, float, float]]):
            if not boxes:
                return None
            xs = statistics.median(b[0] for b in boxes)
            ys = statistics.median(b[1] for b in boxes)
            x2 = statistics.median(b[0] + b[2] for b in boxes)
            y2 = statistics.median(b[1] + b[3] for b in boxes)
            return [round(xs, 4), round(ys, 4), round(x2 - xs, 4), round(y2 - ys, 4)]

        title_boxes, body_boxes, image_boxes = [], [], []
        for s in content_slides:
            if s.label in ("section",):
                continue
            if s.title_shape:
                title_boxes.append(s.title_shape.bbox)
            others = [sh for sh in s.shapes if sh is not s.title_shape and sh.signature() not in decoration_sigs
                      and sh.kind in ("text", "autoshape", "picture", "table", "chart", "group")
                      and sh.bbox[2] * sh.bbox[3] > 0.01 and sh.bbox[2] < 0.98]
            if s.title_shape and others:
                x = min(o.bbox[0] for o in others)
                y = min(o.bbox[1] for o in others)
                x2 = max(o.bbox[0] + o.bbox[2] for o in others)
                y2 = max(o.bbox[1] + o.bbox[3] for o in others)
                body_boxes.append((x, y, x2 - x, y2 - y))
            for sh in others:
                if sh.kind == "picture":
                    image_boxes.append(sh.bbox)
        title_zone = median_box(title_boxes)
        body_zone = median_box(body_boxes)
        image_zone = median_box(image_boxes)

        # ---- decorations: donor slides and shape ids to replicate
        cover = slides[0] if slides and slides[0].label == "cover" else None
        content_donor = None
        best = -1
        for s in content_slides:
            present = sum(1 for sh in s.shapes if sh.signature() in decoration_sigs)
            if present > best:
                best, content_donor = present, s
        decorations: dict[str, Any] = {"content": None, "cover": None}
        if content_donor is not None and best > 0:
            items = []
            for sh in content_donor.shapes:
                sig = sh.signature()
                if sig in decoration_sigs:
                    items.append({"shape_id": sh.shape_id, "kind": sig[0], "bbox": list(sh.bbox),
                                  "fill": sh.fill, "text": sh.text if sig[0] == "text" else None})
            decorations["content"] = {"donor_slide": content_donor.index, "items": items,
                                      "background": content_donor.background}
        if cover is not None:
            items = []
            for sh in cover.shapes:
                if sh.text.strip():
                    continue
                items.append({"shape_id": sh.shape_id, "kind": sh.kind, "bbox": list(sh.bbox), "fill": sh.fill})
            decorations["cover"] = {"donor_slide": 0, "items": items, "background": cover.background}

        # cover text zones
        cover_zone = None
        if cover is not None:
            texts = sorted([sh for sh in cover.shapes if sh.text.strip()], key=lambda s: -s.max_font_pt)
            if texts:
                cover_zone = {
                    "title": list(texts[0].bbox),
                    "title_pt": texts[0].max_font_pt,
                    "title_color": (texts[0].text_colors.most_common(1) or [[title_color]])[0][0],
                    "subtitle": list(texts[1].bbox) if len(texts) > 1 else None,
                    "subtitle_pt": texts[1].max_font_pt if len(texts) > 1 else None,
                    "subtitle_color": (texts[1].text_colors.most_common(1) or [[secondary]])[0][0]
                    if len(texts) > 1 else None,
                    "align": texts[0].alignment or "left",
                }

        # ---- layout usage and placeholder-capable layouts
        layout_usage: dict[int, Counter] = defaultdict(Counter)
        for s in slides:
            layout_usage[s.layout_index][s.label] += 1
        layouts_meta = []
        for li, layout in enumerate(self.layouts):
            phs = []
            for ph in layout.placeholders:
                t = ph.placeholder_format.type
                phs.append({
                    "idx": ph.placeholder_format.idx,
                    "type": "title" if t in TITLE_TYPES else "body" if t in BODY_TYPES else
                    "picture" if t == PP_PLACEHOLDER.PICTURE else "footer" if t in FOOTER_TYPES else "other",
                    "bbox": list(self._norm_bbox(ph)),
                })
            layouts_meta.append({"index": li, "name": layout.name, "placeholders": phs,
                                 "used_by": dict(layout_usage.get(li, {}))})

        content_layout_index = Counter(s.layout_index for s in content_slides).most_common(1)[0][0] \
            if content_slides else 0
        cover_layout_index = cover.layout_index if cover else 0
        section_layout_index = next((s.layout_index for s in slides if s.label == "section"), None)

        corner_shapes = [sh.autoshape for s in slides for sh in s.shapes if sh.autoshape]
        rounded = sum(1 for a in corner_shapes if a and "ROUNDED" in a.upper())
        uses_images = sum(1 for s in content_slides
                          for sh in s.shapes if sh.kind == "picture" and sh.signature() not in decoration_sigs)

        return {
            "source_kind": "pptx",
            "slide_size": {"w": int(self.W), "h": int(self.H)},
            "theme": self.theme,
            "colors": {
                "primary": primary,
                "secondary": secondary,
                "accents": [c for c, _ in fill_rank[:6]] or theme_accents,
                "background": background,
                "text": text_color,
                "title": title_color,
                "palette": fill_rank[:8],
                "dark_background": luminance(background) < 0.25,
            },
            "fonts": {
                "heading": {"family": clean_font_name(heading_font) or "Arial", "fallback": fallback_for(heading_font)},
                "body": {"family": clean_font_name(body_font) or "Arial", "fallback": fallback_for(body_font)},
                "arabic": {"family": self.theme_fonts.get("minor_cs") or "Noto Sans Arabic",
                           "fallback": "Noto Sans Arabic"},
            },
            "typography": {
                "title_pt": title_pt,
                "body_pt": body_pt,
                "min_pt": max(14.0, round(min(body_sizes), 1)) if body_sizes else 16.0,
                "title_bold": (statistics.mean(title_bold) > 0.5) if title_bold else True,
            },
            "zones": {"title": title_zone, "body": body_zone, "image": image_zone,
                      "image_side": ("right" if image_zone and image_zone[0] + image_zone[2] / 2 > 0.5 else "left")
                      if image_zone else "right",
                      "cover": cover_zone},
            "layouts": layouts_meta,
            "layout_map": {
                "content_layout_index": content_layout_index,
                "cover_layout_index": cover_layout_index,
                "section_layout_index": section_layout_index,
            },
            "decorations": decorations,
            "layouts_found": [{"key": k, "count": c} for k, c in
                              Counter(s.label for s in slides).most_common()],
            "slide_labels": [s.label for s in slides],
            "content_style": stats.summary(),
            "visual_rules": {
                "corner_radius": "rounded" if corner_shapes and rounded / len(corner_shapes) >= 0.4 else "square",
                "uses_images": uses_images > 0,
                "image_slides_ratio": round(uses_images / max(1, len(content_slides)), 2),
                "title_align": (next((s.title_shape.alignment for s in content_slides
                                      if s.title_shape and s.title_shape.alignment), None) or "left"),
            },
            "stats": {"slides": n, "decoration_count": len(decoration_sigs)},
        }


def analyze_pptx(path: Path) -> dict[str, Any]:
    return PptxAnalyzer(path).analyze()
