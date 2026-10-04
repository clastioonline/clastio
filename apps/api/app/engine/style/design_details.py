"""Inspect every slide for reusable native table and chart formatting, without data."""
from __future__ import annotations

import json
from collections import Counter
from copy import deepcopy

from lxml import etree
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from app.engine.style.content import walk_shapes


def _xml(element):
    # Relationship-backed fills/links need their source part; never transplant dangling IDs.
    if element is not None and any(key.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")
                                   for node in element.iter() for key in node.attrib):
        return None
    return etree.tostring(element, encoding="unicode") if element is not None else None


def inspect_design(prs) -> dict:
    inventory, tables, charts = [], [], {}
    for number, slide in enumerate(prs.slides, 1):
        shapes = list(walk_shapes(slide.shapes))
        fonts = sorted({run.font.name for shape in shapes if shape.has_text_frame
                        for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.name})
        row = {"slide": number, "layout": slide.slide_layout.name, "fonts": fonts,
               "tables": 0, "charts": 0, "pictures": sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in shapes)}
        for shape in shapes:
            if shape.has_table:
                row["tables"] += 1
                table = shape.table
                # Include every body row in inspection; retain the most frequent cell styling.
                def cell_style(cell):
                    properties = cell._tc.find(qn("a:tcPr"))
                    return [xml for child in properties if etree.QName(child).localname in
                            {"solidFill", "gradFill", "noFill", "pattFill", "lnL", "lnR", "lnT", "lnB"}
                            and (xml := _xml(child)) is not None] if properties is not None else []
                header = cell_style(table.cell(0, 0))
                body = [tuple(cell_style(cell)) for tr in list(table.rows)[1:] for cell in tr.cells]
                tables.append({"header": header, "body": list(Counter(body).most_common(1)[0][0]) if body else header,
                               "header_text": _xml(table.cell(0, 0)._tc.find('.//' + qn('a:rPr'))),
                               "body_text": _xml(table.cell(1 if len(table.rows) > 1 else 0, 0)._tc.find('.//' + qn('a:rPr'))),
                               "properties": _xml(shape._element.find('.//' + qn('a:tblPr')))})
            if shape.has_chart:
                row["charts"] += 1
                chart = shape.chart
                kind = {XL_CHART_TYPE.COLUMN_CLUSTERED: "bar", XL_CHART_TYPE.LINE_MARKERS: "line",
                        XL_CHART_TYPE.LINE: "line", XL_CHART_TYPE.PIE: "pie"}.get(chart.chart_type)
                if kind:
                    charts.setdefault(kind, []).append({
                        "style": chart.chart_style,
                        "series": [_xml(series._element.find(qn("c:spPr"))) for series in chart.series]})
        inventory.append(row)
    def dominant(values):
        counts = Counter(json.dumps(value, sort_keys=True) for value in values)
        return max(values, key=lambda value: counts[json.dumps(value, sort_keys=True)])
    return {"slides_inspected": len(inventory), "slides": inventory,
            "table": dominant(tables) if tables else None,
            "charts": {kind: dominant(values) for kind, values in charts.items()}}


def apply_table_style(table, style: dict) -> None:
    if style.get("properties"):
        properties = table._tbl.find(qn("a:tblPr"))
        original = etree.fromstring(style["properties"].encode())
        for key in ("firstRow", "lastRow", "firstCol", "lastCol", "bandRow", "bandCol"):
            if key in original.attrib:
                properties.set(key, original.get(key))
        identifier = original.find(qn("a:tableStyleId"))
        if identifier is not None:
            for child in list(properties):
                if child.tag == identifier.tag:
                    properties.remove(child)
            properties.append(deepcopy(identifier))
    for index, row in enumerate(table.rows):
        formats = style.get("header" if index == 0 else "body") or []
        for cell in row.cells:
            properties = cell._tc.get_or_add_tcPr()
            for xml in formats:
                element = etree.fromstring(xml.encode())
                for child in list(properties):
                    if child.tag == element.tag or (etree.QName(element).localname.endswith("Fill") and
                                                   etree.QName(child).localname.endswith("Fill")):
                        properties.remove(child)
                properties.append(element)
            text_xml = style.get("header_text" if index == 0 else "body_text")
            if text_xml:
                original = etree.fromstring(text_xml.encode())
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        properties = run._r.get_or_add_rPr()
                        for key in ("b", "i", "u"):
                            if key in original.attrib:
                                properties.set(key, original.get(key))
                        for child in original:
                            if etree.QName(child).localname not in {"solidFill", "latin", "ea", "cs"}:
                                continue
                            for current in list(properties):
                                if current.tag == child.tag:
                                    properties.remove(current)
                            properties.append(deepcopy(child))


def apply_chart_style(chart, style: dict) -> None:
    if style.get("style") is not None:
        chart.chart_style = style["style"]
    formats = style.get("series") or []
    for index, series in enumerate(chart.series):
        if index < len(formats) and formats[index]:
            old = series._element.find(qn("c:spPr"))
            if old is not None:
                series._element.remove(old)
            element = etree.fromstring(formats[index].encode())
            # c:spPr follows the series name and precedes values/markers.
            position = 2 + int(series._element.find(qn("c:tx")) is not None)
            series._element.insert(position, element)
