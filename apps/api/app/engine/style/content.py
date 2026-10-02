"""Read editable slide content, including groups, tables, notes and inherited text.

Picture bytes remain available for reuse; text inside a picture requires visual inspection.
"""
from __future__ import annotations

import hashlib
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def walk_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk_shapes(shape.shapes)
        else:
            yield shape


def shape_image_bytes(shape) -> bytes | None:
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, "image")):
        return shape.image.blob
    blip = shape._element.find("./" + qn("p:spPr") + "/" + qn("a:blipFill") + "/" + qn("a:blip"))
    if blip is not None and blip.get(qn("r:embed")):
        return shape.part.related_part(blip.get(qn("r:embed"))).blob
    return None


def extract_content(prs) -> tuple[list[dict[str, Any]], dict[str, bytes]]:
    pages, images = [], {}
    for number, slide in enumerate(prs.slides, 1):
        shapes = list(walk_shapes(slide.shapes))
        if slide._element.get("showMasterSp", "1") not in ("0", "false"):
            shapes.extend(walk_shapes(sh for sh in slide.slide_layout.shapes if not sh.is_placeholder))
            if slide.slide_layout._element.get("showMasterSp", "1") not in ("0", "false"):
                shapes.extend(walk_shapes(sh for sh in slide.slide_layout.slide_master.shapes
                                          if not sh.is_placeholder))
        texts, pictures, seen = [], [], set()
        for shape in shapes:
            text = ""
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
            elif shape.has_table:
                text = "\n".join(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
            # Narrow sidebar labels are navigation, not teaching material.
            sidebar = shape.left is not None and shape.width is not None and \
                shape.left < prs.slide_width * .05 and shape.width < prs.slide_width * .2
            if text and text not in seen and not sidebar:
                texts.append(text)
                seen.add(text)
            data = shape_image_bytes(shape)
            if data is not None:
                digest = hashlib.sha256(data).hexdigest()
                images[digest] = data
                props = shape._element.find(".//" + qn("p:cNvPr"))
                pictures.append({"image_key": digest, "name": shape.name,
                                 "description": props.get("descr", "") if props is not None else "",
                                 "bbox": [shape.left / prs.slide_width, shape.top / prs.slide_height,
                                          shape.width / prs.slide_width, shape.height / prs.slide_height]})
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        pages.append({"number": number, "layout_name": slide.slide_layout.name,
                      "text": "\n".join(texts), "notes": notes, "pictures": pictures})
    return pages, images
