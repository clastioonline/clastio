"""Recognize small, isolated slide numbers without treating teaching numerals as chrome."""
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn


def is_footer_number_text(text, bbox, font_pt=0):
    text = str(text or '').strip()
    x, y, width, height = bbox
    return (text.isdigit() and len(text) <= 4 and y >= .85 and width <= .18 and height <= .10
            and (x <= .18 or x + width >= .82) and font_pt <= 18)


def is_footer_number_shape(shape, slide_width, slide_height):
    if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
        return True
    if not shape.has_text_frame:
        return False
    if any(field.get('type') == 'slidenum' for field in shape._element.iter(qn('a:fld'))):
        return True
    sizes = [run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
             if run.font.size is not None]
    bbox = (shape.left / slide_width, shape.top / slide_height,
            shape.width / slide_width, shape.height / slide_height)
    return is_footer_number_text(shape.text, bbox, max(sizes, default=0))
