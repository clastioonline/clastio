"""Local, unpaid review of the supplied Grade 1 deck using its native design.

Original pictures with inconsistent object counts are replaced by editable counting diagrams.
The source upload and the earlier live-review artifacts remain untouched.
"""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation

from app.engine.exports.documents import assessment_docx
from app.engine.qc.visual import docx_to_pdf, inspect
from app.engine.render.renderer import DeckRenderer
from app.engine.style.content import shape_image_bytes
from app.engine.style.pptx_analyzer import analyze_pptx
from app.engine.template.builder import build_native
from app.generation.specs import AssessmentDoc, Bullet, CountingGroup, SlideSpec, Visual

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "apps/web/e2e/screenshots/extraction-improvement"


def main():
    source = next(ROOT.glob("W02-G01*.pptx"))
    prs = Presentation(source)
    base, template = build_native(source, analyze_pptx(source))
    images = {}

    def original_picture(page, shape_id):
        shape = next(shape for shape in prs.slides[page - 1].shapes if shape.shape_id == shape_id)
        data = shape_image_bytes(shape)
        if data is None:
            raise ValueError("Expected a source picture or picture fill")
        digest = hashlib.sha256(data).hexdigest()
        images[digest] = data
        return digest

    flag = original_picture(2, 13)
    reflection = original_picture(10, 6)

    def counted(first, second, show_total=True):
        return Visual(kind="diagram", description="Two parts put together",
                      counting_groups=[CountingGroup(count=first, label=f"{first} red cubes", color="red"),
                                       CountingGroup(count=second, label=f"{second} green cubes", color="green")],
                      show_total=show_total)

    def slide(number, title, stage, minutes, lines, notes, visual=None, asset_id=None):
        return SlideSpec(number=number, layout="image_text" if visual else "concept", purpose=stage,
                         title=title, teaching_stage=stage, timing_minutes=minutes,
                         bullets=[Bullet(text=line) for line in lines], speaker_notes=notes,
                         visual=visual or Visual(), asset_id=asset_id)

    slides = [
        SlideSpec(number=1, layout="cover", title="Put Together: Addition", purpose="topic",
                  teaching_stage="topic", subtitle="Grade 1 Mathematics • Two parts make a whole",
                  timing_minutes=1, speaker_notes="Introduce the lesson. We will join two parts and count the whole."),
        slide(2, "Look, Think, Share", "engage", 4,
              ["Which flag is this?", "Name its four colours.", "We will build with coloured cubes."],
              "The UAE flag is red, green, white and black. Count the four colours together. "
              "Connect the familiar colours to today's cube models; colours are not numbers to add.",
              Visual(kind="image", source_image_key=flag, description="The UAE flag from the original deck"), flag),
        slide(3, "Two Parts Make a Whole", "objective", 3,
              ["I can find the two parts.", "I can put the parts together.", "I can count the whole.",
               "I can write an addition equation."],
              "Use the original learning objective and success criteria. This lesson practises putting together "
              "within 10. The teacher supplied standard 1.OA.A.1 covers a broader range within 20. "
              "Support: count concrete objects. Extension: explain how the two parts relate to the whole."),
        slide(4, "Put Together 4 and 4", "explore", 7,
              ["Make a part with 4 red cubes.", "Make a part with 4 green cubes.", "Put the parts together.",
               "Count all 8 cubes: 4 + 4 = 8."],
              "Adapt the original written apple problem (4 red and 4 green) into an accurate cube model. "
              "Its source picture disagrees with the written quantities, so it is not reused. "
              "Ask pupils to identify each part before counting the whole. Answer: 8.", counted(4, 4)),
        slide(5, "Parts 4 and 2, Whole 6", "explain", 6,
              ["The parts are 4 and 2.", "Put both parts together.", "The whole is 6.", "Write: 4 + 2 = 6."],
              "Model the original parts-and-whole example with exactly four red cubes and two green cubes. "
              "Count six altogether. The original fish illustration has inconsistent counts, so use this "
              "editable diagram. Ask: Which numbers name the parts? Which number names the whole?", counted(4, 2)),
        slide(6, "Our Addition Words", "explain", 3,
              ["Parts: the groups we put together.", "Whole: all the objects together.", "Plus (+): we add.",
               "Equals (=): both sides have the same value."],
              "Use 4 + 2 = 6 to explain the symbols. Four and two are the parts; six is the whole. "
              "Say 'four plus two equals six'. Equality means the value on the left matches the value on the right."),
        slide(7, "Build, Turn, Compare, Explain", "elaborate", 7,
              ["Build 4 red cubes and 2 green cubes.", "Count the whole: 6.", "Swap the two groups.",
               "2 + 4 is also 6. Explain why."],
              "Pairs build 4 + 2, then swap the groups to show 2 + 4. No cubes are added or removed, "
              "so the whole stays six. This follows the original comparison activity without its inaccurate "
              "picture. Then choose two new parts with a combined total at most 10. "
              "Support: teacher supplies the groups. Extension: explain why swapping parts keeps the total.", counted(4, 2)),
        slide(8, "Build, Write, Say", "elaborate", 7,
              ["Build a group of 3 cubes.", "Build another group of 5 cubes.", "Put the parts together.",
               "Write the equation. Say the whole."],
              "Guided practice adapts the original 3-and-5 example. Answer: 3 + 5 = 8. "
              "Ask pupils to explain: 'I had three and five. I put them together. The whole is eight.' "
              "Follow-up: 1 white cube and 6 green cubes gives 1 + 6 = 7.", counted(3, 5, False)),
        slide(9, "Exit Ticket: Show What You Know", "evaluate", 4,
              ["Count 3 red cubes and 2 green cubes.", "How many cubes are there altogether?",
               "Write: 3 + 2 = __.", "Check by counting all the cubes."],
              "Each pupil answers independently. Expected answer: 3 + 2 = 5. "
              "Look for identifying both parts, counting the whole and matching the equation. "
              "Provide physical cubes for pupils needing support.", counted(3, 2, False)),
        slide(10, "Self-Reflect", "self_reflect", 3,
              ["Can I find both parts?", "Can I put them together?", "Can I count the whole?",
               "Choose: I got it, getting there, or more practice."],
              "Use the original self-reflection visual. Let pupils choose honestly; each choice is useful. "
              "Plan another concrete model for pupils who need practice. "
              "The original picture is retained, with its full proportions and no cropping.",
              Visual(kind="image", source_image_key=reflection,
                     description="Original three-choice self-reflection poster"), reflection),
    ]
    assert sum(s.timing_minutes for s in slides) == 45
    renderer = DeckRenderer(base, template)
    data = renderer.render(slides, images, core_props={"title": "Put Together: Addition — Revised Review",
                                                     "subject": "Grade 1 Mathematics"})
    report = inspect(data)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "revised-lesson.pptx").write_bytes(data)
    (OUT / "revised-lesson.pdf").write_bytes(report.pdf)
    (OUT / "revised-slides.json").write_text(json.dumps([s.model_dump() for s in slides], indent=2))
    worksheet_data = json.loads((ROOT / "apps/web/e2e/screenshots/live-openai-review/worksheet.json").read_text())
    worksheet = AssessmentDoc.model_validate(worksheet_data["content"])
    worksheet.title = "Put Together: Addition Within 10"
    for section in worksheet.sections:
        for question in section.questions:
            question.options = [option for option in question.options if option.strip()]
    worksheet.sections[0].questions[-1].stem = "What is the total when you add 5 and 4?"
    for name, key in (("revised-worksheet", False), ("revised-answer-key", True)):
        docx = assessment_docx(worksheet, subtitle="Grade 1 Mathematics • Addition within 10",
                               accent="#408332", font="Arial", key=key)
        (OUT / f"{name}.docx").write_bytes(docx)
        (OUT / f"{name}.pdf").write_bytes(docx_to_pdf(docx))
    for number, png in enumerate(report.previews, 1):
        (OUT / f"revised-{number}.png").write_bytes(png)
    document = pymupdf.open(stream=report.pdf, filetype="pdf")
    canvas = Image.new("RGB", (1200, 338 * 5), "white")
    for index, page in enumerate(document):
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(.625, .625))
        image = Image.open(io.BytesIO(pixmap.tobytes("png"))).resize((600, 338))
        canvas.paste(image, ((index % 2) * 600, (index // 2) * 338))
    canvas.save(OUT / "revised-contact-sheet.png")
    summary = {"provider_calls": 0, "slides": len(slides), "duration_minutes": 45,
               "source_pictures_reused": 2, "editable_counting_diagrams": 5,
               "visual_issues": report.issues, "failing_slides": report.failing_slides,
               "text_overflow": [r.number for r in renderer.reports if r.overflow],
               "source_picture_counts_checked": True}
    (OUT / "revised-review.json").write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary))


if __name__ == "__main__":
    main()
